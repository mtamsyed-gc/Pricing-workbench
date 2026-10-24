"""The structure-only cedant example in examples/ (SPEC §8.4). Expected counts are read from the example folder so
no cedant labels appear in tests/."""
import copy
import re

from _util import REPO, load_json, tmpdir
from openpyxl import load_workbook

from pw.gates import v4_round_trip
from pw.template.generator import generate_template
from pw.template.parser import fill_template, parse_submission, to_catalogue
from pw.template.spec import load_template_spec, normalise

EX = REPO / "examples" / "ufic_2027"


def _check(ns, expected):
    sheets = {s["sheet_name"]: s for s in ns["sheets"]}
    assert set(sheets) == set(expected["sheets"])
    for name, exp in expected["sheets"].items():
        blocks = sheets[name]["blocks"]
        if "blocks" in exp:
            assert [b["tpa"] for b in blocks] == list(exp["blocks"]), name          # one block per TPA, in order
            for b in blocks:
                assert [len(b["plans"]), len(b["rows"]), len(b["columns"])] == exp["blocks"][b["tpa"]], (name, b["tpa"])
        if "blocks_by_plan" in exp:
            for b in blocks:
                pid = b["plans"][0]["plan_id"]
                assert [len(b["plans"]), len(b["rows"]), len(b["columns"])] == exp["blocks_by_plan"][pid], (name, pid)
            assert sum(len(b["refer_cells"]) for b in blocks) == exp["refer_cells"]
        if "member_categories" in exp:
            secs = {r["section"] for b in blocks for r in b["rows"]}
            assert len(secs) == exp["member_categories"]
            per = {}
            for r in blocks[0]["rows"]:
                per[r["section"]] = per.get(r["section"], 0) + 1
            assert max(per.values()) == exp["max_bands_per_category"]


def test_ufic_structure_matches_spec_8_4():
    spec = load_template_spec(EX / "template_spec.json")
    ns = normalise(spec)
    _check(ns, load_json(EX / "expected_structure.json"))
    # every rate block belongs to exactly one TPA (no shared tables); owner-confirmed bands: no placeholders left
    for s in ns["sheets"]:
        for b in s["blocks"]:
            assert isinstance(b["tpa"], str)
    assert not any(r["placeholder"] for s in ns["sheets"] for b in s["blocks"] for r in b["rows"])
    # plan letters line up across TPAs within each product x region (master plan table); networks differ per TPA
    for s in ns["sheets"]:
        ids = [[p["plan_id"] for p in b["plans"]] for b in s["blocks"] if len(b["plans"]) > 1]
        for a in ids:
            assert a == sorted(a) and set(a) <= set(max(ids, key=len)), s["sheet_name"]


def test_ufic_splits_option_on():
    spec = load_template_spec(EX / "template_spec.json")
    exp = load_json(EX / "expected_structure.json")["splits_on"]
    on = copy.deepcopy(spec)
    for s in on["sheets"]:
        if s["sheet_name"] in exp:
            for b in s["blocks"]:
                if b["tpa"] in exp[s["sheet_name"]]:
                    b["splits"] = {"marital": True, "member_section": True}
    ns = normalise(on)
    for s in ns["sheets"]:
        if s["sheet_name"] in exp:
            for b in s["blocks"]:
                if b["tpa"] in exp[s["sheet_name"]]:
                    assert [len(b["plans"]), len(b["rows"]), len(b["columns"])] == exp[s["sheet_name"]][b["tpa"]]
    with tmpdir() as d:
        assert v4_round_trip(on, d)["status"] == "PASS"


def test_ufic_v4_round_trip_and_generated_workbook():
    spec = load_template_spec(EX / "template_spec.json")
    with tmpdir() as d:
        r = v4_round_trip(spec, d)
        assert r["status"] == "PASS" and r["hash_ok"]
        g = generate_template(spec, d / "t.xlsx")
        wb = load_workbook(d / "t.xlsx")
        assert set(load_json(EX / "expected_structure.json")["sheets"]) <= set(wb.sheetnames)
        text = " ".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v)
        forbidden_incumbent = load_json(EX / "expected_structure.json")["forbidden_tokens"][1]
        assert not re.search(r"\b" + re.escape(forbidden_incumbent) + r"\b", text)   # F8: never name the incumbent
        assert "lead reinsurer" in text
        assert g.n_refer == 2
        fill_template(d / "t.xlsx", d / "f.xlsx")
        p = parse_submission(d / "f.xlsx", template_spec=spec)
        cat = to_catalogue(p, spec, "SUB", "synthetic")                          # every band has a range now
        assert len(cat.scopes) == sum(len(s["blocks"]) for s in normalise(spec)["sheets"])
