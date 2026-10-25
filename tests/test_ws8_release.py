"""WS8 release hardening: tender folders in the release (blank + synthetic training tender, step files), readable
Excel views (column widths and number formats), one laptop check at a time."""
import importlib.util
import json
import os
import re
import subprocess
import sys
import time
import zipfile

import pandas as pd
from _util import REPO, load_json, requires, tmpdir
from openpyxl import load_workbook

from pw import __version__
from pw.outputs import view_layout, write_table
from pw.tender_kit import STEPS, TEMPLATE_NAME, TRAINING_NAME, step_args, template_files, training_files, write_folder


def _laptop_check():
    spec = importlib.util.spec_from_file_location("laptop_check_ws8", REPO / "tools" / "laptop_check.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_tender_folders_hold_the_step_files_and_inputs():
    blank, train = template_files(), training_files()
    for files in (blank, train):
        for f, _, args in STEPS:
            bat = files[f].decode("ascii")
            assert "\r\n" in bat and "pricing-workbench-*" in bat and "py -3.12" in bat and "pause" in bat
            assert f'%PY% "%PWDIR%\\tools\\pw.py" {args}\r\n' in bat
        assert "README.txt" in files and "run_config.json" in files
    for d in ("inputs", "exposure", "exposure/maps", "submissions", "treaty"):
        assert f"{d}/README.txt" in blank
    for need in ("inputs/template_spec.json", "inputs/cedant_filled.xlsx", "inputs/scorecard_weights.json",
                 "exposure/dim_life.csv", "exposure/run_manifest.csv", "exposure/maps/plan_map.csv",
                 "submissions/SUB_X.xlsx", "submissions/SUB_Y.xlsx", "treaty/treaty_book_map.csv", "terms_scores.csv"):
        assert need in train, need
    cfg = json.loads(train["run_config.json"])
    # every path the training config names is in the folder (or written by steps 2 / 3 into work/)
    named = [cfg["expiring"], *cfg["exposure"].values(), cfg["treaty"]["book_map"], cfg["treaty"]["experience"],
             cfg["scorecard"]["weights_file"], cfg["scorecard"]["terms_scores"],
             *(s[k] for s in cfg["submissions"] for k in ("file", "template_spec", "issued"))]
    for p in named:
        if p in (cfg["exposure"]["snapshot_date"],) or p.startswith("work/"):
            continue
        assert p in train or any(k.startswith(p + "/") for k in train), p
    assert step_args(3) == ["run", "run_config.json", "--runs-root", "runs"]


def test_laptop_check_reads_the_step_commands_from_the_bat_files():
    lc = _laptop_check()
    with tmpdir() as d:
        write_folder(d / "T")
        assert lc.step_commands(d / "T") == [[x.replace("\\", os.sep) for x in a.split()] for _, _, a in STEPS]


def test_write_folder_refuses_a_non_empty_folder():
    with tmpdir() as d:
        write_folder(d / "T")
        try:
            write_folder(d / "T")
        except Exception as e:  # noqa: BLE001
            assert "not empty" in str(e)
        else:
            raise AssertionError("expected a refusal")


@requires("examples/ufic_2027")
def test_tender_folders_carry_no_cedant_tokens():
    tokens = load_json(REPO / "examples" / "ufic_2027" / "expected_structure.json")["forbidden_tokens"]
    pats = [(t, re.compile(r"(?<![A-Za-z0-9])" + re.escape(t) + r"(?![A-Za-z0-9])")) for t in tokens]
    hits = []
    for name, files in ((TEMPLATE_NAME, template_files()), (TRAINING_NAME, training_files())):
        for rel, data in files.items():
            if rel.endswith(".xlsx"):
                import io
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    txt = " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))
            else:
                txt = data.decode("utf-8", "replace")
            hits += [f"{name}/{rel}: {t}" for t, rx in pats if rx.search(txt)]
    assert not hits, hits[:10]


def test_zip_root_has_both_tender_folders_beside_the_tool():
    sys.path.insert(0, str(REPO / "tools"))
    import package
    with tmpdir() as d:
        names = package.build(d / "r.zip")
        for root in (TEMPLATE_NAME, TRAINING_NAME):
            for f, _, _ in STEPS:
                assert f"{root}/{f}" in names, (root, f)
        assert f"pricing-workbench-{__version__}/pw/tender_kit.py" in names


def test_training_tender_runs_all_four_steps_as_the_step_files_do():
    """Each step's command run from the tender folder: read-back PASS, template hash = the quotes', RUN COMPLETE."""
    with tmpdir() as d:
        t = d / TRAINING_NAME
        write_folder(t, training=True)
        outs = []
        for i in range(4):
            r = subprocess.run([sys.executable, str(REPO / "tools" / "pw.py"), *step_args(i)], cwd=str(t),
                               capture_output=True, text=True, timeout=1200)
            assert r.returncode == 0, (i, r.stdout[-2000:], r.stderr[-2000:])
            outs.append(r.stdout)
        assert "CEDANT READ PASS" in outs[1]
        assert "RUN COMPLETE" in outs[3] and "V5 SUB_X PASS, SUB_Y PASS" in outs[3]
        m = json.loads((t / "runs" / "TRAINING" / "run_manifest.json").read_text(encoding="utf-8"))
        assert m["run"]["status"] == "COMPLETE" and m["gates"]["COVERAGE"]["status"] == "PASS"
        assert (t / "runs" / "TRAINING" / "out" / "Recommendation_Pack_TRAINING.xlsx").exists()


def test_excel_view_sets_widths_and_number_formats():
    df = pd.DataFrame({"option_id": ["A", "B"], "lives": [2487.0, 12.0], "uwy": [2025, 2024],
                       "prem": [13304559.32, 224676.97], "index": [1.05509, 0.98117],
                       "basis_note": ["x" * 200, "short"]})
    lay = dict(zip(df.columns, view_layout(df)))
    assert lay["option_id"][0] is None and lay["lives"][0] == "int" and lay["uwy"][0] == "year"
    assert lay["prem"][0] == "amount" and lay["index"][0] == "ratio"
    assert lay["prem"][1] >= len("13,304,559.32") + 2 and lay["basis_note"][1] == 60
    with tmpdir() as d:
        write_table("ViewCheck", df, d / "a")
        write_table("ViewCheck", df, d / "b")
        assert (d / "a" / "ViewCheck.xlsx").read_bytes() == (d / "b" / "ViewCheck.xlsx").read_bytes()
        ws = load_workbook(d / "a" / "ViewCheck.xlsx").active
        fmt = {c.value: ws.cell(2, c.column).number_format for c in ws[1]}
        assert fmt == {"option_id": "General", "lives": "#,##0", "uwy": "0", "prem": "#,##0.00", "index": "0.0000",
                       "basis_note": "General"}
        assert round(ws.column_dimensions["D"].width) >= 15


def test_laptop_check_lock_allows_one_check_at_a_time():
    lc = _laptop_check()
    with tmpdir() as d:
        lock = d / "laptop_check" / "RUNNING.lock"
        assert lc.acquire_lock(lock) is None
        msg = lc.acquire_lock(lock)
        assert msg and "already running" in msg and str(lock) in msg
        old = time.time() - (lc.LOCK_STALE_HOURS + 1) * 3600
        os.utime(lock, (old, old))
        assert lc.acquire_lock(lock) is None                    # a stale lock is taken over


def test_run_laptop_check_bat_streams_to_the_console_file():
    bat = (REPO / "tools" / "RUN_LAPTOP_CHECK.bat").read_text(encoding="ascii")
    assert '--console "%~dp0laptop_check_console.txt"' in bat and '> "%~dp0laptop_check_console.txt"' not in bat


def test_a_refused_second_check_leaves_the_running_checks_console_file_alone():
    """Found on the owner's laptop (25 Oct 2026): the refused start emptied the running check's console file."""
    lc = _laptop_check()
    lock = REPO / "runs" / "laptop_check" / "RUNNING.lock"
    with tmpdir() as d:
        con = d / "console.txt"
        con.write_text("PASS  something\n", encoding="utf-8")
        assert lc.acquire_lock(lock) is None
        try:
            assert lc.main(["--quick", "--console", str(con)]) == lc.EXIT_LOCKED
        finally:
            lock.unlink()
        assert con.read_text(encoding="utf-8") == "PASS  something\n"
