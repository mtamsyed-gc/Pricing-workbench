"""WS2: SME intake adapter (V1 on market files), hand-checked as-sold prices, tier matching, richness, gap register,
CAT_ALPHA vs CAT_GAMMA vs CAT_IPMI comparison with FX, end-to-end CLI. Optional: every real intake file in
PBM_SME_DIR (env) loads and passes V1 (skipped when not set)."""
import json
import os
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from _util import FIX, REPO, catalogue, raises, tmpdir
from openpyxl import load_workbook

from pw import MappingRequired, SpecError
from pw.benchmark import gap_register, plan_price, read_market_basis, read_standard_cells, richness, value_map
from pw.gates import v1_source_reproduction
from pw.market_intake import load_market_folder, load_sme_intake
from pw.tiers import TierMap

MKT = FIX / "MARKET_SYN"
LADDER = ["CN", "GN", "RN", "RN2", "RN3", "VN"]
STD = read_standard_cells(MKT / "standard_cells.csv")


def _intake(name):
    return load_sme_intake(MKT / "intake" / f"{name}.xlsx")


def _rate(name, net, age, member, gender):
    """Base rate straight from the intake cells (B_Rates rows 42.., 4 columns per network), no adapter."""
    wb = load_workbook(MKT / "intake" / f"{name}.xlsx", data_only=True)
    b = wb["B_Rates"]
    j = LADDER.index(net)
    for i in range(20):
        lo, hi = b.cell(16 + i, 2).value, b.cell(16 + i, 3).value
        if lo is not None and lo <= age <= hi:
            k = ["PM", "PF", "DM", "DF"].index(("P" if member == "EMPLOYEE" else "D") + gender)
            return float(b.cell(42 + i, 2 + 4 * j + k).value)
    raise KeyError


def test_intake_adapter_loads_and_passes_v1():
    cats, rep = load_market_folder(MKT / "intake")
    assert [c.catalogue_id for c in cats] == ["MKT_MKT-ALPINE-2026", "MKT_MKT-BIRCH-2026", "MKT_MKT-CEDAR-2026"]
    for c in cats:
        assert c.kind == "MARKET"
        v = v1_source_reproduction(c, MKT / "intake" / c.meta["market"]["source_file"])
        assert v["status"] == "PASS"
    a = cats[0].scopes[0]
    assert [p.plan_id for p in a.plans] == ["CN", "GN", "RN", "RN3"] and a.rating_region == "DXB_NE"
    f = next(c for c in a.components if c.component_id == "f_copay_op_diag")
    vals = {(r["key"]["copay_op_diag"], r["key"]["network"]): r["value"] for r in f.rows}
    assert vals[("0%", "ANY")] == 1.035 and vals[("10%", "GN")] == 1.0          # default option = 1 by definition
    st = {(r["plan_id"], r["area"], r["item"]): r["status"] for r in rep}
    # CR #13 (approved 24 Oct 2026): room type is a registry dimension, so its factors enter the price
    rt = next(c for c in a.components if c.component_id == "f_room_type")
    assert {(r["key"]["room_type"], r["value"]) for r in rt.rows} >= {("Semi Private", 0.955)}
    assert ("MKT-ALPINE-2026", "levers", "Territory") not in st and ("MKT-ALPINE-2026", "factors", "Room type") not in st
    assert st[("MKT-ALPINE-2026", "addons", "Maternity")] == "PRICED"
    assert st[("MKT-BIRCH-2026", "addons", "Optical")] == "NOT_PRICED"          # no priced limits at all: not bridgeable
    assert st[("MKT-CEDAR-2026", "addons", "Dental")] == "BRIDGED"              # 2,000 between the file's 1,000 and 5,000
    dn = next(c for c in cats[2].scopes[0].components if c.component_id == "m_module_dental")
    assert [r["value"] for r in dn.rows] == [250.0 + (450.0 - 250.0) * (2000 - 1000) / (5000 - 1000)]
    assert "bridged" in dn.rows[0]["source_ref"]


def test_as_sold_price_hand_checked_with_modules():
    cat, _ = _intake("MKT-ALPINE-2026")
    s = cat.scopes[0]
    pp = plan_price(cat, s, "GN", STD, "AED", {"AED": 1.0}, {})
    num = 0.0
    for r in STD.itertuples():
        mem = "EMPLOYEE" if r.member_category == "EMPLOYEE" else "DEP"
        v = _rate("MKT-ALPINE-2026", "GN", r.age, mem, r.gender)       # default design: all factors 1
        v += 310.0                                                      # dental, all members
        if r.maternity_eligible == "Y":
            v += 2400.0                                                 # maternity, eligible females
        num += r.weight * round(v, 2)
    assert abs(pp["price"] - num / STD["weight"].sum()) < 1e-9 and pp["priced_share"] == 1.0


def test_monthly_gross_plan_annualised_and_put_on_net_basis():
    cat, _ = _intake("MKT-BIRCH-2026")
    assert cat.premium_basis == "GROSS" and cat.meta["market"]["gross_to_net"] == 0.82
    s = cat.scopes[0]
    one = STD.iloc[[0]]                                                 # 35 M employee
    pp = plan_price(cat, s, "RN", one, "AED", {"AED": 1.0}, {"basis": "NET", "basis_factors": {"GROSS_TO_NET": 0.82}})
    monthly = _rate("MKT-BIRCH-2026", "RN", 35, "EMPLOYEE", "M")
    assert abs(pp["price"] - round(monthly * 12 * 0.82, 2)) < 1e-9


def test_married_eligible_module_network_override():
    cat, _ = _intake("MKT-CEDAR-2026")
    s = cat.scopes[0]
    f = pd.DataFrame([{"age": 32, "gender": "F", "member_category": "SPOUSE", "marital": "MARRIED", "maternity_eligible": "Y",
                       "nationality_group": None, "group_size_band": None, "weight": 1.0}])
    gn = plan_price(cat, s, "GN", f, "AED", {}, {})["price"]
    cn = plan_price(cat, s, "CN", f, "AED", {}, {})["price"]
    dental = 250.0 + (450.0 - 250.0) * 1000 / 4000                     # bridged 2,000 level, all members
    assert abs(gn - round(_rate("MKT-CEDAR-2026", "GN", 32, "DEP", "F") + 3300.0 + dental, 2)) < 1e-9   # GN-specific row
    assert abs(cn - round(_rate("MKT-CEDAR-2026", "CN", 32, "DEP", "F") + 2900.0 + dental, 2)) < 1e-9   # 'All' row
    single = f.assign(marital="SINGLE")
    assert abs(plan_price(cat, s, "GN", single, "AED", {}, {})["price"]
               - round(_rate("MKT-CEDAR-2026", "GN", 32, "DEP", "F") + dental, 2)) < 1e-9


def test_tier_map_rules():
    df = pd.DataFrame({"catalogue_id": ["ANY", "ANY", "X", "X"], "plan": ["P1", "P2", "P2", "P2"],
                       "tier": ["CN", "GN", "RN", "RN2"], "weight": ["1", "1", "0.4", "0.6"]})
    t = TierMap(df, LADDER)
    assert t.tiers("Y", "P2") == [("GN", 1.0)] and t.tiers("X", "P2") == [("RN", 0.4), ("RN2", 0.6)]
    raises(MappingRequired, t.tiers, "X", "P9")
    raises(SpecError, TierMap, df.assign(weight=["1", "1", "0.4", "0.5"]), LADDER)
    raises(SpecError, TierMap, df.assign(tier=["CN", "GN", "RN", "ZZ"]), LADDER)


def test_richness_on_common_basis():
    tm = TierMap.read(MKT / "tier_map.csv", LADDER)
    mb = read_market_basis(MKT / "market_basis.json")
    om = {("ANY", "copay_op_diag", "AED25"): "AED 25", ("ANY", "aml", "1250000"): "1500000"}
    a = catalogue("CAT_ALPHA")
    s = a.scopes[0]
    r6 = richness(a, s, "P6", tm.tiers("CAT_ALPHA", "P6"), mb, ["copay_op_diag", "aml"], om)
    assert r6["richness_status"] == "VALUED" and abs(r6["richness"] - 1.0 * 0.985) < 1e-12     # 10% x AED 750k
    r2 = richness(a, s, "P2", tm.tiers("CAT_ALPHA", "P2"), mb, ["copay_op_diag", "aml"], om)
    assert abs(r2["richness"] - 1.012) < 1e-12                                                  # 1.25m mapped to 1.5m
    r2b = richness(a, s, "P2", tm.tiers("CAT_ALPHA", "P2"), mb, ["copay_op_diag", "aml"], {})
    assert r2b["richness_status"] == "PARTIAL" and "1250000" in r2b["richness_note"]           # never assumed 1
    g = catalogue("CAT_GAMMA")
    rg = richness(g, g.scopes[0], "G1", tm.tiers("CAT_GAMMA", "G1"), mb, ["copay_op_diag", "aml"], om)
    assert rg["richness_status"] == "NOT_VALUED" and np.isnan(rg["richness"])
    assert mb["placeholder"] is True


def test_value_map_alpha_gamma_ipmi_with_fx():
    tm = TierMap.read(MKT / "tier_map.csv", LADDER)
    cats = [catalogue("CAT_ALPHA"), catalogue("CAT_GAMMA"), catalogue("CAT_IPMI")]
    fx = {"AED": 1.0, "USD": 3.6725}
    vm = value_map(cats, {"CAT_ALPHA"}, tm, STD, None, "AED", fx, "NET", {}, None, [], {})
    assert set(vm["catalogue_id"]) == {"CAT_ALPHA", "CAT_GAMMA", "CAT_IPMI"}
    assert (vm["priced_share_standard"] == 1.0).all() and vm["price_standard"].notna().all()
    ip = cats[2]
    usd = plan_price(ip, ip.scopes[0], "PLUS", STD, "USD", fx, {})["price"]
    row = vm[(vm["catalogue_id"] == "CAT_IPMI") & (vm["plan"] == "PLUS")].iloc[0]
    assert abs(row["price_standard"] - usd * 3.6725) < 0.01 * len(STD)            # cell-level rounding only
    cn = vm[vm["tier"] == "CN"]
    med = cn[cn["role"] == "MARKET"]["price_standard"].median()
    a1 = cn[cn["catalogue_id"] == "CAT_ALPHA"].iloc[0]
    assert abs(a1["ratio_to_median_standard"] - a1["price_standard"] / med) < 1e-12
    p6 = vm[(vm["catalogue_id"] == "CAT_ALPHA") & (vm["plan"] == "P6")].iloc[0]
    assert abs(p6["tier_rank"] - (0.5 * 5 + 0.5 * 6)) < 1e-12                       # one-to-many tier weights


def test_gap_register_assessments():
    tm = TierMap(pd.DataFrame({"catalogue_id": ["ANY"], "plan": ["ANY"], "tier": ["GN"], "weight": ["1"]}), LADDER)
    from pw.model import BenefitLine, Catalogue, Component, Plan, RatingScope

    def cat(cid, lines):
        s = RatingScope(scope_id=cid, product_type="SME", regulator="DHA", rating_region="DXB", tpa_set=["ANY"],
                        plans=[Plan(plan_id="GN", label="GN")],
                        components=[Component(component_id="base", type="BASE_TABLE", keys=["age"], rows=[])],
                        benefits=[BenefitLine(plan_id="GN", benefit_code=k, value=v) for k, v in lines.items()])
        return Catalogue(catalogue_id=cid, owner="x", kind="MARKET", version="1", effective_date=None, currency="AED",
                         premium_basis="NET", scopes=[s])
    subj = cat("S", {"aml": 500000, "room_type": "SEMI_PRIVATE", "telehealth": False, "dental_limit": 3000})
    mk = [cat("M1", {"aml": 1000000, "room_type": "PRIVATE", "telehealth": True, "dental_limit": 2000, "mat_scans": 3}),
          cat("M2", {"aml": 1000000, "room_type": "PRIVATE", "telehealth": True, "dental_limit": 1000}),
          cat("M3", {"aml": 250000, "room_type": "SHARED", "telehealth": False, "dental_limit": 0})]
    g = gap_register([subj] + mk, {"S"}, tm).set_index("benefit_code")
    assert g.loc["aml", "assessment"] == "WEAKER" and float(g.loc["aml", "market_median"]) == 1000000
    assert abs(g.loc["aml", "share_peers_richer"] - 2 / 3) < 1e-12
    assert g.loc["room_type", "assessment"] == "WEAKER" and g.loc["room_type", "market_median"] == "PRIVATE"
    assert g.loc["telehealth", "assessment"] == "GAP"                    # not covered while 2/3 of the market covers it
    assert g.loc["dental_limit", "assessment"] == "RICHER"
    assert g.loc["mat_scans", "assessment"] == "NOT_IN_SCHEDULE"         # data completeness, not a benefit gap


def test_cli_benchmark_end_to_end():
    pw = [sys.executable, str(REPO / "tools" / "pw.py")]
    with tmpdir() as d:
        r = subprocess.run(pw + ["benchmark", str(FIX / "RUN_SYN" / "run_config.json"), "--runs-root", str(d)],
                           capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0, r.stderr
        from pw.outputs import read_table
        out = d / "RUN_SYN" / "out"
        vm = read_table("ValueMap", out)
        assert {"MKT_MKT-ALPINE-2026", "CAT_GAMMA", "CAT_IPMI", "CAT_ALPHA"} <= set(vm["catalogue_id"])
        assert vm["price_standard"].notna().all()
        g = vm["catalogue_id"] == "CAT_GAMMA"          # SME group-size bands start at 3; the exposure mix is individual (1)
        assert vm.loc[~g, "price_exposure"].notna().all() and (vm.loc[g, "priced_share_exposure"] == 0).all()
        assert vm.loc[g, "status"].str.contains("NO_PRICED_CELLS_EXPOSURE").all()
        assert vm["richness_basis"].str.startswith("PLACEHOLDER").all()
        assert len(read_table("GapRegister", out)) and len(read_table("TaxonomyReconciliation", out))
        m = json.loads((d / "RUN_SYN" / "run_manifest.json").read_text())
        assert m["benchmark"]["market_plans"] == 5 and m["benchmark"]["market_basis_placeholder"] is True


def test_real_sme_intake_files_if_available():
    """Set PBM_SME_DIR to <PBM clone>/sme_benchmarking to run over the tool's shipped intake files (never in this repo)."""
    root = os.environ.get("PBM_SME_DIR")
    if not root:
        return
    folder = Path(root) / "intake_shipped"
    cats, rep = load_market_folder(folder)
    assert len(cats) >= 1
    for c in cats:
        assert v1_source_reproduction(c, folder / c.meta["market"]["source_file"], raise_on_fail=False)["status"] == "PASS"


def test_import_sme_basis_converts_and_flags_placeholder():
    import types
    sys.path.insert(0, str(REPO / "tools"))
    import import_sme_basis as imp
    from pw.market_intake import load_layout
    tax = types.SimpleNamespace(TOOL_VERSION="fictional", NETS=LADDER,
                                REF_LEVER={"Annual Medical Limit": [(1000000, [1, 1, 1, 1, 1, 1]), (500000, [0.9, 0.91, 0.92, 0.93, 0.94, 0.95])],
                                           "Room type": [("Private", [1, 1, 1, 1, 1, 1])],
                                           "Other adjustment 1": [("x", [1, 1, 1, 1, 1, 1])]})
    d = imp.convert(tax, load_layout())
    assert d["placeholder"] is True and d["skipped_levers"] == ["Other adjustment 1"]   # room type mapped since CR #13
    e = {(x["family"], x["option"], x["network"]): x["relativity"] for x in d["entries"]}
    assert e[("aml", "500000", "RN")] == 0.92 and e[("aml", "1000000", "VN")] == 1.0
    assert e[("room_type", "Private", "GN")] == 1 and {x["family"] for x in d["entries"]} == {"aml", "room_type"}
