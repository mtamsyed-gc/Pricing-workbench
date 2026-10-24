"""WS3: scorecard hand checks, weight validation, ranking flip on a constructed case, lead-only panels, grid,
incumbent-name guard and the end-to-end CLI."""
import json
import subprocess
import sys

import numpy as np
import pandas as pd
from _util import FIX, REPO, raises, tmpdir
from openpyxl import load_workbook

from pw import ProposalRequired, SpecError
from pw.scorecard import (flip_points, grid, normalise, one_at_a_time, panel_inputs, score, score_inputs,
                          validate_weights, write_pack)


def _inputs(rows):
    """rows: {option_id: {criterion: value}} -> ScoreInputs long table."""
    out = []
    for oid, d in rows.items():
        for c, v in d.items():
            out.append({"option_id": oid, "criterion": c, "metric": c.lower(), "value": v, "source": "test",
                        "submitted_text": None})
    return pd.DataFrame(out)


# A cheaper, B better commission, C in between on both
CASE = _inputs({"A": {"PRICE": 1.00, "COMMISSION": 12.0},
                "B": {"PRICE": 1.06, "COMMISSION": 18.0},
                "C": {"PRICE": 1.02, "COMMISSION": 14.0}})


def test_weights_must_sum_to_one_and_be_known():
    validate_weights({"PRICE": 0.6, "COMMISSION": 0.4})
    raises(SpecError, validate_weights, {"PRICE": 0.6, "COMMISSION": 0.5})      # 1.1: rejected, not rescaled
    raises(SpecError, validate_weights, {"PRICE": 0.6, "COMMISSION": 0.39})
    raises(SpecError, validate_weights, {"PRICE": 1.2, "COMMISSION": -0.2})
    raises(SpecError, validate_weights, {"PRICE": 0.5, "QUALITY": 0.5})
    raises(ProposalRequired, validate_weights, {})


def test_min_max_normalisation():
    s = normalise(pd.Series([1.00, 1.06, 1.02], index=list("ABC")), "LOWER")
    assert np.allclose(s.to_numpy(), [100.0, 0.0, 100 * 0.04 / 0.06])
    s = normalise(pd.Series([12.0, 18.0, np.nan], index=list("ABC")), "HIGHER")
    assert s["A"] == 0.0 and s["B"] == 100.0 and np.isnan(s["C"])
    assert (normalise(pd.Series([5.0, 5.0], index=list("AB")), "HIGHER") == 100.0).all()


def test_weighted_scores_hand_checked():
    card, rank = score(CASE, {"PRICE": 0.7, "COMMISSION": 0.3})
    # normalised: PRICE A 100, B 0, C 66.667; COMMISSION A 0, B 100, C 33.333
    exp = {"A": 0.7 * 100 + 0.3 * 0, "B": 0.7 * 0 + 0.3 * 100, "C": 0.7 * (200 / 3) + 0.3 * (100 / 3)}
    got = rank.set_index("option_id")["score"].to_dict()
    for k, v in exp.items():
        assert abs(got[k] - v) < 1e-9, (k, got[k], v)
    assert list(rank["option_id"]) == ["A", "C", "B"] and list(rank["rank"]) == [1, 2, 3]
    c = card[(card["option_id"] == "C") & (card["criterion"] == "PRICE")].iloc[0]
    assert abs(c["contribution"] - 0.7 * 200 / 3) < 1e-9


def test_flip_weight_on_constructed_case():
    """PRICE weight t, COMMISSION 1 - t. A = 100t, B = 100(1-t), C = (200t + 100(1-t))/3 = (100 + 100t)/3.
    Leader: B for t < 1/2 (B > C until t = 0.5: 100 - 100t = (100 + 100t)/3 -> t = 0.5, where A = B = C = 50),
    A for t > 1/2. So the leader flips from B to A at exactly t = 0.5 on PRICE."""
    w = {"PRICE": 0.7, "COMMISSION": 0.3}
    oat = one_at_a_time(CASE, w)
    p = oat[oat["criterion"] == "PRICE"].sort_values("t_from")
    assert list(p["leader"]) == ["B", "A"]
    assert abs(p.iloc[0]["t_to"] - 0.5) < 1e-12
    fl = flip_points(oat)
    f = fl[fl["criterion"] == "PRICE"].iloc[0]
    assert abs(f["flip_weight"] - 0.5) < 1e-12 and f["leader_below"] == "B" and f["leader_above"] == "A"
    assert abs(f["distance"] - 0.2) < 1e-12
    # the mirror criterion flips at 1 - 0.5
    c = fl[fl["criterion"] == "COMMISSION"].iloc[0]
    assert abs(c["flip_weight"] - 0.5) < 1e-12 and c["leader_below"] == "A" and c["leader_above"] == "B"
    # confirm by brute force either side of the flip
    for t, lead in ((0.49, "B"), (0.51, "A")):
        _, r = score(CASE, {"PRICE": t, "COMMISSION": 1 - t})
        assert r.iloc[0]["option_id"] == lead


def test_grid_win_shares():
    g = grid(CASE, {"PRICE": 0.7, "COMMISSION": 0.3}, step=0.1).set_index("option_id")
    # 11 points t = 0..1: B leads t = 0..0.4 (5 pts), 3-way tie at 0.5 split equally, A leads t = 0.6..1 (5 pts)
    assert g["grid_points"].iloc[0] == 11
    assert abs(g.loc["A", "win_share"] - (5 + 1 / 3) / 11) < 1e-12
    assert abs(g.loc["B", "win_share"] - (5 + 1 / 3) / 11) < 1e-12
    assert abs(g.loc["C", "win_share"] - (1 / 3) / 11) < 1e-12
    raises(SpecError, grid, CASE, {"PRICE": 0.7, "COMMISSION": 0.3}, step=0.3)


def test_missing_weighted_value_is_unranked_not_zero():
    case = _inputs({"A": {"PRICE": 1.0, "COMMISSION": 12.0}, "B": {"PRICE": 1.1, "COMMISSION": np.nan}})
    _, r = score(case, {"PRICE": 0.5, "COMMISSION": 0.5})
    b = r.set_index("option_id").loc["B"]
    assert b["status"] == "UNRANKED" and np.isnan(b["score"]) and "COMMISSION" in b["reason"]
    _, r = score(case, {"PRICE": 1.0, "COMMISSION": 0.0})          # weight 0: missing value does not matter
    assert (r["status"] == "RANKED").all()


def _ri():
    rows = []
    for oid, idx in (("X", [1.03, 1.04, 1.02]), ("Y", [0.99, 0.98, 1.00])):
        for sc, v in zip(("STANDARD", "WRITTEN_LY", "UWY_MINUS_1"), idx):
            rows.append({"option_id": oid, "scenario": sc, "cut": "TOTAL", "index": v,
                         "index_benefit_adj": v / 1.01, "unpriced_share": 0.01 if oid == "X" else 0.02})
    return pd.DataFrame(rows)


def test_score_inputs_and_lead_only_panel():
    ts = pd.DataFrame({"option_id": ["X", "Y"], "commission_pct": [17.5, 15.0], "profit_share_pct": [20, 25],
                       "corridor_pts": [5, 0], "cap_pct": [120, 110], "share_pct": [50, 30]})
    reg = {"X": {"rate_cells": 100, "v5_counts": {"OUTLIER": 2, "HASH_MISMATCH": 1}}, "Y": {"rate_cells": 100, "v5_counts": {}}}
    si = score_inputs(_ri(), reg, ts)
    v = si.set_index(["option_id", "criterion"])["value"]
    assert v[("X", "PRICE")] == 1.03 and abs(v[("X", "ROBUSTNESS_SPREAD")] - 0.02) < 1e-12
    assert v[("X", "V5_FLAG_SHARE")] == 0.02                      # structural flags are not cell flags
    assert v[("Y", "COMMISSION")] == 15.0 and v[("X", "CAP")] == 120
    pi = panel_inputs(si, _ri(), [{"panel_id": "P", "lead": "Y", "members": {"X": 0.25, "Y": 0.75}}])
    p = pi.set_index("criterion")["value"]
    assert abs(p["PRICE"] - (0.25 * 1.03 + 0.75 * 0.99)) < 1e-12
    s = [0.25 * a + 0.75 * b for a, b in zip([1.03, 1.04, 1.02], [0.99, 0.98, 1.00])]
    assert abs(p["ROBUSTNESS_SPREAD"] - (max(s) - min(s))) < 1e-12
    assert p["COMMISSION"] == 15.0 and p["SHARE"] == 30           # lead-only terms
    raises(SpecError, panel_inputs, si, _ri(), [{"panel_id": "P", "lead": "Y", "members": {"X": 0.3, "Y": 0.6}}])
    raises(SpecError, panel_inputs, si, _ri(), [{"panel_id": "P", "lead": "Z", "members": {"X": 0.5, "Y": 0.5}}])
    raises(SpecError, panel_inputs, si, _ri(), [{"panel_id": "X", "lead": "Y", "members": {"X": 0.5, "Y": 0.5}}])


def test_pack_names_and_incumbent_guard():
    card, rank = score(CASE, {"PRICE": 0.7, "COMMISSION": 0.3})
    oat = one_at_a_time(CASE, {"PRICE": 0.7, "COMMISSION": 0.3})
    fl, g = flip_points(oat), grid(CASE, {"PRICE": 0.7, "COMMISSION": 0.3})
    ot = pd.DataFrame({"option_id": ["A", "B", "C"], "reinsurer": ["Incumbent Re (fictional)", "Beta Re", "Gamma Re"]})
    with tmpdir() as d:
        cfg = {"show_reinsurer_names": False, "market_facing": True, "incumbent_names": ["Incumbent Re"]}
        p = write_pack(d / "pack.xlsx", "T", CASE, card, rank, oat, fl, g, {"PRICE": 0.7, "COMMISSION": 0.3}, cfg, ot)
        wb = load_workbook(p)
        assert wb.sheetnames == ["Summary", "Scorecard", "Inputs", "Sensitivity"]
        text = " ".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v)
        assert "Incumbent Re" not in text and "Beta Re" not in text
        cfg = {**cfg, "show_reinsurer_names": True}
        raises(SpecError, write_pack, d / "pack2.xlsx", "T", CASE, card, rank, oat, fl, g,
               {"PRICE": 0.7, "COMMISSION": 0.3}, cfg, ot)
        assert not (d / "pack2.xlsx").exists()                     # a refused market-facing pack is not left behind
        cfg = {**cfg, "market_facing": False}
        p = write_pack(d / "pack3.xlsx", "T", CASE, card, rank, oat, fl, g, {"PRICE": 0.7, "COMMISSION": 0.3}, cfg, ot)
        text = " ".join(str(v) for ws in load_workbook(p).worksheets for row in ws.iter_rows(values_only=True) for v in row if v)
        assert "B (Beta Re)" in text


def test_cli_score_end_to_end():
    pw = [sys.executable, str(REPO / "tools" / "pw.py")]
    cfg = str(FIX / "RUN_SYN" / "run_config.json")
    with tmpdir() as d:
        r = subprocess.run(pw + ["score", cfg, "--runs-root", str(d)], capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 2 and "index" in r.stderr                       # needs WS1 outputs first
        for o in ("SUB_X", "SUB_Y"):
            subprocess.run(pw + ["t", "ingest", str(FIX / o / "submission.xlsx"), "--spec", str(FIX / "T_ALPHA" / "template_spec.json"),
                                 "--issued", str(FIX / "T_ALPHA" / "template.xlsx"), "--run", cfg, "--id", o, "--owner", o,
                                 "--runs-root", str(d)], check=True, capture_output=True, cwd=str(REPO / "docs"))
        subprocess.run(pw + ["index", cfg, "--runs-root", str(d)], check=True, capture_output=True, cwd=str(REPO / "docs"))
        r = subprocess.run(pw + ["score", cfg, "--runs-root", str(d)], capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0, r.stderr
        out = d / "RUN_SYN" / "out"
        from pw.outputs import read_table
        rk = read_table("Ranking", out)
        assert set(rk["option_id"]) == {"SUB_X", "SUB_Y", "PANEL_XY"} and (rk["status"] == "RANKED").all()
        assert rk.iloc[0]["option_id"] == "SUB_Y"                              # cheaper option wins at example weights
        si = read_table("ScoreInputs", out)
        assert list(si.columns[:5]) == ["option_id", "criterion", "metric", "value", "source"]
        assert (out / "Recommendation_Pack_RUN_SYN.xlsx").exists()
        m = json.loads((d / "RUN_SYN" / "run_manifest.json").read_text())
        assert abs(sum(m["scorecard"]["weights"].values()) - 1) < 1e-12 and "Ranking" in m["outputs"]
        h = m["outputs"]["Ranking"]["sha256_parquet"]
        subprocess.run(pw + ["score", cfg, "--runs-root", str(d)], check=True, capture_output=True, cwd=str(REPO / "docs"))
        assert json.loads((d / "RUN_SYN" / "run_manifest.json").read_text())["outputs"]["Ranking"]["sha256_parquet"] == h
