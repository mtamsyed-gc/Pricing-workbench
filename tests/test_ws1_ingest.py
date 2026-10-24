"""WS1: T ingest + V5 (SPEC §6, §8.3) and the TPA factor basis (HANDOVER item 4, provisional option A)."""
import json

from _util import FIX, catalogue, tmpdir
from openpyxl import load_workbook

from pw.evaluator import evaluate
from pw.model import Catalogue
from pw.template.ingest import ingest, issued_hash_of, registered_options
from pw.template.parser import TPA_FACTOR_BASIS, make_reproducible, parse_submission, to_catalogue

SPEC = json.loads((FIX / "T_ALPHA" / "template_spec.json").read_text(encoding="utf-8"))
ISSUED = FIX / "T_ALPHA" / "template.xlsx"


def _keys(wb):
    out = {}
    for row in wb["_KEYS"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            out[row[0]] = (row[5], row[6], row[2])
    return out


def _tampered(d):
    """SUB_X with: a blank, a zero, a text value, an outlier (x2.5) and a tampered _META hash."""
    wb = load_workbook(FIX / "SUB_X" / "submission.xlsx")
    rate = sorted(k for k, v in _keys(wb).items() if v[2] == "RATE")
    picks = {"BLANK": rate[0], "ZERO": rate[1], "TEXT": rate[2], "OUTLIER": rate[3]}
    ks = _keys(wb)
    for what, ck in picks.items():
        sh, addr, _ = ks[ck]
        c = wb[sh][addr]
        c.value = {"BLANK": None, "ZERO": 0, "TEXT": "TBC", "OUTLIER": c.value * 2.5}[what]
    for row in wb["_META"].iter_rows(min_row=2):
        if row[0].value == "structure_hash":
            row[1].value = "f" * 64
    p = d / "tampered.xlsx"
    wb.save(p)
    make_reproducible(p)
    return p, picks


def test_v5_flags_hash_blank_zero_text_outlier_and_ingest_excludes_them():
    with tmpdir() as d:
        p, picks = _tampered(d)
        r = ingest(p, SPEC, "T1", "SUB_T", "synthetic", issued_template=ISSUED, expiring=catalogue("CAT_ALPHA"), runs_root=d)
        fl = r["v5"]["flags"]
        by = {f: set(fl[fl["flag"] == f]["cell_key"]) for f in fl["flag"].unique()}
        assert "_META" in by["HASH_MISMATCH"]
        assert picks["BLANK"] in by["BLANK"] and picks["ZERO"] in by["NOT_POSITIVE"]
        assert picks["TEXT"] in by["NON_NUMERIC"] and picks["OUTLIER"] in by["OUTLIER"]
        o = fl[fl["cell_key"] == picks["OUTLIER"]].iloc[0]
        assert 2.2 < o["ratio"] < 2.8 and abs(o["value"] / o["expiring"] - o["ratio"]) < 1e-12
        assert r["v5"]["status"] == "FLAGGED" and not r["v5"]["hash_ok"]
        assert r["v5"]["invalid"] == {picks["BLANK"], picks["ZERO"], picks["TEXT"]}
        # the option is still registered (V5 flags, never stops); invalid cells are not in the catalogue
        e = json.loads((d / "T1" / "options.json").read_text(encoding="utf-8"))["SUB_T"]
        assert e["excluded_cells"] == 3 and e["v5_status"] == "FLAGGED"
        cat = registered_options("T1", d)["SUB_T"]
        refs = {r_["source_ref"] for s in cat.scopes for r_ in s.components[0].rows}
        ks = _keys(load_workbook(p))
        for ck in r["v5"]["invalid"]:
            assert f"{ks[ck][0]}!{ks[ck][1]}" not in refs
        assert (d / "T1" / "options" / "SUB_T" / "v5_flags.csv").exists()
        assert json.loads((d / "T1" / "run_manifest.json").read_text())["options"]["SUB_T"]["v5_status"] == "FLAGGED"


def test_v5_tampered_keys_and_missing_issued_hash():
    from pw.gates import v5_submission_integrity
    p = parse_submission(FIX / "SUB_Y" / "submission.xlsx", template_spec=SPEC)
    r = v5_submission_integrity(p, SPEC, issued_hash_of(ISSUED))
    assert r["status"] == "PASS"                                 # no expiring: no outlier test, structure fine
    r = v5_submission_integrity(p, SPEC, "0" * 64)
    assert {"HASH_MISMATCH", "KEYS_TAMPERED"} <= set(r["flags"]["flag"])
    r = v5_submission_integrity(p, SPEC, None)
    assert "no issued hash" in r["flags"].iloc[0]["detail"]


def test_tpa_factor_basis_pre_is_pinned_and_post_neutralises():
    """Option A (PRE): price = submitted block value x issued TPA factor (0.93 for TPA_GREEN) x other factors."""
    assert TPA_FACTOR_BASIS == "PRE"
    alpha = catalogue("CAT_ALPHA")
    p = parse_submission(FIX / "SUB_X" / "submission.xlsx", template_spec=SPEC)
    pre = to_catalogue(p, SPEC, "X", "x", expiring=alpha)
    post = to_catalogue(p, SPEC, "X", "x", expiring=alpha, tpa_basis="POST")
    cell = {"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_GREEN", "plan": "P4", "age": 41, "gender": "M",
            "copay_op_diag": "0%", "nationality_group": "OTHER"}
    scope = next(s for s in pre.scopes if s.tpa_set == ["TPA_GREEN"])
    v = next(r["value"] for r in scope.components[0].rows if r["key"]["plan"] == "P4" and r["key"]["gender"] == "M"
             and r["key"]["age"][0] <= 41 <= r["key"]["age"][1])
    a, _ = evaluate(pre, [cell])
    b, _ = evaluate(post, [cell])
    assert a.loc[0, "price"] == round(v * 0.93, 2)
    assert b.loc[0, "price"] == round(v, 2)
    assert pre.meta["tpa_factor_basis"] == "PRE" and post.meta["tpa_factor_basis"] == "POST"


def test_registered_options_round_trip():
    with tmpdir() as d:
        for o in ("SUB_X", "SUB_Y"):
            owner = Catalogue.read(FIX / o / "catalogue.json").owner
            ingest(FIX / o / "submission.xlsx", SPEC, "T2", o, owner, issued_template=ISSUED, expiring=catalogue("CAT_ALPHA"),
                   runs_root=d)
        opts = registered_options("T2", d)
        assert list(opts) == ["SUB_X", "SUB_Y"]
        assert opts["SUB_X"].content_hash() == Catalogue.read(FIX / "SUB_X" / "catalogue.json").content_hash()
