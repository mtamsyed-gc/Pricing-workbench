"""WS9 Spec Builder (SPEC §8.6, CR #26): spec -> builder -> spec round trips on every repo spec, a builder filled from
blank reproduces the fixture spec, every read-back gate stops with its cell, determinism, fields kept from the base."""
import json
import shutil

from _util import FIX, REPO, raises, requires, tmpdir
from openpyxl import load_workbook

from pw import SpecError
from pw.gates import v4_round_trip
from pw.template.builder import (APPLIES_SHEET, HEADER_ROW, T, TOP_KEYS, generate_builder, read_builder, write_spec)
from pw.template.cedant import generate_cedant_workbook
from pw.template.generator import generate_template
from pw.template.spec import CEDANT_SHEET_FIELDS, cedant_model, cedant_spec_hash, load_template_spec, normalise, spec_hash

SPECS = {"T_ALPHA": FIX / "T_ALPHA" / "template_spec.json", "C_ALPHA": FIX / "C_ALPHA" / "template_spec.json"}
EX = REPO / "examples" / "ufic_2027" / "template_spec.json"


def _strip(d):
    return {k: v for k, v in d.items() if k != "_base_dir"}


def _semantic(spec):
    """Meaning of a spec: the normalised tender layout and the cedant model (raw forms and text-only fields left out)."""
    n = normalise(spec)
    n = {k: v for k, v in n.items() if k in TOP_KEYS and k not in ("age_bands_file", "cedant_workbook")}
    n["sheets"] = [{k: v for k, v in s.items() if k not in CEDANT_SHEET_FIELDS} for s in n["sheets"]]
    extra = [{k: s.get(k) for k in ("formula", "premium_basis", "scope_id")} for s in spec["sheets"]]
    return json.dumps({"n": n, "m": cedant_model(spec)[0]["sheets"], "x": extra}, sort_keys=True, default=str)


def _round_trip(path, d):
    spec = load_template_spec(path)
    x = d / "builder.xlsx"
    generate_builder(spec, x)
    got, rep = read_builder(x, spec)
    assert rep.status == "PASS", rep.problems[:10]
    # same JSON (content and key order), same hashes
    assert json.dumps(_strip(got)) == json.dumps(_strip(spec))
    assert spec_hash(normalise(got)) == spec_hash(normalise(spec))
    assert cedant_spec_hash(got) == cedant_spec_hash(spec)
    # the issued template and the cedant workbook are byte-identical
    generate_template(spec, d / "t1.xlsx")
    generate_template(got, d / "t2.xlsx")
    assert (d / "t1.xlsx").read_bytes() == (d / "t2.xlsx").read_bytes()
    if cedant_model(spec)[0]["sheets"] and any(s.get(k) for s in spec["sheets"] for k in ("factor_tables", "riders", "benefits")):
        generate_cedant_workbook(spec, d / "c1.xlsx")
        generate_cedant_workbook(got, d / "c2.xlsx")
        assert (d / "c1.xlsx").read_bytes() == (d / "c2.xlsx").read_bytes()
    # without the base: same meaning
    got2, rep2 = read_builder(x, None)
    assert rep2.status == "PASS", rep2.problems[:10]
    assert _semantic(got2) == _semantic(spec)
    return spec, got


def test_round_trip_fixture_specs():
    for name, path in SPECS.items():
        with tmpdir() as d:
            _round_trip(path, d)


@requires("examples/ufic_2027")
def test_round_trip_cedant_example_keeps_age_bands_file_and_owner_records():
    with tmpdir() as d:
        spec, got = _round_trip(EX, d)
        assert got["age_bands_file"] == spec["age_bands_file"] and "age_band_sets" not in got
        for k in ("_note", "source_flags", "pending"):
            assert got[k] == spec[k]
        assert got["cedant_workbook"]["decided"] == spec["cedant_workbook"]["decided"]


def _values(xlsx):
    """{sheet: [ {header: value} ]} of a builder's long tables, and the field / value cells."""
    wb = load_workbook(xlsx, data_only=True)
    out = {}
    for n, cols in T.items():
        if cols is None:
            continue
        ws = wb[n]
        hdr = [str(ws.cell(HEADER_ROW + 1, c).value or "").rstrip(" *") for c in range(1, len(cols) + 1)]
        rows = []
        for r in range(HEADER_ROW + 2, ws.max_row + 1):
            v = {h: ws.cell(r, c + 1).value for c, h in enumerate(hdr)}
            if any(x is not None for x in v.values()):
                rows.append(v)
        out[n] = rows
    fv = {}
    for n in ("Tender", "Cedant workbook"):
        ws = wb[n]
        for r in range(HEADER_ROW + 2, ws.max_row + 1):
            if ws.cell(r, 2).value is not None:
                fv[(n, ws.cell(r, 1).value, ws.cell(r, 2).value)] = ws.cell(r, 4).value
    lists = {c: [ws_c.value for ws_c in col if ws_c.value is not None]
             for c, col in enumerate(wb["Lists"].iter_cols(min_row=HEADER_ROW + 2))}
    return out, fv, lists


def test_blank_builder_filled_by_hand_reproduces_the_fixture_spec():
    """A user fills a blank builder (here: openpyxl types the fixture's rows into it); read without a base spec it
    gives the fixture spec's meaning, and the template passes V4."""
    spec = load_template_spec(SPECS["C_ALPHA"])
    with tmpdir() as d:
        generate_builder(spec, d / "pre.xlsx")
        rows, fv, lists = _values(d / "pre.xlsx")
        generate_builder(None, d / "blank.xlsx")
        assert _values(d / "blank.xlsx")[0] == {n: [] for n in rows}
        wb = load_workbook(d / "blank.xlsx")
        for n, rws in rows.items():
            ws = wb[n]
            hdr = {str(ws.cell(HEADER_ROW + 1, c).value or "").rstrip(" *"): c for c in range(1, ws.max_column + 1)}
            for i, row in enumerate(rws):
                for h, v in row.items():
                    if v is not None:
                        ws.cell(HEADER_ROW + 2 + i, hdr[h]).value = v
        for n in ("Tender", "Cedant workbook"):
            ws = wb[n]
            for r in range(HEADER_ROW + 2, ws.max_row + 1):
                k = (n, ws.cell(r, 1).value, ws.cell(r, 2).value)
                if k in fv:
                    ws.cell(r, 4).value = fv[k]
        wb.save(d / "filled.xlsx")
        got, rep = read_builder(d / "filled.xlsx", None)
        assert rep.status == "PASS", rep.problems[:10]
        assert _semantic(got) == _semantic(spec)
        assert spec_hash(normalise(got)) == spec_hash(normalise(spec))
        r = v4_round_trip(got, d / "v4")
        assert r["status"] == "PASS"


def test_generate_and_read_are_deterministic():
    spec = load_template_spec(SPECS["C_ALPHA"])
    with tmpdir() as d:
        generate_builder(spec, d / "a.xlsx")
        generate_builder(spec, d / "b.xlsx")
        assert (d / "a.xlsx").read_bytes() == (d / "b.xlsx").read_bytes()
        generate_builder(None, d / "c.xlsx")
        generate_builder(None, d / "e.xlsx")
        assert (d / "c.xlsx").read_bytes() == (d / "e.xlsx").read_bytes()
        s1, _ = read_builder(d / "a.xlsx", spec)
        s2, _ = read_builder(d / "a.xlsx", spec)
        write_spec(s1, d / "1.json")
        write_spec(s2, d / "2.json")
        assert (d / "1.json").read_bytes() == (d / "2.json").read_bytes()
        assert (d / "1.json").read_bytes().endswith(b"\n") and b"\r\n" not in (d / "1.json").read_bytes()


# ---------------------------------------------------------------------------------------------- gates
def _edit(src, dst, edits):
    """edits: [(sheet, header, row index (0 = first data row) | ('field', section, code), value)]"""
    wb = load_workbook(src)
    for sheet, header, idx, value in edits:
        ws = wb[sheet]
        if isinstance(idx, tuple):
            for r in range(1, ws.max_row + 1):
                if ws.cell(r, 1).value == idx[1] and ws.cell(r, 2).value == idx[2]:
                    ws.cell(r, 4).value = value
            continue
        hdr = {str(ws.cell(HEADER_ROW + 1, c).value or "").rstrip(" *"): c for c in range(1, ws.max_column + 1)}
        ws.cell(HEADER_ROW + 2 + idx, hdr[header]).value = value
    wb.save(dst)
    return {s: {h: c for h, c in [(str(wb[s].cell(HEADER_ROW + 1, c).value or "").rstrip(" *"), c)
                                  for c in range(1, wb[s].max_column + 1)]} for s, *_ in edits}


def _cell(sheet_cols, sheet, header, idx):
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(sheet_cols[sheet][header])}{HEADER_ROW + 2 + idx}"


def _gate(edits, check, base=True, cell=None):
    spec = load_template_spec(SPECS["C_ALPHA"])
    with tmpdir() as d:
        generate_builder(spec, d / "b.xlsx")
        cols = _edit(d / "b.xlsx", d / "e.xlsx", edits)
        got, rep = read_builder(d / "e.xlsx", spec if base else None)
        assert got is None, f"{check}: the read must stop"
        hits = [p for p in rep.problems if p["check"] == check]
        assert hits, (check, rep.problems[:5])
        if cell:
            want = _cell(cols, *cell)
            assert any(p["cell"] == want for p in hits), (want, hits)
        return rep


def test_gate_blank_mandatory_and_not_in_list():
    _gate([("Sheets", "Product type", 0, None)], "blank_mandatory", cell=("Sheets", "Product type", 0))
    _gate([("Sheets", "Product type", 0, "NOT_A_PRODUCT")], "not_in_list", cell=("Sheets", "Product type", 0))
    _gate([("Columns", "Gender", 0, "X")], "not_in_list", cell=("Columns", "Gender", 0))
    _gate([("Benefits", "Benefit code", 0, "no_such_benefit")], "not_in_list", cell=("Benefits", "Benefit code", 0))
    _gate([("Tender", None, ("field", "tender", "premium_basis"), "NETT")], "not_in_list")
    _gate([("Tender", None, ("field", "tender", "cedant"), None)], "spec_invalid")       # validator: missing cedant


def test_gate_numbers_and_ages():
    _gate([("Plans", "Plan number", 0, "one")], "not_a_number", cell=("Plans", "Plan number", 0))
    _gate([("Age bands", "Age from", 0, 1.5)], "not_whole", cell=("Age bands", "Age from", 0))
    _gate([("Age bands", "Age from", 1, 99)], "age_range", cell=("Age bands", "Age from", 1))
    _gate([("Age bands", "Age to", 0, 40)], "overlap")
    _gate([("Age bands", "Age to", 0, None)], "age_range", cell=("Age bands", "Age to", 0))


def test_gate_duplicates():
    _gate([("Blocks", "Block id", 1, "ALPHA_IND_A|TPA_RED")], "duplicate", cell=("Blocks", "Block id", 1))
    _gate([("Age bands", "Label", 1, "0-4")], "duplicate", cell=("Age bands", "Label", 1))
    _gate([("Factor tables", "Component id", 1, "f_opdiag")], "duplicate", cell=("Factor tables", "Component id", 1))
    _gate([("Options", "Option", 1, "0%")], "duplicate", cell=("Options", "Option", 1))


def test_gate_undefined_references():
    _gate([("Blocks", "Age band set", 0, "NO_SUCH_SET")], "undefined", cell=("Blocks", "Age band set", 0))
    _gate([("Blocks", "Column set", 0, "NO_SUCH_COLS")], "undefined", cell=("Blocks", "Column set", 0))
    _gate([("Plans", "Sheet", 0, "No such sheet")], "undefined", cell=("Plans", "Sheet", 0))
    _gate([("Plans", "Block id", 0, "NO_SUCH_BLOCK")], "undefined", cell=("Plans", "Block id", 0))
    _gate([("Options", "Component id", 0, "f_nope")], "undefined", cell=("Options", "Component id", 0))
    _gate([("Refer cells", "Plan", 0, "P99")], "undefined", cell=("Refer cells", "Plan", 0))


def test_gate_row_sources_and_tpa_plans():
    # custom rows on a block that already has an age band set
    rep = _gate([("Custom rows", "Sheet", 0, "INDIVIDUAL RGN_A"), ("Custom rows", "Applies to", 0, "ALPHA_IND_A|TPA_RED"),
                 ("Custom rows", "Row label", 0, "All ages")], "row_source")
    assert "block 'ALPHA_IND_A|TPA_RED'" in rep.problems[0]["message"]
    # the same TPA x plan in two blocks: the spec validator stops it (SPEC §8.0)
    _gate([("Blocks", "TPA", 1, "TPA_RED")], "spec_invalid")


def test_gate_factor_reference_rows():
    spec = load_template_spec(SPECS["C_ALPHA"])
    with tmpdir() as d:
        generate_builder(spec, d / "b.xlsx")
        rows, _, _ = _values(d / "b.xlsx")
        first = [i for i, r in enumerate(rows["Options"]) if r["Component id"] == "f_opdiag"]
        ref = [i for i in first if rows["Options"][i]["Reference"] == "Y"][0]
        other = [i for i in first if i != ref][0]
    _gate([("Options", "Reference", other, "Y")], "reference_count", cell=("Options", "Reference", max(ref, other)))
    _gate([("Options", "Reference", ref, None)], "reference_count")


def test_gate_spec_changed_and_not_a_builder():
    spec = load_template_spec(SPECS["C_ALPHA"])
    with tmpdir() as d:
        generate_builder(spec, d / "b.xlsx")
        changed = json.loads(json.dumps(_strip(spec)))
        changed["sheets"][0]["title"] = "changed elsewhere"
        changed["_base_dir"] = spec["_base_dir"]
        got, rep = read_builder(d / "b.xlsx", changed)
        assert got is None and rep.problems[0]["check"] == "spec_changed"
        got, rep = read_builder(FIX / "C_ALPHA" / "cedant_workbook.xlsx", spec)
        assert got is None and rep.problems[0]["check"] == "not_a_builder"
        wb = load_workbook(d / "b.xlsx")
        del wb["Plans"]
        wb.save(d / "c.xlsx")
        got, rep = read_builder(d / "c.xlsx", spec)
        assert got is None and rep.problems[0]["check"] == "sheet_missing"


@requires("examples/ufic_2027")
def test_gate_bands_of_the_age_bands_file_are_not_overwritten():
    spec = load_template_spec(EX)
    with tmpdir() as d:
        generate_builder(spec, d / "b.xlsx")
        got, rep = read_builder(_edit_bands(d), spec)
        assert got is None and rep.problems[0]["check"] == "file_bands_changed"


def _edit_bands(d):
    wb = load_workbook(d / "b.xlsx")
    ws = wb["Age bands"]
    ws.cell(HEADER_ROW + 2, 2).value = "renamed"
    wb.save(d / "e.xlsx")
    return d / "e.xlsx"


# ---------------------------------------------------------------------------------------------- edits and kept fields
def test_an_edit_changes_only_its_object_and_unknown_fields_are_kept():
    spec = load_template_spec(SPECS["C_ALPHA"])
    base = json.loads(json.dumps(_strip(spec)))
    base["_note"] = "kept"
    base["options"] = {"outlier_band": [0.6, 1.4], "language": "en"}
    base["sheets"][0]["x_future"] = {"a": 1}
    base["sheets"][0]["blocks"][0]["x_future"] = [1]
    base["_base_dir"] = spec["_base_dir"]
    with tmpdir() as d:
        g = generate_builder(base, d / "b.xlsx")
        assert "top-level '_note'" in g.not_carried and "sheet 'INDIVIDUAL RGN_A': 'x_future'" in g.not_carried
        _edit(d / "b.xlsx", d / "e.xlsx", [("Blocks", "Label", 1, "Green TPA (renamed)"),
                                            ("Tender", None, ("field", "options", "outlier_band_high"), 1.5)])
        got, rep = read_builder(d / "e.xlsx", base)
        assert rep.status == "PASS", rep.problems
        assert got["_note"] == "kept" and got["options"] == {"outlier_band": [0.6, 1.5], "language": "en"}
        assert got["sheets"][0]["x_future"] == {"a": 1} and got["sheets"][0]["blocks"][0]["x_future"] == [1]
        b = got["sheets"][0]["blocks"]
        assert b[1]["label"] == "Green TPA (renamed)"
        assert b[0] == base["sheets"][0]["blocks"][0] and b[2] == base["sheets"][0]["blocks"][2]
        assert got["sheets"][0]["factor_tables"] == base["sheets"][0]["factor_tables"]
        # without the base the unknown fields are dropped with a builder read that still passes
        got2, rep2 = read_builder(d / "e.xlsx", None)
        assert rep2.status == "PASS" and "_note" not in got2


def test_new_rows_extend_the_spec():
    """A new sheet, block, plan, section and refer cell typed into spare rows are read (header scan)."""
    spec = load_template_spec(SPECS["T_ALPHA"])
    with tmpdir() as d:
        generate_builder(spec, d / "b.xlsx")
        rows, _, _ = _values(d / "b.xlsx")
        ns, nb, npl = len(rows["Sheets"]), len(rows["Blocks"]), len(rows["Plans"])
        nsec, nrc = len(rows["Sections"]), len(rows["Refer cells"])
        band = rows["Age bands"][0]["Band set"]
        edits = [("Sheets", "Sheet name", ns, "SME RGN_B"), ("Sheets", "Product type", ns, "SME"),
                 ("Sheets", "Regulator", ns, "DHA"), ("Sheets", "Rating region", ns, "RGN_B"),
                 ("Blocks", "Sheet", nb, "SME RGN_B"), ("Blocks", "Block id", nb, "SME_B|TPA_RED"), ("Blocks", "TPA", nb, "TPA_RED"),
                 ("Plans", "Sheet", npl, "SME RGN_B"), ("Plans", "Block id", npl, "SME_B|TPA_RED"), ("Plans", "Plan id", npl, "S1"),
                 ("Plans", "Network", npl, "N_STD"), ("Plans", "Plan label", npl, "SME plan 1"),
                 ("Sections", "Sheet", nsec, "SME RGN_B"), ("Sections", "Applies to", nsec, APPLIES_SHEET),
                 ("Sections", "Section label", nsec, "Employee"), ("Sections", "Member category", nsec, "EMPLOYEE"),
                 ("Sections", "Age band set", nsec, band),
                 ("Refer cells", "Sheet", nrc, "SME RGN_B"), ("Refer cells", "Block id", nrc, "SME_B|TPA_RED"),
                 ("Refer cells", "Gender", nrc, "F")]
        _edit(d / "b.xlsx", d / "e.xlsx", edits)
        got, rep = read_builder(d / "e.xlsx", spec)
        assert rep.status == "PASS", rep.problems
        sh = got["sheets"][-1]
        assert sh["sheet_name"] == "SME RGN_B" and sh["sections"][0]["key"] == {"member_category": "EMPLOYEE"}
        assert sh["blocks"][0]["plans"] == [{"plan_id": "S1", "label": "SME plan 1", "network": "N_STD"}]
        assert sh["blocks"][0]["refer_cells"] == [{"gender": "F"}]
        assert got["sheets"][:-1] == spec["sheets"]
        generate_template(got, d / "t.xlsx")


def test_generate_refuses_an_invalid_spec():
    spec = load_template_spec(SPECS["T_ALPHA"])
    bad = {**spec, "premium_basis": "BOTH"}
    with tmpdir() as d:
        raises(SpecError, generate_builder, bad, d / "b.xlsx")


def test_cli_spec_builder_and_read_keep_the_previous_spec():
    import subprocess
    import sys
    pw = [sys.executable, str(REPO / "tools" / "pw.py")]
    with tmpdir() as d:
        shutil.copy(SPECS["C_ALPHA"], d / "template_spec.json")
        r = subprocess.run(pw + ["t", "spec-builder", str(d / "template_spec.json"), "--out", str(d / "b.xlsx")],
                           capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0 and "Spec Builder" in r.stdout, r.stdout + r.stderr
        r = subprocess.run(pw + ["t", "spec-builder", str(d / "missing.json"), "--blank-if-missing", "--out", str(d / "blank.xlsx")],
                           capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0 and "blank" in r.stdout, r.stdout + r.stderr
        before = (d / "template_spec.json").read_bytes()
        r = subprocess.run(pw + ["t", "spec-read", str(d / "b.xlsx"), "--base", str(d / "template_spec.json"), "--out",
                                 str(d / "template_spec.json")], capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0 and "SPEC READ PASS" in r.stdout, r.stdout + r.stderr
        assert (d / "template_spec.previous.json").read_bytes() == before
        assert json.loads((d / "template_spec.json").read_text(encoding="utf-8")) == json.loads(before)
        # a failing read leaves the spec untouched and exits 2
        _edit(d / "b.xlsx", d / "bad.xlsx", [("Sheets", "Product type", 0, None)])
        (d / "template_spec.previous.json").unlink()
        r = subprocess.run(pw + ["t", "spec-read", str(d / "bad.xlsx"), "--base", str(d / "template_spec.json"), "--out",
                                 str(d / "template_spec.json")], capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 2 and "SPEC READ FAIL" in r.stdout and "Sheets!" in r.stdout, r.stdout + r.stderr
        assert (d / "template_spec.json").read_bytes() == before and not (d / "template_spec.previous.json").exists()
        assert (d / "template_spec_read_report.csv").exists()


def test_training_tender_spec_steps_run_as_the_step_files_do():
    """Steps 0a / 0b from the tender folder: a pre-filled builder, then a read-back to the same spec (previous kept);
    steps 1 to 3 still run on the rewritten spec. The laptop check reads the same commands from the .bat files."""
    import importlib.util
    import os
    import subprocess
    import sys
    from pw.tender_kit import SPEC_STEPS, TRAINING_NAME, step_args, write_folder
    spec = importlib.util.spec_from_file_location("laptop_check_ws9", REPO / "tools" / "laptop_check.py")
    lc = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(lc)
    with tmpdir() as d:
        t = d / TRAINING_NAME
        write_folder(t, training=True)
        cmds = lc.spec_step_commands(t)
        assert cmds == [[x.replace("\\", os.sep) for x in a.split()] for _, _, a in SPEC_STEPS]
        before = (t / "inputs" / "template_spec.json").read_bytes()
        for args, text in zip(cmds, ("Spec Builder", "SPEC READ PASS")):
            r = subprocess.run([sys.executable, str(REPO / "tools" / "pw.py"), *args], cwd=str(t), capture_output=True,
                               text=True, timeout=600)
            assert r.returncode == 0 and text in r.stdout, (args, r.stdout[-1500:], r.stderr[-1500:])
        assert (t / "work" / "spec_builder.xlsx").exists()
        assert (t / "inputs" / "template_spec.previous.json").read_bytes() == before
        assert json.loads((t / "inputs" / "template_spec.json").read_text(encoding="utf-8")) == json.loads(before)
        for i, text in ((0, "structure hash"), (1, "CEDANT READ PASS"), (2, "structure hash")):
            r = subprocess.run([sys.executable, str(REPO / "tools" / "pw.py"), *step_args(i)], cwd=str(t),
                               capture_output=True, text=True, timeout=600)
            assert r.returncode == 0 and text in r.stdout, (i, r.stdout[-1500:], r.stderr[-1500:])


def test_blank_tender_folder_step_0a_makes_a_blank_builder_and_0b_reads_a_new_spec():
    import subprocess
    import sys
    from pw.tender_kit import SPEC_STEPS, write_folder
    with tmpdir() as d:
        t = d / "NEW"
        write_folder(t)
        args = [x.replace("\\", "/") for x in SPEC_STEPS[0][2].split()]
        (t / "work").mkdir(exist_ok=True)
        r = subprocess.run([sys.executable, str(REPO / "tools" / "pw.py"), *args], cwd=str(t), capture_output=True, text=True)
        assert r.returncode == 0 and "blank" in r.stdout, r.stdout + r.stderr
        # a user fills the blank builder with the fixture's content (as in the blank-fill test) and runs 0b
        shutil.copy(FIX / "C_ALPHA" / "spec_builder.xlsx", t / "inputs" / "spec_builder.xlsx")
        args = [x.replace("\\", "/") for x in SPEC_STEPS[1][2].split()]
        r = subprocess.run([sys.executable, str(REPO / "tools" / "pw.py"), *args], cwd=str(t), capture_output=True, text=True)
        assert r.returncode == 0 and "SPEC READ PASS" in r.stdout and "reading without a base" in r.stdout, r.stdout + r.stderr
        got = load_template_spec(t / "inputs" / "template_spec.json")
        assert _semantic(got) == _semantic(load_template_spec(SPECS["C_ALPHA"]))
        assert not (t / "inputs" / "template_spec.previous.json").exists()


def test_a_builder_can_be_read_again_after_a_fix_but_not_over_someone_elses_change():
    import subprocess
    import sys
    pw = [sys.executable, str(REPO / "tools" / "pw.py")]
    with tmpdir() as d:
        shutil.copy(SPECS["C_ALPHA"], d / "template_spec.json")
        subprocess.run(pw + ["t", "spec-builder", str(d / "template_spec.json"), "--out", str(d / "b.xlsx")], check=True,
                       capture_output=True)
        read = pw + ["t", "spec-read", str(d / "e.xlsx"), "--base", str(d / "template_spec.json"), "--out",
                     str(d / "template_spec.json")]
        _edit(d / "b.xlsx", d / "e.xlsx", [("Blocks", "Label", 0, "first edit")])
        r = subprocess.run(read, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout
        _edit(d / "e.xlsx", d / "e.xlsx", [("Blocks", "Label", 0, "second edit")])     # same builder, read again
        r = subprocess.run(read, capture_output=True, text=True)
        assert r.returncode == 0 and "read before" in r.stdout, r.stdout
        assert json.loads((d / "template_spec.json").read_text(encoding="utf-8"))["sheets"][0]["blocks"][0]["label"] == "second edit"
        # someone else changes the spec: the builder no longer matches the current spec or the previous one
        spec = json.loads((d / "template_spec.json").read_text(encoding="utf-8"))
        spec["sheets"][0]["title"] = "edited by someone else"
        (d / "template_spec.json").write_text(json.dumps(spec), encoding="utf-8")
        (d / "template_spec.previous.json").write_text(json.dumps(spec), encoding="utf-8")
        r = subprocess.run(read, capture_output=True, text=True)
        assert r.returncode == 2 and "spec_changed" in r.stdout, r.stdout
