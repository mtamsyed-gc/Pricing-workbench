"""WS5 on the structure-only cedant example (SPEC §8.4 / §8.5): the cedant workbook generates with every block, factor
family, rider and plan default, carries no rates or factor values, and a synthetic fill reads back and prices.
Expected structure is read from the example folder so no cedant labels appear in tests/."""
import json

import pandas as pd
from _util import REPO, load_json, tmpdir
from openpyxl import load_workbook

from pw.evaluator import evaluate
from pw.template.cedant import FACTORS, RIDERS, generate_cedant_workbook, read_cedant
from pw.template.generator import generate_template
from pw.template.parser import fill_template, synthetic_value
from pw.template.spec import cedant_model, load_template_spec, spec_hash

EX = REPO / "examples" / "ufic_2027"
REQUIRES = ("examples/ufic_2027",)          # repo only: not shipped in a release zip


def _spec():
    return load_template_spec(EX / "template_spec.json")


def test_example_cedant_workbook_has_every_block_family_rider_and_default():
    spec = _spec()
    exp = load_json(EX / "expected_structure.json")["cedant"]
    model, errs = cedant_model(spec)
    assert errs == []
    got = {}
    for m in model["sheets"]:
        got[m["sheet_name"]] = {
            "groups": {g["label"]: {"factor_tables": g["factor_tables"], "riders": g["riders"],
                                    "plan_defaults": g["default_families"], "plans": len(g["plans"])} for g in m["groups"]},
            "factor_rows": sum(len(t["rows"]) for t in m["factor_tables"]),
            "rider_rows": sum(len(t["rows"]) for t in m["riders"]), "benefit_codes": len(m["benefits"])}
    assert got == exp["sheets"]
    with tmpdir() as d:
        g = generate_cedant_workbook(spec, d / "c.xlsx")
        t = generate_template(spec, d / "t.xlsx")
        # every tender block / rate cell is in the cedant workbook at the same place (one layout, two modes)
        rate_c = sorted((c.cell_key, c.sheet, c.address, c.refer) for c in g.layout.cells)
        rate_t = sorted((c.cell_key, c.sheet, c.address, c.refer) for c in t.layout.cells)
        assert rate_c == rate_t and g.counts["RATE"] == exp["rate_cells"] and g.n_refer == exp["refer_cells"]
        wb = load_workbook(d / "c.xlsx")
        # structure only: every input cell is blank; the long tables hold labels only
        for r in wb["_KEYS"].iter_rows(min_row=2, values_only=True):
            if r[2] != "REFER":
                assert wb[r[5]][r[6]].value is None, r[0]
        for s in (FACTORS, RIDERS, "Plan defaults", "Benefits"):
            for row in wb[s].iter_rows(min_row=5, values_only=True):
                assert not any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in row), (s, row)
        refs = [r for r in wb[FACTORS].iter_rows(min_row=5, values_only=True) if r[8] == "Y"]
        assert len(refs) == sum(r["reference"] for m in model["sheets"] for t2 in m["factor_tables"] for r in t2["rows"])
        for m in model["sheets"]:                       # every key has its own reference or inherits the ANY row's
            for t2 in m["factor_tables"]:
                for k in {r["key"] for r in t2["rows"]}:
                    own = sum(r["reference"] for r in t2["rows"] if r["key"] == k)
                    assert own == 1 or (own == 0 and any(r["reference"] for r in t2["rows"] if r["key"] == "ANY"))
    # the cedant sections never change the reinsurer template (spec hash unchanged)
    bare = json.loads((EX / "template_spec.json").read_text(encoding="utf-8"))
    bare.pop("cedant_workbook")
    for sh in bare["sheets"]:
        for k in ("scope_id", "factor_tables", "riders", "plan_defaults", "benefits", "premium_basis"):
            sh.pop(k, None)
    assert spec_hash(bare) == spec_hash(spec)


def test_example_synthetic_fill_reads_back_and_prices():
    """Synthetic values only (never the cedant's): the example workbook reads back into one scope per TPA group and
    prices a cell of every scope."""
    spec = _spec()
    model, _ = cedant_model(spec)
    opts = {}
    for m in model["sheets"]:
        for g in m["groups"]:
            for fam in g["default_families"]:
                t = next(x for x in m["factor_tables"] + m["riders"] if x["family"] == fam and set(g["tpas"]) <= set(x["tpas"]))
                ref = next((r["option"] for r in t["rows"] if r.get("reference")), None)
                opts[(m["sheet_name"], tuple(g["tpas"]), fam)] = ref or "NONE"

    def value_fn(ck, keys, kind):
        k = keys[0]
        if kind == "RATE":
            return synthetic_value(ck, 7)
        if kind == "FACTOR":
            t = next(x for m in model["sheets"] if m["sheet_name"] == k["sheet"] for x in m["factor_tables"]
                     if x["component_id"] == k["component_id"] and x["tpas"] == k["tpas"])
            r = next(r for r in t["rows"] if r["key"] == k["key"] and r["option"] == k["option"])
            return 1.0 if r["reference"] else 0.95
        if kind == "RIDER":
            return 123.45
        if kind == "DEFAULT":
            return opts[(k["sheet"], tuple(k["tpas"]), k["family"])]
        return None
    with tmpdir() as d:
        generate_cedant_workbook(spec, d / "c.xlsx")
        fill_template(d / "c.xlsx", d / "f.xlsx", value_fn=value_fn)
        cat, rep = read_cedant(d / "f.xlsx", spec)
        assert rep.status == "PASS", rep.problems[:5]
        assert cat.kind == "EXPIRING" and cat.scopes
        assert len(cat.scopes) == sum(len(m["groups"]) for m in model["sheets"])
        cells = []
        for s in cat.scopes:
            p = s.plans[0]
            t = s.tpa_set[0]
            base = next(c for c in s.components if c.type == "BASE_TABLE")
            row = next(r for r in base.rows if r["status"] == "VALUE")
            cell = {"cell_id": len(cells), "product_type": s.product_type, "rating_region": s.rating_region, "tpa": t,
                    "plan": p.plan_id, "nationality_group": "OTHER", "maternity_eligible": "N"}
            for dim, v in row["key"].items():
                if dim not in ("plan", "tpa") and v != "ANY":
                    cell[dim] = v[0] if isinstance(v, list) else v
            cells.append(cell)
        res, _ = evaluate(cat, pd.DataFrame(cells))
        assert (res["status"] == "PRICED").all(), res[res["status"] != "PRICED"][["scope_id", "reason"]].to_string()
