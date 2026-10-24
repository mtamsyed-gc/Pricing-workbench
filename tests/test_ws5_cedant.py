"""WS5 cedant rate workbook (SPEC §8.5): authoring from a catalogue's structure, generation, scripted fill, gated
read-back into the EXPIRING catalogue, cell-level pricing identity on EXPO_SYN, the reinsurer template from the
read-back, and every read-back gate one by one."""
import copy
import hashlib
import json
import shutil
import subprocess
import sys

import numpy as np
import pandas as pd
from _util import FIX, REPO, catalogue, raises, tmpdir
from openpyxl import load_workbook

from pw import GateFailed, SpecError
from pw.evaluator import evaluate
from pw.exposure import build_exposure
from pw.io_excel import make_reproducible
from pw.model import Catalogue
from pw.template.cedant import (BENEFITS, DEFAULTS, FACTORS, RIDERS, fill_from_catalogue, generate_cedant_workbook,
                                read_cedant, read_cedant_workbook)
from pw.template.generator import generate_template
from pw.template.spec import cedant_spec_from_catalogue, spec_hash, validate_template_spec

C = FIX / "C_ALPHA"
PW = [sys.executable, str(REPO / "tools" / "pw.py")]


def _spec():
    return json.loads((C / "template_spec.json").read_text(encoding="utf-8"))


def _cells():
    d = FIX / "EXPO_SYN"
    return build_exposure(d / "dim_life.csv", d / "maps", d / "run_manifest.csv", strict=False).cells


def _h(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


# ------------------------------------------------------------------------------------------------ authoring / layout
def test_spec_from_catalogue_keeps_the_tender_layout():
    spec = _spec()
    assert validate_template_spec(spec) == []
    tender = json.loads((FIX / "T_ALPHA" / "template_spec.json").read_text(encoding="utf-8"))
    rate_part = {k: v for k, v in spec.items() if k != "cedant_workbook"}
    rate_part["sheets"] = [{k: v for k, v in sh.items() if k not in ("scope_id", "formula", "premium_basis", "factor_tables",
                                                                     "riders", "plan_defaults", "benefits")}
                           for sh in spec["sheets"]]
    assert rate_part == tender                                   # one layout: same blocks, plans, bands, columns
    assert spec_hash(spec) == spec_hash(tender)                  # cedant sections never change the tender's hash
    sh = spec["sheets"][0]
    assert [t["family"] for t in sh["factor_tables"]] == ["copay_op_diag", "ded_consult", "copay_pharmacy", "aml",
                                                          "network_option", "tpa", "nationality"]
    assert [t["family"] for t in sh["riders"]] == ["module_maternity"]
    assert sh["plan_defaults"] == ["copay_op_diag", "ded_consult", "copay_pharmacy", "aml", "network_option", "module_maternity"]
    assert "rules (3)" in spec["cedant_workbook"]["not_carried"]


def test_generated_workbook_layout_keys_and_determinism():
    spec = _spec()
    with tmpdir() as d:
        g1 = generate_cedant_workbook(spec, d / "a.xlsx")
        g2 = generate_cedant_workbook(spec, d / "b.xlsx")
        assert _h(d / "a.xlsx") == _h(d / "b.xlsx") and g1.structure_hash == g2.structure_hash
        wb = load_workbook(d / "a.xlsx")
        assert wb.sheetnames == ["Cover", "INDIVIDUAL RGN_A", FACTORS, RIDERS, DEFAULTS, BENEFITS, "_LISTS", "_KEYS", "_META"]
        assert all(wb[s].sheet_state == "hidden" for s in ("_LISTS", "_KEYS", "_META"))
        meta = {r[0]: r[1] for r in wb["_META"].iter_rows(min_row=2, values_only=True)}
        assert meta["mode"] == "CEDANT" and meta["structure_hash"] == g1.structure_hash
        kinds = pd.Series([r[2] for r in wb["_KEYS"].iter_rows(min_row=2, values_only=True)]).value_counts().to_dict()
        assert kinds == {"RATE": 618, "FACTOR": 49, "RIDER": 21, "DEFAULT": 42, "BENEFIT": 7, "BENEFIT_TEXT": 7, "REFER": 12}
        assert g1.counts == {"RATE": 618, FACTORS: 49, RIDERS: 21, DEFAULTS: 42, BENEFITS: 7}
        ws = wb[FACTORS]
        hdr = [c.value for c in ws[4]]
        assert hdr[:9] == ["Product", "Emirate", "TPAs", "Family", "Key on", "Key", "Option", "Factor", "Reference"]
        f = ws["H5"]
        assert f.value is None and f.fill.fgColor.rgb.endswith("FFFF99") and f.protection.locked is False
        assert ws["G5"].protection.locked is not False and ws.protection.sheet
        refs = [(r[3], r[5], r[6]) for r in ws.iter_rows(min_row=5, values_only=True) if r[8] == "Y"]
        assert ("copay_op_diag", "ANY", "0%") in refs and ("aml", "N_PLAT", "2000000") in refs
        assert ("tpa", "ANY", "TPA_RED") in refs and ("tpa", "ANY", "TPA_BLUE") not in refs   # one reference per key
        types = {dv.type for s in (FACTORS, DEFAULTS, BENEFITS) for dv in wb[s].data_validations.dataValidation}
        assert {"decimal", "list"} <= types
        assert "TPA basis" in [c.value for c in wb["Cover"]["A"]]
        # no value is pre-filled anywhere an input lives (structure only)
        krows = list(wb["_KEYS"].iter_rows(min_row=2, values_only=True))
        assert all(wb[r[5]][r[6]].value in (None,) for r in krows if r[2] != "REFER")
        assert _h(d / "a.xlsx") == _h(C / "cedant_workbook.xlsx")


# ------------------------------------------------------------------------------------------------ round trip
def test_round_trip_prices_expo_syn_cell_for_cell():
    """SPEC §8.5 / WS5 acceptance: CAT_ALPHA -> cedant workbook -> filled -> read back prices EXPO_SYN exactly as
    CAT_ALPHA (tolerance 0), at the catalogue defaults, at seeded explicit options, and with fill = reference."""
    alpha = catalogue("CAT_ALPHA")
    back = Catalogue.read(C / "expiring.json")
    assert back.kind == "EXPIRING" and back.catalogue_id == "CAT_ALPHA" and back.meta["mode"] == "CEDANT"
    with tmpdir() as d:                                          # the fixture is what the reader returns today
        fresh = read_cedant_workbook(C / "filled.xlsx", _spec())
        fresh.write(d / "x.json")
        assert (d / "x.json").read_text(encoding="utf-8") == (C / "expiring.json").read_text(encoding="utf-8")
    cells = _cells()
    rng = np.random.default_rng(5)
    opt = cells.copy()
    choices = {"copay_op_diag": ["0%", "10%", "20%", "AED25"], "ded_consult": ["NIL", "D30", "D60", "P15_C40"],
               "copay_pharmacy": ["0%", "15%", "25%"], "aml": ["300000", "750000", "1250000", "2000000"],
               "network_option": ["OPT1", "OPT2"], "module_maternity": ["M10", "M20", "MAML", "NONE"]}
    for col, vals in choices.items():
        pick = rng.choice(vals + [None], size=len(opt))
        opt[col] = pd.Series(pick, index=opt.index, dtype=object)
    cols = ["cell_id", "status", "reason", "base", "factor_product", "modules_add", "premium_pre_loading", "price",
            "defaulted", "base_status", "price_base", "price_riders"]
    n_priced = 0
    for cs, o in ((cells, {}), (opt, {}), (cells, {"fill": "reference"}), (opt, {"fill": "reference"})):
        a, ta = evaluate(alpha, cs, {**o, "trace": True})
        b, tb = evaluate(back, cs, {**o, "trace": True})
        a, b = a[cols], b[cols]
        for c in cols:
            if a[c].dtype == "float64":
                assert np.array_equal(a[c].to_numpy(), b[c].to_numpy(), equal_nan=True), (o, c)
            elif c != "reason":
                assert (a[c].astype(str) == b[c].astype(str)).all(), (o, c)
        n_priced += int((a["status"] == "PRICED").sum())
        # same components touch the same cells with the same values (component ids are carried over)
        key = ["cell_id", "component_id", "value"]
        ta = ta[ta["type"].isin(["FACTOR_TABLE", "MODULE", "FORMULA"])][key].sort_values(key).reset_index(drop=True)
        tb = tb[tb["type"].isin(["FACTOR_TABLE", "MODULE", "FORMULA"])][key].sort_values(key).reset_index(drop=True)
        assert ta.equals(tb), o
    assert n_priced > 4 * 2000
    # benefit lines, plan defaults, networks and reference levels carried over
    sa, sb = alpha.scopes[0], back.scopes[0]
    assert [(p.plan_id, p.tier_rank, p.defaults, p.network_by_tpa) for p in sa.plans] == \
        [(p.plan_id, p.tier_rank, p.defaults, p.network_by_tpa) for p in sb.plans]
    assert [(b.plan_id, b.benefit_code, b.value, b.condition_text) for b in sa.benefits] == \
        [(b.plan_id, b.benefit_code, b.value, b.condition_text) for b in sb.benefits]
    for ca in sa.components:
        if ca.type in ("FACTOR_TABLE", "MODULE"):
            cb = next(c for c in sb.components if c.component_id == ca.component_id)
            assert (ca.keys, ca.reference_level, ca.reference_level_by, ca.reference_level_by_key, ca.unit, ca.eligibility) == \
                (cb.keys, cb.reference_level, cb.reference_level_by, cb.reference_level_by_key, cb.unit, cb.eligibility)
            assert [(r["key"], r["value"]) for r in ca.rows] == [(r["key"], r["value"]) for r in cb.rows]


def test_reinsurer_template_from_the_read_back_equals_t_alpha():
    """Reinsurer mode unchanged: the tender template generated from the read-back catalogue (same spec, cedant
    sections included) is byte-identical to T_ALPHA (Existing Rates, Factor Tables, Add-ons from the read-back)."""
    with tmpdir() as d:
        generate_template(_spec(), d / "r.xlsx", expiring=Catalogue.read(C / "expiring.json"))
        assert _h(d / "r.xlsx") == _h(FIX / "T_ALPHA" / "template.xlsx")


def test_read_back_drives_v5_like_the_source_catalogue():
    from pw.gates import v5_submission_integrity
    from pw.template.parser import parse_submission, read_meta
    tspec = json.loads((FIX / "T_ALPHA" / "template_spec.json").read_text(encoding="utf-8"))
    parsed = parse_submission(FIX / "SUB_Y" / "submission.xlsx", template_spec=tspec)
    issued = read_meta(FIX / "T_ALPHA" / "template.xlsx")["structure_hash"]
    a = v5_submission_integrity(parsed, tspec, issued, catalogue("CAT_ALPHA"))
    b = v5_submission_integrity(parsed, tspec, issued, Catalogue.read(C / "expiring.json"))
    assert a["counts"] == b["counts"] and a["flags"].equals(b["flags"])


# ------------------------------------------------------------------------------------------------ gates
def _find(ws, **want):
    """1-based row of the first long-table row whose columns equal the given values."""
    hdr = {c.value: c.column for c in ws[4] if c.value}
    for r in range(5, ws.max_row + 1):
        if all(str(ws.cell(r, hdr[k]).value) == str(v) for k, v in want.items()):
            return r, hdr
    raise AssertionError(f"row {want} not found")


def _spare(ws):
    hdr = {c.value: c.column for c in ws[4] if c.value}
    for r in range(5, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").startswith("Spare rows"):
            return r + 1, hdr
    raise AssertionError("no spare rows")


def _mutated(d, fn, name="m.xlsx"):
    wb = load_workbook(C / "filled.xlsx")
    fn(wb)
    wb.save(d / name)
    make_reproducible(d / name)
    return read_cedant(d / name, _spec())


def _rate_addr(wb, label_part):
    for r in wb["_KEYS"].iter_rows(min_row=2, values_only=True):
        if r[2] == "RATE" and label_part in r[0]:
            return r[5], r[6]
    raise AssertionError(label_part)


def test_gate_failures_one_by_one():
    cases = []

    def case(name, fn, check, text, cell=None):
        cases.append((name, fn, check, text, cell))

    def set_rate(v):
        def f(wb):
            s, a = _rate_addr(wb, "plan=P3|age=40-44|gender=F")
            wb[s][a].value = v
        return f
    case("blank mandatory rate", set_rate(None), "rate", "blank mandatory rate Plan 3 / 40-44 / Female", "INDIVIDUAL RGN_A!")
    case("non-numeric rate", set_rate("n/a"), "rate", "not a number")
    case("rate <= 0", set_rate(0), "rate", "must be > 0")

    def set_factor(col, v, **at):
        def f(wb):
            ws = wb[FACTORS]
            r, h = _find(ws, **at)
            ws.cell(r, h[col]).value = v
        return f
    case("blank factor", set_factor("Factor", None, Family="copay_pharmacy", Option="15%"), "factor", "blank factor")
    case("non-numeric factor", set_factor("Factor", "x", Family="copay_pharmacy", Option="15%"), "factor", "not a number")
    case("factor <= 0", set_factor("Factor", -0.5, Family="copay_pharmacy", Option="15%"), "factor", "must be > 0")
    case("no reference row", set_factor("Reference", None, Family="copay_op_diag", Option="0%"), "reference",
         "key 'ANY' has no reference row")
    case("two reference rows", set_factor("Reference", "Y", Family="tpa", Option="TPA_BLUE"), "reference",
         "has 2 reference rows")
    case("reference factor not 1.0", set_factor("Factor", 1.02, Family="aml", Key="N_GOLD", Option="1250000"),
         "reference", "must read 1.0")
    case("bad reference flag", set_factor("Reference", "maybe", Family="ded_consult", Option="D60"), "reference",
         "must be Y or blank")

    def set_default(v):
        def f(wb):
            ws = wb[DEFAULTS]
            r, h = _find(ws, Plan="P4", Family="copay_op_diag")
            ws.cell(r, h["Option"]).value = v
        return f
    case("default names unknown option", set_default("15%"), "default", "is not an option of its table")
    case("blank plan default", set_default(None), "default", "blank plan default P4 / copay_op_diag")

    def spare_row(sheet, **vals):
        def f(wb):
            ws = wb[sheet]
            r, h = _spare(ws)
            for k, v in vals.items():
                ws.cell(r, h[k]).value = v
        return f
    pe = {"Product": "INDIVIDUAL", "Emirate": "RGN_A", "TPAs": "ALL"}
    case("unknown benefit code", spare_row(BENEFITS, **pe, Plan="P2", **{"Benefit code": "spa_days"}, Value=3), "benefit",
         "unknown benefit code 'spa_days'")

    def ladder(wb):
        ws = wb[BENEFITS]
        r, h = _find(ws, Plan="P1", **{"Benefit code": "room_type"})
        ws.cell(r, h["Value"]).value = "PALACE"
    case("unknown ladder value", ladder, "benefit", "is not a ladder value of room_type")
    case("duplicate factor row", spare_row(FACTORS, **pe, Family="copay_op_diag", Key="ANY", Option="10%", Factor=0.95),
         "duplicate", "duplicate row copay_op_diag key=ANY option=10%")
    case("unknown product x emirate", spare_row(FACTORS, Product="SME", Emirate="RGN_A", TPAs="ALL", Family="aml",
                                                Option="1", Factor=1.0), "unknown", "no rate sheet for product 'SME'")
    case("no such table", spare_row(FACTORS, **pe, Family="copay_ip", Option="10%", Factor=0.9), "unknown",
         "no factor table of family 'copay_ip'")

    def rider_blank(wb):
        ws = wb[RIDERS]
        r, h = _find(ws, Key="N_GOLD", Limit="M20")
        ws.cell(r, h["Rate"]).value = None
    case("blank rider rate", rider_blank, "rate", "blank rate for module_maternity N_GOLD/M20")
    case("duplicate plan default", spare_row(DEFAULTS, **pe, Plan="P1", Family="aml", Option="2000000"), "duplicate",
         "duplicate plan default P1 / aml")

    def delete_row(wb):
        ws = wb[FACTORS]
        r, h = _find(ws, Family="ded_consult", Option="D60")
        for c in range(1, len(h) + 1):
            ws.cell(r, c).value = None
    case("deleted factor row", delete_row, "missing", "row key=ANY option=D60 is missing")
    with tmpdir() as d:
        for i, (name, fn, check, text, cell) in enumerate(cases):
            cat, rep = _mutated(d, fn, f"m{i}.xlsx")
            assert cat is None, name
            hit = [p for p in rep.problems if p["check"] == check and text in p["message"]]
            assert hit, (name, rep.problems[:5])
            assert all(p["sheet"] for p in rep.problems), name
            if cell:
                assert any(f"{p['sheet']}!{p['cell']}".startswith(cell) and p["cell"] for p in hit), (name, hit)
            assert len(rep.problems) == len(hit) or name in ("no such table",), (name, rep.problems)
        # the wrapper raises with the report attached
        e = raises(GateFailed, read_cedant_workbook, d / "m0.xlsx", _spec())
        assert e.gate == "CEDANT_READ" and e.details.problems
        # an empty (unfilled) workbook stops with every mandatory cell listed
        cat, rep = read_cedant(C / "cedant_workbook.xlsx", _spec())
        assert cat is None and sum("blank mandatory rate" in p["message"] for p in rep.problems) == 618
        assert sum(p["message"].startswith("blank rate for module_maternity") for p in rep.problems) == 21
        assert sum(p["message"].startswith("blank factor") for p in rep.problems) == 49
        assert sum(p["check"] == "default" for p in rep.problems) == 42
        # a reinsurer template is not a cedant workbook
        cat, rep = read_cedant(FIX / "T_ALPHA" / "template.xlsx", _spec())
        assert cat is None and any(p["check"] == "mode" for p in rep.problems)


def test_spare_rows_extend_existing_tables():
    def add(wb):
        ws = wb[FACTORS]
        r, h = _spare(ws)
        for k, v in {"Product": "INDIVIDUAL", "Emirate": "RGN_A", "TPAs": "ALL", "Family": "copay_op_diag", "Option": "30%",
                     "Factor": 0.88}.items():
            ws.cell(r, h[k]).value = v
        ws = wb[BENEFITS]
        r, h = _spare(ws)
        for k, v in {"Product": "INDIVIDUAL", "Emirate": "RGN_A", "TPAs": "ALL", "Plan": "P3", "Benefit code": "territory",
                     "Value": "ww", "Condition text": "worldwide"}.items():
            ws.cell(r, h[k]).value = v
    with tmpdir() as d:
        cat, rep = _mutated(d, add)
        assert rep.status == "PASS", rep.problems
        f = next(c for c in cat.scopes[0].components if c.component_id == "f_opdiag")
        assert f.rows[-1]["key"] == {"copay_op_diag": "30%"} and f.rows[-1]["value"] == 0.88
        b = [x for x in cat.scopes[0].benefits if x.benefit_code == "territory"]
        assert len(b) == 1 and b[0].value == "WW" and b[0].plan_id == "P3"
        res, _ = evaluate(cat, pd.DataFrame([{"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_RED",
                                              "plan": "P1", "age": 33, "gender": "M", "copay_op_diag": "30%",
                                              "nationality_group": "OTHER", "maternity_eligible": "N"}]))
        res0, _ = evaluate(cat, pd.DataFrame([{"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_RED",
                                               "plan": "P1", "age": 33, "gender": "M", "copay_op_diag": "0%",
                                               "nationality_group": "OTHER", "maternity_eligible": "N"}]))
        assert res["status"].iloc[0] == "PRICED" and abs(res["factor_product"].iloc[0] / res0["factor_product"].iloc[0] - 0.88) < 1e-12


def test_lost_keys_fall_back_to_position_and_text_numbers_warn():
    def edit(wb):
        s, a = _rate_addr(wb, "plan=P1|age=0-4|gender=M")
        wb[s][a].value = str(wb[s][a].value)            # a rate typed as text
        del wb["_KEYS"]
    with tmpdir() as d:
        cat, rep = _mutated(d, edit)
        assert rep.status == "PASS" and rep.hash_ok is False
        assert any("positional" in w or "by position" in w for w in rep.warnings)
        assert any("read as the number" in w for w in rep.warnings)
        a, _ = evaluate(catalogue("CAT_ALPHA"), _cells())
        b, _ = evaluate(cat, _cells())
        assert np.array_equal(a["price"].to_numpy(), b["price"].to_numpy(), equal_nan=True)


def test_spec_validation_of_cedant_sections():
    spec = _spec()
    sh = spec["sheets"][0]

    def bad(mut, text):
        s = copy.deepcopy(spec)
        mut(s["sheets"][0])
        errs = validate_template_spec(s)
        assert any(text in e for e in errs), (text, errs)
        raises(SpecError, generate_cedant_workbook, s, REPO / "runs" / "never.xlsx")
    bad(lambda h: h["factor_tables"][0].update(reference="5%"), "has no reference row")
    bad(lambda h: h["factor_tables"][5].update(reference={"ANY": "TPA_RED"}, rows=["TPA_RED", "TPA_RED"]), "duplicate row")
    bad(lambda h: h["plan_defaults"].append("nationality"), "is not an option family")
    bad(lambda h: h["plan_defaults"].append("copay_ip"), "has no factor table or rider")
    bad(lambda h: h["benefits"].append("spa_days"), "not in registry/benefits.json")
    bad(lambda h: h["factor_tables"][0].update(tpas=["TPA_PINK"]), "have no block on the sheet")
    bad(lambda h: h["riders"][0].update(unit="PER_POLICY"), "unit must be one of")
    bad(lambda h: h["factor_tables"][3].update(keys=["aml"]), "must be exactly")
    bad(lambda h: h["riders"][0]["rows"].append({"key": "N_GOLD", "option": "NONE"}), "reserved")
    assert not (REPO / "runs" / "never.xlsx").exists()
    # TPA-scoped tables split the read-back into TPA groups (one scope per group)
    s = copy.deepcopy(spec)
    s["sheets"][0]["factor_tables"][6]["tpas"] = ["TPA_RED", "TPA_BLUE"]
    s["sheets"][0]["plan_defaults"] = [f for f in s["sheets"][0]["plan_defaults"]]
    from pw.template.spec import cedant_model
    m, errs = cedant_model(s)
    assert errs == []
    assert [(g["scope_id"], g["tpas"]) for g in m["sheets"][0]["groups"]] == \
        [("ALPHA_IND_A|TPA_RED+TPA_BLUE", ["TPA_RED", "TPA_BLUE"]), ("ALPHA_IND_A|TPA_GREEN", ["TPA_GREEN"])]
    assert sh["scope_id"] == "ALPHA_IND_A"


def test_tpa_groups_round_trip():
    """A factor table restricted to some TPAs: two read-back scopes; each prices its TPAs as the source would."""
    s = _spec()
    s["sheets"][0]["factor_tables"][6]["tpas"] = ["TPA_RED", "TPA_BLUE"]       # nationality factor on two TPAs only
    with tmpdir() as d:
        generate_cedant_workbook(s, d / "c.xlsx")
        fill_from_catalogue(d / "c.xlsx", d / "f.xlsx", catalogue("CAT_ALPHA"))
        cat = read_cedant_workbook(d / "f.xlsx", s)
        assert [sc.scope_id for sc in cat.scopes] == ["ALPHA_IND_A|TPA_RED+TPA_BLUE", "ALPHA_IND_A|TPA_GREEN"]
        assert "f_nat" not in [c.component_id for c in cat.scopes[1].components]
        cells = _cells()
        a, _ = evaluate(catalogue("CAT_ALPHA"), cells)
        b, _ = evaluate(cat, cells)
        red_blue = cells["tpa"].isin(["TPA_RED", "TPA_BLUE"]).to_numpy()
        assert np.array_equal(a["price"].to_numpy()[red_blue], b["price"].to_numpy()[red_blue], equal_nan=True)
        green = ~red_blue & (a["status"] == "PRICED").to_numpy()
        other = (cells["nationality_group"] != "OTHER").to_numpy()
        assert (b["factor_product"].to_numpy()[green & other] != a["factor_product"].to_numpy()[green & other]).all()


def test_cli_cedant_and_cedant_read():
    with tmpdir() as d:
        r = subprocess.run(PW + ["t", "cedant", str(C / "template_spec.json"), "--out", str(d / "c.xlsx")],
                           capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0 and "structure hash" in r.stdout, r.stderr
        r = subprocess.run(PW + ["t", "cedant-read", str(d / "c.xlsx"), "--spec", str(C / "template_spec.json"), "--out",
                                 str(d / "exp.json")], capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 2 and "CEDANT READ FAIL" in r.stdout and "blank mandatory rate" in r.stdout
        assert not (d / "exp.json").exists() and (d / "exp_read_report.csv").exists()
        shutil.copy(C / "filled.xlsx", d / "filled.xlsx")
        r = subprocess.run(PW + ["t", "cedant-read", str(d / "filled.xlsx"), "--spec", str(C / "template_spec.json"), "--out",
                                 str(d / "exp.json"), "--id", "CAT_ALPHA"], capture_output=True, text=True, cwd=str(REPO / "docs"))
        assert r.returncode == 0 and "CEDANT READ PASS" in r.stdout, r.stdout + r.stderr
        assert (d / "exp.json").read_text(encoding="utf-8") == (C / "expiring.json").read_text(encoding="utf-8")


# ------------------------------------------------------------------------------------------------ review fixes
def test_plan_default_must_exist_for_the_plans_own_key():
    def edit(wb):
        ws = wb[DEFAULTS]
        r, h = _find(ws, Plan="P1", Family="network_option")
        ws.cell(r, h["Option"]).value = "OPT2"          # OPT2 rows exist for P2 / P3 only
    with tmpdir() as d:
        cat, rep = _mutated(d, edit)
        assert cat is None and any("for key P1" in p["message"] for p in rep.problems), rep.problems


def test_inherited_reference_is_written_per_key_for_fill_reference():
    s = _spec()
    t = next(x for x in s["sheets"][0]["factor_tables"] if x["component_id"] == "f_netopt")
    t["reference"] = {"ANY": "OPT1", "P2": "OPT2"}             # P2 has its own reference; the others inherit ANY's

    def vf(ck, keys, kind, v):
        k = keys[0]
        return 1.0 if kind == "FACTOR" and k["component_id"] == "f_netopt" and k["key"] == "P2" else v
    with tmpdir() as d:
        generate_cedant_workbook(s, d / "c.xlsx")
        fill_from_catalogue(d / "c.xlsx", d / "f.xlsx", catalogue("CAT_ALPHA"), value_fn=vf)
        cat = read_cedant_workbook(d / "f.xlsx", s)
        c = next(x for x in cat.scopes[0].components if x.component_id == "f_netopt")
        assert c.reference_level is None and c.reference_level_by == "plan"
        assert c.reference_level_by_key == {**{f"P{i}": "OPT1" for i in range(1, 8)}, "P2": "OPT2"}
        res, _ = evaluate(cat, pd.DataFrame([{"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_RED",
                                              "plan": p, "age": 40, "gender": "M", "nationality_group": "OTHER",
                                              "maternity_eligible": "N"} for p in ("P2", "P3")]), {"fill": "reference"})
        assert (res["status"] == "PRICED").all()


def test_formula_bounds_and_component_defaults_are_carried():
    alpha = copy.deepcopy(catalogue("CAT_ALPHA"))
    alpha.scopes[0].formula.min_premium = 5000.0
    tender = {k: _spec()[k] for k in ("cedant", "tender_id", "deadline", "contact", "issue_date", "effective_date")}
    s = cedant_spec_from_catalogue(alpha, tender)
    assert s["sheets"][0]["formula"]["min_premium"] == 5000.0
    with tmpdir() as d:
        generate_cedant_workbook(s, d / "c.xlsx")
        fill_from_catalogue(d / "c.xlsx", d / "f.xlsx", alpha)
        cat = read_cedant_workbook(d / "f.xlsx", s)
        cells = _cells()
        a, _ = evaluate(alpha, cells)
        b, _ = evaluate(cat, cells)
        assert np.array_equal(a["price"].to_numpy(), b["price"].to_numpy(), equal_nan=True)
    # two scopes on one sheet with different formulas cannot share one cedant sheet
    two = copy.deepcopy(catalogue("CAT_ALPHA"))
    sc2 = copy.deepcopy(two.scopes[0])
    sc2.scope_id, sc2.tpa_set, sc2.formula.rounding = "ALPHA_B", ["TPA_PINK"], 0
    two.scopes[0].tpa_set = ["TPA_RED"]
    two.scopes.append(sc2)
    from pw import ProposalRequired
    raises(ProposalRequired, cedant_spec_from_catalogue, two, tender)


def test_spec_mismatch_missing_benefits_rename_and_negative_benefit():
    s = _spec()
    s["sheets"][0]["blocks"][0]["plans"][0]["network"] = "N_BASIC"
    cat, rep = read_cedant(C / "filled.xlsx", s)
    assert cat is None and any(p["check"] == "spec" for p in rep.problems)
    with tmpdir() as d:
        cat, rep = _mutated(d, lambda wb: wb.remove(wb[BENEFITS]), "nb.xlsx")
        assert cat is None and any("Benefits' is missing" in p["message"] for p in rep.problems)

        def rename(wb):                                  # as Excel does: sheet renamed, _KEYS formulas follow
            wb[FACTORS].title = "Factors X"
            for row in wb["_KEYS"].iter_rows(min_row=2):
                if isinstance(row[1].value, str) and row[1].value.startswith("='Factors'!"):
                    row[1].value = row[1].value.replace("='Factors'!", "='Factors X'!")
        cat, rep = _mutated(d, rename, "rn.xlsx")
        assert rep.status == "PASS", rep.problems[:3]

        def neg(wb):
            ws = wb[BENEFITS]
            r, h = _find(ws, Plan="P1", **{"Benefit code": "aml"})
            ws.cell(r, h["Value"]).value = -5
        cat, rep = _mutated(d, neg, "ng.xlsx")
        assert cat is None and any("is negative" in p["message"] for p in rep.problems)
