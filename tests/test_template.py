"""Template generator (T) core: spec validation, generation, protection/_KEYS/_META, V4 round trip on every fixture,
byte determinism, tolerance to renamed sheets and lost keys, splits, SUB_X / SUB_Y."""
import copy
import hashlib
import json
import re

import pandas as pd
from _util import CATS, FIX, TENDER, catalogue, raises, tmpdir
from openpyxl import load_workbook

from pw import SpecError
from pw.evaluator import evaluate
from pw.gates import v4_round_trip
from pw.model import Catalogue
from pw.template.generator import generate_template
from pw.template.parser import fill_template, parse_submission, to_catalogue
from pw.template.spec import normalise, spec_from_catalogue, validate_template_spec


def _spec(name):
    return spec_from_catalogue(catalogue(name), {**TENDER, "tender_id": f"T_{name}"})


def test_v4_round_trip_every_fixture():
    with tmpdir() as d:
        for name in CATS:
            r = v4_round_trip(_spec(name), d, expiring=catalogue(name))
            assert r["status"] == "PASS" and r["hash_ok"] and r["n_diffs"] == 0, name


def test_one_block_per_tpa_dynamic_networks_and_refer():
    ns = normalise(_spec("CAT_ALPHA"))
    blocks = ns["sheets"][0]["blocks"]
    assert [b["tpa"] for b in blocks] == ["TPA_RED", "TPA_BLUE", "TPA_GREEN"]   # never shared between TPAs
    assert all(len(b["plans"]) == 7 and len(b["rows"]) == 15 and len(b["columns"]) == 2 for b in blocks)
    ns = normalise(_spec("CAT_DELTA"))
    b = ns["sheets"][0]["blocks"][0]
    assert len(b["rows"]) == 12 and len(b["columns"]) == 1 and len(b["refer_cells"]) == 3


def test_generated_workbook_structure_and_determinism():
    with tmpdir() as d:
        g1 = generate_template(_spec("CAT_BETA"), d / "a.xlsx", expiring=catalogue("CAT_BETA"))
        g2 = generate_template(_spec("CAT_BETA"), d / "b.xlsx", expiring=catalogue("CAT_BETA"))
        h = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()  # noqa: E731
        assert h(d / "a.xlsx") == h(d / "b.xlsx") and g1.structure_hash == g2.structure_hash
        wb = load_workbook(d / "a.xlsx")
        assert wb.sheetnames[:4] == ["Cover", "Existing Rates", "Factor Tables", "Add-ons"]
        assert wb.sheetnames[-3:] == ["Terms", "_KEYS", "_META"]
        assert wb["_KEYS"].sheet_state == "hidden" and wb["_META"].sheet_state == "hidden"
        rs = wb[wb.sheetnames[4]]
        assert rs.protection.sheet
        c = rs["C7"]
        assert c.fill.fgColor.rgb.endswith("FFFF99") and c.protection.locked is False
        assert rs.data_validations.dataValidation
        meta = {r[0]: r[1] for r in wb["_META"].iter_rows(min_row=2, values_only=True)}
        assert meta["structure_hash"] == g1.structure_hash
        er = [v for row in wb["Existing Rates"].iter_rows(values_only=True) for v in row if isinstance(v, str)]
        assert sum("member_category" in v for v in er) == 5   # the 5 V2 flags listed on Existing Rates
        ref = wb["_KEYS"]["B2"].value
        assert re.match(r"^='[^']+'!\$C\$7$", ref)


def test_filled_template_is_byte_reproducible():
    with tmpdir() as d:
        generate_template(_spec("CAT_GAMMA"), d / "t.xlsx", expiring=catalogue("CAT_GAMMA"))
        fill_template(d / "t.xlsx", d / "f1.xlsx")
        fill_template(d / "t.xlsx", d / "f2.xlsx")
        assert (d / "f1.xlsx").read_bytes() == (d / "f2.xlsx").read_bytes()


def test_renamed_sheet_tolerated_through_keys():
    spec = _spec("CAT_GAMMA")
    with tmpdir() as d:
        generate_template(spec, d / "t.xlsx")
        written = fill_template(d / "t.xlsx", d / "f.xlsx")
        wb = load_workbook(d / "f.xlsx")
        old = spec["sheets"][0]["sheet_name"]
        wb[old].title = "Renamed by reinsurer"
        for row in wb["_KEYS"].iter_rows(min_row=2):     # what Excel does to formulas on a rename
            if row[1].value:
                row[1].value = row[1].value.replace(f"'{old}'", "'Renamed by reinsurer'")
        wb.save(d / "g.xlsx")
        p = parse_submission(d / "g.xlsx", template_spec=spec)
        assert p.hash_ok and all(p.values[k] == v for k, v in written.items())
        assert any("moved" in w for w in p.warnings)


def test_lost_keys_fall_back_to_positional_with_warning():
    spec = _spec("CAT_DELTA")
    with tmpdir() as d:
        generate_template(spec, d / "t.xlsx")
        written = fill_template(d / "t.xlsx", d / "f.xlsx")
        wb = load_workbook(d / "f.xlsx")
        del wb["_KEYS"]
        wb.save(d / "g.xlsx")
        p = parse_submission(d / "g.xlsx", template_spec=spec)
        assert p.positional and not p.hash_ok and p.warnings
        rates = {k: v for k, v in written.items() if not k.startswith("term=")}
        assert all(p.values[k] == v for k, v in rates.items())
        raises(SpecError, parse_submission, d / "g.xlsx")


def test_spec_validator():
    spec = _spec("CAT_ALPHA")
    assert validate_template_spec(spec) == []
    bad = copy.deepcopy(spec)
    bad["sheets"][0]["blocks"].append(copy.deepcopy(bad["sheets"][0]["blocks"][0]))
    assert any("TPA x plan" in e for e in validate_template_spec(bad))
    bad = copy.deepcopy(spec)
    bad["sheets"][0]["blocks"][0]["age_bands"] = "NOPE"
    assert any("unknown age band set" in e for e in validate_template_spec(bad))
    bad = copy.deepcopy(spec)
    bad["sheets"][0]["sheet_name"] = "_bad"
    bad["options"] = {"protect_structure": "yes"}
    errs = validate_template_spec(bad)
    assert any("sheet_name" in e for e in errs) and any("protect_structure" in e for e in errs)
    raises(SpecError, normalise, bad)


def test_split_options_marital_and_member_section():
    spec = {**TENDER, "tender_id": "T_SPLIT", "currency": "AED", "premium_basis": "NET",
            "age_band_sets": {"B": [{"label": "0-39", "range": [0, 39]}, {"label": "40+", "range": [40, 120]}]},
            "column_sets": {"MSF": [{"label": "M/SF", "keys": [{"gender": "M"}, {"gender": "F", "marital": "SINGLE"}]},
                                    {"label": "MF", "keys": [{"gender": "F", "marital": "MARRIED"}]}]},
            "sheets": [{"sheet_name": "S", "product_type": "INDIVIDUAL", "regulator": "DOH", "rating_region": "RGN_A",
                        "age_bands": "B", "column_set": "MF",
                        "split_options": {"marital": {"column_set": "MSF"},
                                          "member_section": {"sections": [
                                              {"label": "Principal", "key": {"member_category": "PRINCIPAL"}, "age_bands": "B"},
                                              {"label": "Dependant", "key": {"member_category": "DEPENDANT"}, "age_bands": "B"}]}},
                        "blocks": [{"tpa": "TPA_RED", "plans": [{"plan_id": "P1", "label": "P1"}]},
                                   {"tpa": "TPA_BLUE", "plans": [{"plan_id": "P1", "label": "P1"}],
                                    "splits": {"marital": True, "member_section": True}}]}]}
    ns = normalise(spec)
    off, on = ns["sheets"][0]["blocks"]
    assert len(off["rows"]) * len(off["columns"]) == 4 and len(on["rows"]) * len(on["columns"]) == 8
    with tmpdir() as d:
        assert v4_round_trip(spec, d)["status"] == "PASS"
        generate_template(spec, d / "t.xlsx")
        fill_template(d / "t.xlsx", d / "f.xlsx")
        cat = to_catalogue(parse_submission(d / "f.xlsx", template_spec=spec), spec, "SPLIT", "synthetic")
        blue = next(s for s in cat.scopes if s.tpa_set == ["TPA_BLUE"])
        rows = blue.components[0].rows
        assert len(rows) == 12      # M/SF expands to two canonical keys
        res, _ = evaluate(cat, [{"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_BLUE", "plan": "P1",
                                 "age": 30, "gender": "F", "marital": "SINGLE", "member_category": "SPOUSE"},
                                {"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_BLUE", "plan": "P1",
                                 "age": 30, "gender": "M", "marital": "MARRIED", "member_category": "SPOUSE"}])
        assert res.loc[0, "price"] == res.loc[1, "price"]     # both in the M/SF column, dependant section


def test_sub_x_sub_y_fixtures():
    alpha = catalogue("CAT_ALPHA")
    spec = json.loads((FIX / "T_ALPHA" / "template_spec.json").read_text(encoding="utf-8"))
    for name in ("SUB_X", "SUB_Y"):
        p = parse_submission(FIX / name / "submission.xlsx", template_spec=spec)
        assert p.hash_ok and not [w for w in p.warnings if "non-numeric" in w]
        cat = Catalogue.read(FIX / name / "catalogue.json")
        assert cat.kind == "SUBMISSION" and len(cat.scopes) == 3 and cat.terms["reinsurer"].startswith("Reinsurer")
        assert to_catalogue(p, spec, name, cat.owner, expiring=alpha).to_json() == cat.to_json()
        cell = {"product_type": "INDIVIDUAL", "rating_region": "RGN_A", "tpa": "TPA_GREEN", "plan": "P4", "age": 41,
                "gender": "M", "copay_op_diag": "20%", "nationality_group": "GRP_EAST"}
        e, _ = evaluate(alpha, [cell])
        s, _ = evaluate(cat, [cell])
        ratio = s.loc[0, "price"] / e.loc[0, "price"]
        assert 0.8 < ratio < 1.25 and s.loc[0, "factor_product"] == e.loc[0, "factor_product"]
    x = Catalogue.read(FIX / "SUB_X" / "catalogue.json").scopes[0].components[0].rows
    y = Catalogue.read(FIX / "SUB_Y" / "catalogue.json").scopes[0].components[0].rows
    assert pd.Series([a["value"] for a in x if a["value"]]).sum() > pd.Series([b["value"] for b in y if b["value"]]).sum()
