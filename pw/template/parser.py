"""Submission parser (SPEC §8.3): read a filled tender template back by key into a SUBMISSION catalogue.

Input cells are located through the live reference formulas in _KEYS (Excel rewrites them when sheets are renamed
or rows/columns inserted), falling back to the issued address, and - if _KEYS is lost - to positional parsing from
the template_spec with warnings. WS1 adds V5 (mandatory/outlier checks) and multiplier proposals handling.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .. import ProposalRequired, SpecError
from ..io_excel import make_reproducible
from ..model import Catalogue, Component, Plan, RatingScope
from . import KEYS_HEADER, KEYS_SHEET, META_SHEET
from .generator import compute_layout, hash_keys_rows
from .spec import normalise

_REF = re.compile(r"^=\s*(?:'((?:[^']|'')+)'|([^!']+))!\$?([A-Z]{1,3})\$?(\d+)\s*$")


def read_meta(path: Path) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if META_SHEET not in wb.sheetnames:
            return {}
        ws = wb[META_SHEET]
        out = {}
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or row[0] is None:
                continue
            out[str(row[0])] = "" if row[1] is None else str(row[1])
        return out
    finally:
        wb.close()


def _keys_table(wb_f) -> list[dict]:
    ws = wb_f[KEYS_SHEET]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            if list(r[:len(KEYS_HEADER)]) != KEYS_HEADER:
                raise SpecError(f"{KEYS_SHEET}: unexpected header {r}")
            continue
        if r[0] is None:
            continue
        rows.append(dict(zip(KEYS_HEADER, r)))
    return rows


def _locate(row: dict, sheetnames: list, warnings: list):
    m = _REF.match(str(row["ref"] or ""))
    if m:
        sheet = (m.group(1) or m.group(2)).replace("''", "'")
        addr = f"{m.group(3)}{m.group(4)}"
        if sheet in sheetnames:
            if sheet != row["issued_sheet"] or addr != row["issued_address"]:
                warnings.append(f"{row['cell_key']}: moved {row['issued_sheet']}!{row['issued_address']} -> {sheet}!{addr}")
            return sheet, addr
    warnings.append(f"{row['cell_key']}: reference formula unusable ({row['ref']!r}); using issued address")
    if row["issued_sheet"] not in sheetnames:
        return None, None
    return row["issued_sheet"], row["issued_address"]


@dataclass
class ParsedSubmission:
    path: Path
    values: dict                       # cell_key -> float | str | None
    kinds: dict                        # cell_key -> kind
    keys: dict                         # cell_key -> expansion (list of key dicts)
    meta: dict
    hash_ok: bool
    recomputed_hash: str | None
    positional: bool = False
    warnings: list = field(default_factory=list)

    def terms(self) -> dict:
        return {k.split("=", 1)[1]: v for k, v in self.values.items() if self.kinds.get(k) == "TERM"}


def parse_submission(path: Path, template_spec: dict | None = None) -> ParsedSubmission:
    path = Path(path)
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    warnings = []
    meta = read_meta(path)
    positional = False
    if KEYS_SHEET in wb_f.sheetnames:
        krows = _keys_table(wb_f)
    else:
        if template_spec is None:
            raise SpecError(f"{path.name}: {KEYS_SHEET} sheet missing and no template_spec given for positional parsing")
        warnings.append(f"{KEYS_SHEET} missing: positional parsing from the template_spec (rate cells only)")
        positional = True
        lay = compute_layout(normalise(template_spec))
        krows = [dict(zip(KEYS_HEADER, [r[0], None, r[1], r[2], r[3], r[4], r[5]])) for r in lay.keys_rows()]
    recomputed = hash_keys_rows([[r["cell_key"], r["kind"], bool(r["mandatory"]), r["key_json"], r["issued_sheet"],
                                  r["issued_address"]] for r in krows],
                                _rate_sheets_from_keys(krows), meta.get("tender_id"))
    hash_ok = bool(meta.get("structure_hash")) and meta.get("structure_hash") == recomputed and not positional
    if template_spec is not None and not positional:
        issued = compute_layout(normalise(template_spec))
        if issued.structure_hash() != recomputed and not any(r["kind"] in ("TERM", "MULTIPLIER", "ADDON") for r in krows):
            warnings.append("structure differs from the template_spec layout")
    values, kinds, keys = {}, {}, {}
    names = wb_v.sheetnames
    for r in krows:
        ck = r["cell_key"]
        kinds[ck] = r["kind"]
        keys[ck] = json.loads(r["key_json"])
        if r["kind"] == "REFER":
            continue
        if positional:
            sheet, addr = r["issued_sheet"], r["issued_address"]
            if sheet not in names:
                warnings.append(f"{ck}: sheet {sheet} missing")
                values[ck] = None
                continue
        else:
            sheet, addr = _locate(r, names, warnings)
            if sheet is None:
                values[ck] = None
                continue
        v = wb_v[sheet][addr].value
        if r["kind"] == "TERM":
            values[ck] = None if v is None else str(v)
        elif v is None or (isinstance(v, str) and not v.strip()):
            values[ck] = None
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            values[ck] = float(v)
        else:
            try:
                values[ck] = float(str(v).replace(",", ""))
                warnings.append(f"{ck}: text value '{v}' read as a number")
            except ValueError:
                values[ck] = str(v)
                warnings.append(f"{ck}: non-numeric value '{v}'")
    wb_f.close()
    wb_v.close()
    return ParsedSubmission(path=path, values=values, kinds=kinds, keys=keys, meta=meta, hash_ok=hash_ok,
                            recomputed_hash=recomputed, positional=positional, warnings=warnings)


def _rate_sheets_from_keys(krows) -> list:
    out = []
    for r in krows:
        if r["kind"] in ("RATE", "REFER") and r["issued_sheet"] not in out:
            out.append(r["issued_sheet"])
    return out


def synthetic_value(cell_key: str, seed: int = 4) -> float:
    h = int(hashlib.sha256(f"{seed}|{cell_key}".encode()).hexdigest()[:8], 16)
    return round(100.0 + (h % 900000) / 100.0, 2)


def fill_template(template_path: Path, out_path: Path, seed: int = 4, value_fn=None) -> dict:
    """Fill every input cell (synthetic or value_fn(cell_key, keys, kind) -> value) and save a copy.
    Returns {cell_key: value written}. Used for V4 and the SUB_X / SUB_Y fixtures."""
    wb = load_workbook(template_path, data_only=False)
    krows = _keys_table(wb)
    written = {}
    for r in krows:
        if r["kind"] == "REFER":
            continue
        ck = r["cell_key"]
        keys = json.loads(r["key_json"])
        if value_fn is not None:
            v = value_fn(ck, keys, r["kind"])
        elif r["kind"] == "TERM":
            v = f"synthetic {ck.split('=', 1)[1]}"
        else:
            v = synthetic_value(ck, seed)
        if v is None:
            continue
        m = _REF.match(str(r["ref"]))
        sheet = (m.group(1) or m.group(2)).replace("''", "'") if m else r["issued_sheet"]
        addr = f"{m.group(3)}{m.group(4)}" if m else r["issued_address"]
        wb[sheet][addr].value = v
        written[ck] = (float(v) if not isinstance(v, str) else v)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    wb.close()
    make_reproducible(out_path)
    return written


# ---------------------------------------------------------------------------------------------- to catalogue
_BLOCK_DIMS = ("product_type", "regulator", "rating_region", "tpa")

# TPA factor basis of submitted per-TPA block values (HANDOVER open item 4; owner decision 24 Oct 2026, provisional).
#   "PRE"  (option A): block value = base before the catalogue's TPA factor; the issued TPA factor still applies.
#   "POST" (option B): block value = the TPA's final base; TPA factor tables (family "tpa") are neutralised to 1.0.
# Switching to POST needs the owner's confirmation and a change request (SPEC §8.0 wording, Existing Rates, Cover).
TPA_FACTOR_BASIS = "PRE"
TPA_FAMILY = "tpa"


def _safe_id(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.|-]+", "_", s)


def to_catalogue(parsed: ParsedSubmission, template_spec: dict, catalogue_id: str, owner: str,
                 expiring: Catalogue | None = None, version: str = "1", exclude: set | None = None,
                 tpa_basis: str | None = None) -> Catalogue:
    """SUBMISSION catalogue: one scope per template block (product x region x TPA) with the submitted BASE_TABLE.
    With an expiring catalogue, the issued factor tables / modules / fees / loadings / formula of the matching
    expiring scope are attached (proposed multipliers replace issued ones when the template allowed them and
    values were supplied). The submitted per-TPA base is read on the same basis as the expiring base table."""
    ns = normalise(template_spec)
    allow_prop = ns["options"]["allow_proposed_multipliers"]
    exclude = set(exclude or ())
    tpa_basis = tpa_basis or TPA_FACTOR_BASIS
    if tpa_basis not in ("PRE", "POST"):
        raise SpecError(f"tpa_basis must be PRE or POST, got {tpa_basis!r}")
    cat = Catalogue(catalogue_id=catalogue_id, owner=owner, kind="SUBMISSION", version=version,
                    effective_date=ns.get("effective_date"), currency=ns["currency"], premium_basis=ns["premium_basis"],
                    source_file=parsed.path.name, spec_id=ns["tender_id"], terms=parsed.terms(),
                    meta={"structure_hash": parsed.meta.get("structure_hash"), "hash_ok": parsed.hash_ok,
                          "warnings": parsed.warnings[:200], "tpa_factor_basis": tpa_basis,
                          "excluded_cells": sorted(exclude)})
    by_block = {}
    lay = compute_layout(ns)
    for c in lay.cells:
        by_block.setdefault(c.block_id, []).append(c)
    proposed = {k: v for k, v in parsed.values.items() if parsed.kinds.get(k) == "MULTIPLIER" and v is not None}
    addons = {k: v for k, v in parsed.values.items() if parsed.kinds.get(k) == "ADDON" and v is not None}
    for sh in ns["sheets"]:
        for b in sh["blocks"]:
            scope = RatingScope(scope_id=_safe_id(b["block_id"]), product_type=sh["product_type"], regulator=sh.get("regulator"),
                                rating_region=sh.get("rating_region"), tpa_set=[b["tpa"]], currency=ns["currency"],
                                premium_basis=ns["premium_basis"])
            exp_scope = None
            if expiring is not None:
                from .generator import _match_scope
                exp_scope = _match_scope(expiring, sh, b["tpa"])
            for i, p in enumerate(b["plans"]):
                ep = None
                if exp_scope is not None:
                    try:
                        ep = exp_scope.plan(p["plan_id"])
                    except SpecError:
                        ep = None
                scope.plans.append(Plan(plan_id=p["plan_id"], label=p.get("label", p["plan_id"]),
                                        tier_rank=(ep.tier_rank if ep else i + 1),
                                        network_by_tpa={b["tpa"]: p.get("network") or (ep.network_by_tpa.get(b["tpa"]) if ep else None)},
                                        attributes=dict(ep.attributes) if ep else {}, defaults=dict(ep.defaults) if ep else {}))
            rows = []
            for c in by_block.get(b["block_id"], []):
                v = None if c.refer else parsed.values.get(c.cell_key)
                if not c.refer and (v is None or c.cell_key in exclude):
                    continue
                if not c.refer and not isinstance(v, float):
                    raise ProposalRequired(f"{c.cell_key}: non-numeric submitted value {v!r}")
                for k in c.keys:
                    if isinstance(k.get("age"), str):
                        raise ProposalRequired(f"{c.cell_key}: placeholder age band '{k['age']}' has no range; "
                                               "the owner supplies band boundaries (SPEC §8.4)")
                    rows.append({"key": {d: v2 for d, v2 in k.items() if d not in _BLOCK_DIMS},
                                 "value": None if c.refer else v, "status": "REFER" if c.refer else "VALUE",
                                 "labels": {}, "source_ref": f"{c.sheet}!{c.address}"})
            dims = sorted({d for r in rows for d in r["key"]})
            from ..registry import load_registry
            reg = load_registry()
            dims = reg.sort_dims(dims)
            for r in rows:
                r["key"] = {d: r["key"].get(d, "ANY") for d in dims}
            rows.sort(key=lambda r: json.dumps(r["key"], sort_keys=True))
            scope.components.append(Component(component_id="base", type="BASE_TABLE", keys=dims, rows=rows,
                                              currency=ns["currency"], premium_basis=ns["premium_basis"]))
            if exp_scope is not None:
                scope.formula = copy.deepcopy(exp_scope.formula)
                scope.defaults = dict(exp_scope.defaults)
                scope.rating_basis_age = exp_scope.rating_basis_age
                for comp in exp_scope.components:
                    if comp.type in ("BASE_TABLE", "EXPERIENCE_RATED"):
                        continue
                    c2 = copy.deepcopy(comp)
                    c2.source_blocks = {}
                    c2.block_check = {}
                    if comp.type == "FACTOR_TABLE" and allow_prop:
                        _apply(c2, proposed, "multiplier", exp_scope.scope_id)
                    if comp.type == "FACTOR_TABLE" and comp.family == TPA_FAMILY and tpa_basis == "POST":
                        for r in c2.rows:
                            r["value"] = 1.0
                            r["source_ref"] = "TPA factor neutralised (submitted block is post-TPA-factor)"
                    if comp.type == "MODULE":
                        _apply(c2, addons, "addon", exp_scope.scope_id)
                    scope.components.append(c2)
                scope.benefits = copy.deepcopy(exp_scope.benefits)
            cat.scopes.append(scope)
    if expiring is not None:
        cat.rules = copy.deepcopy(expiring.rules)
    cat.validate()
    return cat


def _apply(comp: Component, supplied: dict, prefix: str, scope_id: str):
    from ..model import canonical_key_str
    from ..registry import load_registry
    order = load_registry().order
    for r in comp.rows:
        ck = f"{prefix}|scope={scope_id}|component={comp.component_id}|" + canonical_key_str(r["key"], order)
        if ck in supplied and isinstance(supplied[ck], float):
            r["value"] = supplied[ck]
            r["source_ref"] = f"submitted {prefix}"
