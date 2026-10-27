"""Spec Builder workbook (SPEC §8.6; work order WS9, CR #26): author template_spec.json in Excel.

    generate_builder(spec | None, out)          -> GeneratedBuilder         (py tools\\pw.py t spec-builder ...)
    read_builder(xlsx, base_spec=None)          -> (spec | None, BuilderReadReport)   (py tools\\pw.py t spec-read ...)

JSON is the master (owner, 26 Oct 2026): the builder is an editing view of one template_spec.json. Every field of
SPEC §8.1 and the cedant sections of §8.5 is a structured row (no cell holds JSON). The read-back is gated like the
cedant read: any problem stops it with the sheet and cell, and no spec is returned. The result must then pass the
template spec validator and the cedant model (spec.py), which stay the only judges of a valid spec.

With a base spec (the spec the builder was made from) the read keeps what the builder does not carry: notes, source
flags, pending items, owner records, unknown keys, the age bands file, and the raw form of any table or block whose
meaning did not change. So spec -> builder -> spec gives back the same JSON.
"""
from __future__ import annotations

import copy
import datetime as _dt
import json
import math
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .. import ProposalRequired, SpecError
from ..io_excel import a1, new_workbook, quote_sheet
from ..registry import ANY, load_registry
from .spec import (DEFAULT_COLUMN_SETS, cedant_model, cedant_spec_hash, normalise, spec_hash,
                   validate_template_spec)

BUILDER_VERSION = "1"
MODE = "SPEC_BUILDER"
HEADER_ROW = 3                      # 0-based row of each table's header (title, help, blank above)
SPARE = 60                          # formatted spare rows below each long table
LIST_ROWS = 300                     # rows of each user list on the Lists sheet
YN = ["Y", "N"]
APPLIES_SHEET, APPLIES_SPLIT = "Sheet rows", "Member split"
ORDERS = {"BASE, FACTORS, MODULES": ["BASE", "FACTORS", "MODULES"], "BASE, MODULES, FACTORS": ["BASE", "MODULES", "FACTORS"]}
UNITS = ["PER_MEMBER", "PER_ELIGIBLE_FEMALE", "PER_FAMILY"]
MODES = ["ADD", "MULTIPLY"]
BASES = ["NET", "GROSS"]
ADDONS = ["READ_ONLY", "FILL"]
COL_DIMS = ("gender", "marital", "member_category")
REFER_DIMS = ("plan", "gender", "marital", "member_category", "age")

# ---------------------------------------------------------------------------------------------- field-value sheets
# (code, description, kind, list): kind str | int | bool | date; list = _LISTS list name for a dropdown
TENDER_FIELDS = [
    ("cedant", "Cedant short name", "str", None),
    ("tender_id", "Tender id (no spaces), e.g. CEDANT_MED_QS_2027", "str", None),
    ("currency", "Currency, e.g. AED", "str", None),
    ("premium_basis", "Premium basis of the base rates", "str", "basis"),
    ("deadline", "Quote deadline (YYYY-MM-DD)", "date", None),
    ("issue_date", "Issue date (YYYY-MM-DD); also the workbooks' document date", "date", None),
    ("effective_date", "Effective date (YYYY-MM-DD)", "date", None),
    ("contact", "Contact shown on the Cover", "str", None),
    ("template_spec_version", "Template spec version (leave as is)", "str", None),
    ("source_catalogue", "Catalogue the spec was authored from (optional)", "str", None),
]
REFERENCE_FIELDS = [
    ("existing_rates", "Show the Existing Rates sheet to reinsurers", "bool", "yn"),
    ("factor_tables", "Show the co-pay / factor tables (read only)", "bool", "yn"),
    ("addons", "Riders: READ_ONLY, or FILL (reinsurers quote their own rider rates)", "str", "addons"),
    ("tob_summary", "Show a table of benefits summary", "bool", "yn"),
    ("terms", "Include the Terms sheet", "bool", "yn"),
]
OPTION_FIELDS = [
    ("allow_proposed_multipliers", "Reinsurers may propose their own co-pay multipliers", "bool", "yn"),
    ("allow_benefit_comments", "Reinsurers may comment on benefits", "bool", "yn"),
    ("protect_structure", "Lock the template structure", "bool", "yn"),
    ("language", "Language code, e.g. en", "str", None),
    ("rounding", "Decimals for rates", "int", None),
    ("outlier_shading", "Shade quotes far from expiring", "bool", "yn"),
    ("outlier_band_low", "Outlier band, lower ratio (e.g. 0.7); blank = default", "num", None),
    ("outlier_band_high", "Outlier band, upper ratio (e.g. 1.3); blank = default", "num", None),
]
CEDANT_FIELDS = [
    ("catalogue_id", "Id of the expiring catalogue read back from the cedant workbook", "str", None),
    ("version", "Version of the cedant's catalogue", "str", None),
    ("effective_date", "Effective date of the cedant's catalogue (YYYY-MM-DD)", "date", None),
    ("spare_rows", "Spare rows below each long table of the cedant workbook", "int", None),
]
FV_SECTIONS = {"Tender": [("tender", TENDER_FIELDS), ("reference_sheets", REFERENCE_FIELDS), ("options", OPTION_FIELDS)],
               "Cedant workbook": [("cedant_workbook", CEDANT_FIELDS)]}
FV_HEADER = ["Section", "Field", "Description", "Value"]

# ---------------------------------------------------------------------------------------------- long tables
# column: (header, kind, list, mandatory, width); kind: in (text input) | num | int | ro (grey, informative)
# list: fixed list name (gated: the value must be in it) | "~name" (fixed suggestions, not gated) | "@name" (user list)
T = {
    "Lists": None,                  # special: one column per user list
    "Age bands": [("Band set", "in", None, True, 16), ("Label", "in", None, True, 14), ("Age from", "int", None, False, 10),
                  ("Age to", "int", None, False, 10), ("Source", "ro", None, False, 26)],
    "Columns": [("Column set", "in", None, True, 16), ("Column label", "in", None, True, 16), ("Gender", "in", "gender", False, 10),
                ("Marital", "in", "marital", False, 12), ("Member category", "in", "member_category", False, 26)],
    "Sheets": [("Sheet name", "in", None, True, 24), ("Title", "in", None, False, 50), ("Product type", "in", "product_type", True, 16),
               ("Regulator", "in", "regulator", True, 12), ("Rating region", "in", "@rating_region", False, 14),
               ("Mandatory", "in", "yn", False, 11), ("Age band set", "in", "@band_set", False, 14),
               ("Column set", "in", "@column_set", False, 14), ("Marital split columns", "in", "@column_set", False, 14),
               ("Refer text", "in", None, False, 28), ("Premium basis", "in", "basis", False, 12), ("Scope id", "in", None, False, 16),
               ("Formula order", "in", "order", False, 24), ("Formula rounding", "int", None, False, 10),
               ("Formula min premium", "num", None, False, 12), ("Formula max premium", "num", None, False, 12)],
    "Co-pay basis": [("Sheet", "in", "@sheet", True, 24), ("Family", "in", "~dim", True, 18), ("Option", "in", None, True, 28)],
    "Blocks": [("Sheet", "in", "@sheet", True, 24), ("Block id", "in", None, True, 24), ("TPA", "in", "@tpa", True, 14),
               ("Label", "in", None, False, 22), ("Marital split", "in", "yn", False, 9), ("Member split", "in", "yn", False, 9),
               ("Age band set", "in", "@band_set", False, 14), ("Column set", "in", "@column_set", False, 14),
               ("Refer text", "in", None, False, 24), ("Comment", "in", None, False, 50)],
    "Plans": [("Sheet", "in", "@sheet", True, 24), ("Block id", "in", "@block", True, 24), ("Plan id", "in", None, True, 9),
              ("Plan label", "in", None, True, 16), ("Network", "in", "@network", False, 16), ("Plan number", "int", None, False, 9)],
    "Sections": [("Sheet", "in", "@sheet", True, 24), ("Applies to", "in", "@applies", True, 24), ("Section label", "in", None, True, 26),
                 ("Member category", "in", "member_category", False, 26), ("Age band set", "in", "@band_set", False, 14),
                 ("Row label (no bands)", "in", None, False, 16)],
    "Custom rows": [("Sheet", "in", "@sheet", True, 24), ("Applies to", "in", "@applies", True, 24), ("Row label", "in", None, True, 12),
                    ("Section label", "in", None, False, 30), ("Member category", "in", "member_category", False, 26),
                    ("Age from", "int", None, False, 9), ("Age to", "int", None, False, 9)],
    "Refer cells": [("Sheet", "in", "@sheet", True, 24), ("Block id", "in", "@block", True, 24), ("Plan", "in", None, False, 8),
                    ("Gender", "in", "gender", False, 8), ("Marital", "in", "marital", False, 10),
                    ("Member category", "in", "member_category", False, 26), ("Age from", "int", None, False, 9),
                    ("Age to", "int", None, False, 9)],
    "Factor tables": [("Sheet", "in", "@sheet", True, 24), ("Component id", "in", None, True, 18), ("Family", "in", "~dim", True, 18),
                      ("Option dimension", "in", "dim", False, 16), ("Key on", "in", "dim", False, 14),
                      ("Key order", "in", None, False, 22), ("TPAs", "in", None, False, 22), ("Default option", "in", None, False, 16),
                      ("Label", "in", None, False, 20)],
    "Riders": [("Sheet", "in", "@sheet", True, 24), ("Component id", "in", None, True, 16), ("Family", "in", "~module", True, 18),
               ("Option dimension", "in", "dim", False, 16), ("Key on", "in", "dim", False, 12), ("Key order", "in", None, False, 22),
               ("TPAs", "in", None, False, 18), ("Unit", "in", "unit", True, 20), ("Mode", "in", "mode", False, 10),
               ("Eligibility", "in", None, False, 22), ("Premium basis", "in", "basis", False, 10), ("Mandatory", "in", "yn", False, 9),
               ("Default option", "in", None, False, 14), ("Label", "in", None, False, 18)],
    "Options": [("Sheet", "in", "@sheet", True, 24), ("Component id", "in", None, True, 18), ("Key", "in", None, False, 12),
                ("Option", "in", None, True, 30), ("Reference", "in", "yn", False, 10), ("Key label", "in", None, False, 18)],
    "Plan defaults": [("Sheet", "in", "@sheet", True, 24), ("Family", "in", "~dim", True, 20)],
    "Benefits": [("Sheet", "in", "@sheet", True, 24), ("Benefit code", "in", "benefit", True, 26), ("Plans", "in", None, False, 30)],
    "Cedant notes": [("Note", "in", None, True, 140)],
}
HELP = {
    "Lists": "Values offered in the dropdowns. Add a TPA, rating region, network, band set or column set here before "
             "using it elsewhere. MF and RATE are built-in column sets. This sheet is not read back.",
    "Age bands": "One row per band. Band set = the name used by sheets, blocks and sections. Age from / to are whole "
                 "years (open top band: to = 120). Both blank = a label-only row. Source 'file ...' = kept in the age "
                 "bands file: edit that file, not these rows.",
    "Columns": "Rate columns. One row per key; two rows with the same set and label form one column (e.g. M/SF = male, "
               "and single female). Blank dimension = any. MF (Male, Female) and RATE (one rate) are built in.",
    "Sheets": "One row per rate sheet (product x emirate). Give an age band set, or define the rows on the Sections or "
              "Custom rows sheets. Marital split columns: the column set used when a block switches the marital split "
              "on. Formula (cedant read-back): blank = default (BASE, FACTORS, MODULES; options rounding).",
    "Co-pay basis": "The basis the base rates are quoted at, per sheet: family and option (e.g. copay_op_diag = 0%). "
                    "Order is kept.",
    "Blocks": "One row per TPA block (product x emirate x TPA). A TPA has several blocks on one sheet only for "
              "different plans. Splits Y = the sheet's marital columns / member sections apply. Age band set / column "
              "set here override the sheet's.",
    "Plans": "One row per plan (network column group) of a block, in column order. Network = the TPA's network label.",
    "Sections": "Row sections. Applies to: 'Sheet rows' (every block of the sheet), 'Member split' (blocks with the "
                "member split on), or a block id (that block only). Each section has its own band set, or one row "
                "label when it has no bands.",
    "Custom rows": "Hand-built rows for irregular tables (e.g. 00-64 Employee & Dependants). Applies to 'Sheet rows' "
                   "or a block id. Two lines with the same row label (and section) form one row covering both keys.",
    "Refer cells": "Cells reinsurers do not quote (shown as Refer). Blank = any; a cell matching every filled field is "
                   "a refer cell (e.g. Member category PARENT = the whole parents row).",
    "Factor tables": "Cedant workbook factor tables (one row per table). Key on = the dimension the factors vary by "
                     "(blank = one set of factors). TPAs: comma list, blank = every TPA of the sheet. Key order: leave "
                     "blank (registry order). Options and the reference option are on the Options sheet.",
    "Riders": "Cedant workbook riders (add-on modules), one row per rider. Unit: PER_MEMBER / PER_ELIGIBLE_FEMALE / "
              "PER_FAMILY. Eligibility: e.g. maternity_eligible=Y (blank = the unit's default). Limits are on the "
              "Options sheet.",
    "Options": "Options of each factor table and rider, one row per key x option, in the order shown. Key blank = no "
               "key. Reference Y = the option with factor 1.0 (one per key; factor tables only). Key label: text shown "
               "beside the key (e.g. the source column a plan takes).",
    "Plan defaults": "Option families asked per plan in the cedant workbook (the option each plan is sold with).",
    "Benefits": "Benefit codes asked per plan in the cedant workbook. Plans: comma list, blank = every plan of the sheet.",
    "Cedant notes": "Notes shown on the cedant workbook's Cover, one per row.",
}
FV_HELP = {"Tender": "Tender fields, reference sheets and options. Blank = not set (the tool's default applies).",
           "Cedant workbook": "Cedant workbook (SPEC §8.5) settings. Notes are on the Cedant notes sheet."}
ORDER_SHEETS = ["Cover", "Tender", "Lists", "Age bands", "Columns", "Sheets", "Co-pay basis", "Blocks", "Plans", "Sections",
                "Custom rows", "Refer cells", "Factor tables", "Riders", "Options", "Plan defaults", "Benefits",
                "Cedant workbook", "Cedant notes"]
LISTS_SHEET, META_SHEET = "_LISTS", "_META"
USER_LISTS = [("TPAs", "tpa"), ("Rating regions", "rating_region"), ("Networks", "network"), ("Band sets", "band_set"),
              ("Column sets", "column_set")]

# keys of each spec object the builder carries; anything else is copied from the base spec (SPEC §8.6)
TOP_KEYS = {"template_spec_version", "cedant", "tender_id", "currency", "premium_basis", "deadline", "contact", "issue_date",
            "effective_date", "source_catalogue", "age_band_sets", "age_bands_file", "column_sets", "sheets",
            "reference_sheets", "options", "cedant_workbook"}
SHEET_KEYS = {"sheet_name", "title", "product_type", "regulator", "rating_region", "mandatory", "age_bands", "column_set",
              "copay_basis", "sections", "rows", "split_options", "refer_text", "premium_basis", "scope_id", "formula",
              "factor_tables", "riders", "plan_defaults", "benefits", "blocks"}
BLOCK_KEYS = {"block_id", "tpa", "label", "comment", "plans", "splits", "refer_cells", "refer_text", "age_bands", "column_set",
              "sections", "rows"}
PLAN_KEYS = {"plan_id", "label", "network", "plan_number"}
TABLE_KEYS = {"component_id", "family", "option_dim", "key_dim", "keys", "tpas", "rows", "options", "key_values", "reference",
              "default_option", "key_labels", "label"}
RIDER_KEYS = TABLE_KEYS - {"reference"} | {"unit", "mode", "eligibility", "premium_basis", "mandatory"}
CW_KEYS = {"catalogue_id", "version", "effective_date", "spare_rows", "notes"}
FORMULA_KEYS = {"order", "rounding", "min_premium", "max_premium"}
OPTIONS_KEYS = {c for c, *_ in OPTION_FIELDS if not c.startswith("outlier_band")} | {"outlier_band"}
REF_KEYS = {c for c, *_ in REFERENCE_FIELDS}


# ============================================================================================== helpers
def _s(v) -> str | None:
    """Cell value -> stripped text (None for blank). Whole floats lose '.0'; dates become YYYY-MM-DD."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "Y" if v else "N"
    if isinstance(v, _dt.datetime):
        return v.date().isoformat() if v.time() == _dt.time(0) else v.isoformat()
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return str(v)
        return str(int(v)) if v.is_integer() else repr(v)
    t = str(v).strip()
    return t or None


def _csv(v) -> list[str]:
    t = _s(v)
    return [x.strip() for x in t.split(",") if x.strip()] if t else []


def _yn(v) -> bool:
    return v is True or (isinstance(v, str) and v.upper() == "Y")


def _order_like(new: dict, base: dict | None) -> dict:
    """Keys in the base object's order first (minimal JSON diffs), then the rest in their own order."""
    if not base:
        return new
    out = {k: new[k] for k in base if k in new}
    out.update({k: v for k, v in new.items() if k not in out})
    return out


@dataclass
class GeneratedBuilder:
    path: Path
    counts: dict
    not_carried: list
    source_hash: str


@dataclass
class BuilderReadReport:
    problems: list = field(default_factory=list)      # [{sheet, cell, check, message, value}]
    warnings: list = field(default_factory=list)
    counts: dict = field(default_factory=dict)

    @property
    def status(self) -> str:
        return "PASS" if not self.problems else "FAIL"

    def add(self, sheet, cell, check, message, value=None):
        self.problems.append({"sheet": sheet, "cell": cell or "", "check": check, "message": message,
                              "value": "" if value is None else str(value)})

    def warn(self, msg: str):
        if msg not in self.warnings:
            self.warnings.append(msg)

    def to_frame(self):
        import pandas as pd
        rows = [{"level": "PROBLEM", **p} for p in self.problems] + \
               [{"level": "WARNING", "sheet": "", "cell": "", "check": "warning", "message": w, "value": ""} for w in self.warnings]
        return pd.DataFrame(rows, columns=["level", "sheet", "cell", "check", "message", "value"])

    def summary(self) -> str:
        return (f"SPEC READ {self.status}: {len(self.problems)} problem(s), {len(self.warnings)} warning(s); "
                + ", ".join(f"{k} {v}" for k, v in sorted(self.counts.items())))


def _fixed_lists() -> dict[str, list[str]]:
    reg = load_registry()

    def vals(code):
        d = reg.dim(code)
        vals = list(d.values or ()) if isinstance(d.values, (list, tuple)) else []
        return vals + [x for x in (d.parents or ()) if x not in vals]
    dims = sorted(reg.dims)
    return {"yn": YN, "basis": BASES, "addons": ADDONS, "order": list(ORDERS), "unit": UNITS, "mode": MODES,
            "product_type": vals("product_type"), "regulator": vals("regulator"), "gender": vals("gender"),
            "marital": vals("marital"), "member_category": vals("member_category"), "dim": dims,
            "module": [d for d in dims if d.startswith("module_")], "benefit": sorted(reg.benefits)}


def _file_band_sets(spec: dict) -> dict:
    f = spec.get("age_bands_file")
    if not f:
        return {}
    base = Path(spec.get("_base_dir") or ".")
    doc = json.loads((base / f).read_text(encoding="utf-8"))
    return dict(doc.get("age_band_sets") or {})


# ============================================================================================== spec -> rows
def _key_cols(key: dict, dims, where: str) -> dict:
    bad = [d for d in key if d not in dims]
    if bad:
        raise ProposalRequired(f"{where}: dimension(s) {bad} have no column in the Spec Builder (edit the JSON for this)")
    out = {}
    for d, v in key.items():
        if d == "age":
            if not (isinstance(v, list) and len(v) == 2):
                raise ProposalRequired(f"{where}: age must be [from, to] for the Spec Builder")
            out["Age from"], out["Age to"] = v
        else:
            out[{"gender": "Gender", "marital": "Marital", "member_category": "Member category", "plan": "Plan"}[d]] = v
    return out


def _not_carried(spec: dict) -> list[str]:
    out = [f"top-level '{k}'" for k in spec if k not in TOP_KEYS and k != "_base_dir"]
    for k in (spec.get("options") or {}):
        if k not in OPTIONS_KEYS:
            out.append(f"options.{k}")
    for sh in spec.get("sheets") or []:
        n = sh.get("sheet_name")
        out += [f"sheet '{n}': '{k}'" for k in sh if k not in SHEET_KEYS]
        out += [f"sheet '{n}': formula.{k}" for k in (sh.get("formula") or {}) if k not in FORMULA_KEYS]
        for b in sh.get("blocks") or []:
            out += [f"block '{b.get('block_id')}': '{k}'" for k in b if k not in BLOCK_KEYS]
            for p in b.get("plans") or []:
                out += [f"block '{b.get('block_id')}' plan '{p.get('plan_id')}': '{k}'" for k in p if k not in PLAN_KEYS]
        for kind, keys in (("factor_tables", TABLE_KEYS), ("riders", RIDER_KEYS)):
            for t in sh.get(kind) or []:
                out += [f"sheet '{n}' {kind} '{t.get('component_id')}': '{k}'" for k in t if k not in keys]
    out += [f"cedant_workbook.{k}" for k in (spec.get("cedant_workbook") or {}) if k not in CW_KEYS]
    return out


def spec_rows(spec: dict | None) -> dict:
    """The builder's rows for a spec (or empty tables for None): {sheet: [ {header: value} ]} plus field values."""
    rows = {n: [] for n in T if T[n] is not None}
    fv = {"tender": {}, "reference_sheets": {}, "options": {}, "cedant_workbook": {}}
    lists = {code: [] for _, code in USER_LISTS}
    lists["column_set"] += list(DEFAULT_COLUMN_SETS)
    if spec is None:
        return {"rows": rows, "fv": fv, "lists": lists}
    for c, *_ in TENDER_FIELDS:
        if spec.get(c) is not None:
            fv["tender"][c] = spec[c]
    for c, *_ in REFERENCE_FIELDS:
        if c in (spec.get("reference_sheets") or {}):
            fv["reference_sheets"][c] = spec["reference_sheets"][c]
    opts = spec.get("options") or {}
    for c, *_ in OPTION_FIELDS:
        if c in opts:
            fv["options"][c] = opts[c]
    if isinstance(opts.get("outlier_band"), list) and len(opts["outlier_band"]) == 2:
        fv["options"]["outlier_band_low"], fv["options"]["outlier_band_high"] = opts["outlier_band"]
    cw = spec.get("cedant_workbook") or {}
    for c, *_ in CEDANT_FIELDS:
        if cw.get(c) is not None:
            fv["cedant_workbook"][c] = cw[c]
    rows["Cedant notes"] = [{"Note": n} for n in cw.get("notes") or []]
    # age bands (spec first, then the file's sets, marked)
    fsets = _file_band_sets(spec)
    for src, sets in (("", spec.get("age_band_sets") or {}), (f"file {spec.get('age_bands_file')}", fsets)):
        for name, bands in sets.items():
            lists["band_set"].append(name)
            for b in bands:
                if set(b) - {"label", "range"}:
                    raise ProposalRequired(f"age band set '{name}': band fields {sorted(set(b) - {'label', 'range'})} have no "
                                           "column in the Spec Builder")
                r = {"Band set": name, "Label": b.get("label"), "Source": src or None}
                if b.get("range"):
                    r["Age from"], r["Age to"] = b["range"]
                rows["Age bands"].append(r)
    for name, cols in (spec.get("column_sets") or {}).items():
        if name not in lists["column_set"]:
            lists["column_set"].append(name)
        for c in cols:
            for k in c["keys"]:
                rows["Columns"].append({"Column set": name, "Column label": c["label"],
                                        **_key_cols(k, COL_DIMS, f"column set '{name}' column '{c['label']}'")})
    for sh in spec.get("sheets") or []:
        n = sh["sheet_name"]
        for k in ("age_bands", "column_set"):
            if k in sh and not isinstance(sh[k], str):
                raise ProposalRequired(f"sheet '{n}': inline {k} have no Spec Builder form; name them in the spec first")
        f = sh.get("formula") or {}
        order = f.get("order")
        olab = next((lab for lab, o in ORDERS.items() if o == order), None) if order is not None else None
        if order is not None and olab is None:
            raise ProposalRequired(f"sheet '{n}': formula order {order} has no Spec Builder option")
        so = sh.get("split_options") or {}
        if set(so) - {"marital", "member_section"}:
            raise ProposalRequired(f"sheet '{n}': split_options {sorted(set(so) - {'marital', 'member_section'})} unknown")
        if sh.get("rating_region") and sh["rating_region"] not in lists["rating_region"]:
            lists["rating_region"].append(sh["rating_region"])
        rows["Sheets"].append({"Sheet name": n, "Title": sh.get("title"), "Product type": sh.get("product_type"),
                               "Regulator": sh.get("regulator"), "Rating region": sh.get("rating_region"),
                               "Mandatory": None if "mandatory" not in sh else ("Y" if sh["mandatory"] else "N"),
                               "Age band set": sh.get("age_bands"), "Column set": sh.get("column_set"),
                               "Marital split columns": (so.get("marital") or {}).get("column_set"),
                               "Refer text": sh.get("refer_text"), "Premium basis": sh.get("premium_basis"),
                               "Scope id": sh.get("scope_id"), "Formula order": olab, "Formula rounding": f.get("rounding"),
                               "Formula min premium": f.get("min_premium"), "Formula max premium": f.get("max_premium")})
        for fam, opt in (sh.get("copay_basis") or {}).items():
            rows["Co-pay basis"].append({"Sheet": n, "Family": fam, "Option": opt})

        def sections(secs, applies, where):
            for sec in secs:
                bad = set(sec) - {"label", "key", "age_bands", "row_label"}
                if bad:
                    raise ProposalRequired(f"{where}: section fields {sorted(bad)} have no column in the Spec Builder")
                key = sec.get("key") or {}
                rows["Sections"].append({"Sheet": n, "Applies to": applies, "Section label": sec.get("label"),
                                         **_key_cols(key, ("member_category",), where), "Age band set": sec.get("age_bands"),
                                         "Row label (no bands)": sec.get("row_label")})

        def custom(rws, applies, where):
            for r in rws:
                bad = set(r) - {"label", "section", "key", "keys"}
                if bad:
                    raise ProposalRequired(f"{where}: row fields {sorted(bad)} have no column in the Spec Builder")
                for k in (r["keys"] if r.get("keys") else [r.get("key") or {}]):
                    rows["Custom rows"].append({"Sheet": n, "Applies to": applies, "Row label": r.get("label"),
                                                "Section label": r.get("section"),
                                                **_key_cols(k, ("member_category", "age"), where)})
        sections(sh.get("sections") or [], APPLIES_SHEET, f"sheet '{n}' sections")
        sections((so.get("member_section") or {}).get("sections") or [], APPLIES_SPLIT, f"sheet '{n}' member split")
        custom(sh.get("rows") or [], APPLIES_SHEET, f"sheet '{n}' rows")
        for b in sh.get("blocks") or []:
            bid = b.get("block_id")
            if not bid:
                raise ProposalRequired(f"sheet '{n}': a block without block_id (the Spec Builder needs one)")
            for k in ("age_bands", "column_set"):
                if k in b and not isinstance(b[k], str):
                    raise ProposalRequired(f"block '{bid}': inline {k} have no Spec Builder form")
            sp = b.get("splits") or {}
            if b.get("tpa") and b["tpa"] not in lists["tpa"]:
                lists["tpa"].append(b["tpa"])
            rows["Blocks"].append({"Sheet": n, "Block id": bid, "TPA": b.get("tpa"), "Label": b.get("label"),
                                   "Marital split": None if "marital" not in sp else ("Y" if sp["marital"] else "N"),
                                   "Member split": None if "member_section" not in sp else ("Y" if sp["member_section"] else "N"),
                                   "Age band set": b.get("age_bands"), "Column set": b.get("column_set"),
                                   "Refer text": b.get("refer_text"), "Comment": b.get("comment")})
            for p in b.get("plans") or []:
                if p.get("network") and p["network"] not in lists["network"]:
                    lists["network"].append(p["network"])
                rows["Plans"].append({"Sheet": n, "Block id": bid, "Plan id": p.get("plan_id"), "Plan label": p.get("label"),
                                      "Network": p.get("network"), "Plan number": p.get("plan_number")})
            sections(b.get("sections") or [], bid, f"block '{bid}' sections")
            custom(b.get("rows") or [], bid, f"block '{bid}' rows")
            for rc in b.get("refer_cells") or []:
                rows["Refer cells"].append({"Sheet": n, "Block id": bid, **_key_cols(rc, REFER_DIMS, f"block '{bid}' refer cell")})
        model, _ = cedant_model({"sheets": [sh]})
        msh = model["sheets"][0] if model["sheets"] else {"factor_tables": [], "riders": []}
        mt = {m["component_id"]: m for m in msh["factor_tables"] + msh["riders"]}
        for kind in ("factor_tables", "riders"):
            for t in sh.get(kind) or []:
                cid = t.get("component_id")
                if not cid:
                    raise ProposalRequired(f"sheet '{n}': a {kind} entry without component_id (the Spec Builder needs one)")
                m = mt.get(cid)
                if m is None:
                    raise ProposalRequired(f"sheet '{n}' {kind} '{cid}': not a valid table (fix the spec first)")
                common = {"Sheet": n, "Component id": cid, "Family": t.get("family"), "Option dimension": t.get("option_dim"),
                          "Key on": t.get("key_dim"), "Key order": ", ".join(t["keys"]) if t.get("keys") else None,
                          "TPAs": ", ".join(t["tpas"]) if t.get("tpas") else None, "Default option": t.get("default_option"),
                          "Label": t.get("label")}
                if kind == "factor_tables":
                    rows["Factor tables"].append(common)
                else:
                    el = t.get("eligibility")
                    rows["Riders"].append({**common, "Unit": t.get("unit"), "Mode": t.get("mode"),
                                           "Eligibility": None if el is None else ", ".join(f"{d}={v}" for d, v in el.items()),
                                           "Premium basis": t.get("premium_basis"),
                                           "Mandatory": None if "mandatory" not in t else ("Y" if t["mandatory"] else "N")})
                labelled = set()
                for r in m["rows"]:
                    lab = m["key_labels"].get(r["key"]) if r["key"] not in labelled else None
                    labelled.add(r["key"])
                    rows["Options"].append({"Sheet": n, "Component id": cid, "Key": None if r["key"] == ANY else r["key"],
                                            "Option": r["option"], "Reference": "Y" if r.get("reference") else None,
                                            "Key label": lab})
        for fam in sh.get("plan_defaults") or []:
            rows["Plan defaults"].append({"Sheet": n, "Family": fam})
        for bd in sh.get("benefits") or []:
            code = bd if isinstance(bd, str) else bd.get("code")
            plans = None if isinstance(bd, str) else bd.get("plans")
            rows["Benefits"].append({"Sheet": n, "Benefit code": code, "Plans": None if plans is None else ", ".join(plans)})
    return {"rows": rows, "fv": fv, "lists": lists}


# ============================================================================================== generate
def generate_builder(spec: dict | None, out_path: Path) -> GeneratedBuilder:
    """Write the Spec Builder workbook: blank (spec None) or pre-filled from a spec (structure only)."""
    if spec is not None:
        errs = validate_template_spec(spec)
        if errs:
            raise SpecError("template_spec invalid (fix it before making a builder):\n  " + "\n  ".join(errs))
    data = spec_rows(spec)
    nc = _not_carried(spec) if spec is not None else []
    src_hash = spec_hash(normalise(spec)) if spec is not None else ""
    src_chash = cedant_spec_hash(spec) if spec is not None else ""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = new_workbook(out_path, title="Spec Builder" + (f" {spec['tender_id']}" if spec else ""))
    F = {"title": wb.add_format({"bold": True, "font_size": 14}),
         "help": wb.add_format({"italic": True, "font_color": "#555555", "text_wrap": True, "valign": "top"}),
         "hdr": wb.add_format({"bold": True, "border": 1, "bg_color": "#DDEBF7", "text_wrap": True, "valign": "top"}),
         "in": wb.add_format({"bg_color": "#FFFF99", "border": 1}),
         "num": wb.add_format({"bg_color": "#FFFF99", "border": 1, "num_format": "General"}),
         "ro": wb.add_format({"bg_color": "#F2F2F2", "border": 1, "font_color": "#555555"}),
         "bold": wb.add_format({"bold": True}), "wrap": wb.add_format({"text_wrap": True, "valign": "top"})}
    fixed = _fixed_lists()
    ws = {n: wb.add_worksheet(n) for n in ORDER_SHEETS}
    for w in ws.values():                       # printable: landscape, one page wide
        w.set_landscape()
        w.set_paper(9)
        w.fit_to_pages(1, 0)
    ws_lists = wb.add_worksheet(LISTS_SHEET)
    ws_meta = wb.add_worksheet(META_SHEET)
    # hidden fixed lists
    refs = {}
    for c, (name, vals) in enumerate(sorted(fixed.items())):
        ws_lists.write_string(0, c, name)
        for i, v in enumerate(vals):
            ws_lists.write_string(1 + i, c, v)
        refs[name] = f"={quote_sheet(LISTS_SHEET)}!{a1(2, c + 1, True)}:{a1(max(len(vals), 1) + 1, c + 1, True)}"
    ws_lists.hide()
    # user lists (visible)
    wl = ws["Lists"]
    wl.write(0, 0, "Lists", F["title"])
    wl.merge_range(1, 0, 1, len(USER_LISTS) * 2 - 2, HELP["Lists"], F["help"])
    wl.set_row(1, 45)
    for j, (lab, code) in enumerate(USER_LISTS):
        c = j * 2
        wl.set_column(c, c, 20)
        wl.write(HEADER_ROW, c, lab, F["hdr"])
        vals = data["lists"][code]
        for i in range(LIST_ROWS):
            v = vals[i] if i < len(vals) else None
            if v is None:
                wl.write_blank(HEADER_ROW + 1 + i, c, None, F["in"])
            else:
                wl.write_string(HEADER_ROW + 1 + i, c, v, F["in"])
        refs["@" + code] = (f"={quote_sheet('Lists')}!{a1(HEADER_ROW + 2, c + 1, True)}:"
                            f"{a1(HEADER_ROW + 1 + LIST_ROWS, c + 1, True)}")
    # cross references to the defining tables
    n_sheets = len(data["rows"]["Sheets"]) + SPARE
    n_blocks = len(data["rows"]["Blocks"]) + SPARE
    refs["@sheet"] = f"={quote_sheet('Sheets')}!$A${HEADER_ROW + 2}:$A${HEADER_ROW + 1 + n_sheets}"
    refs["@block"] = f"={quote_sheet('Blocks')}!$B${HEADER_ROW + 2}:$B${HEADER_ROW + 1 + n_blocks}"
    applies = [APPLIES_SHEET, APPLIES_SPLIT]
    ws_lists.write_string(0, len(fixed), "applies")
    for i, v in enumerate(applies):
        ws_lists.write_string(1 + i, len(fixed), v)
    # 'Applies to' accepts a block id too: warning-style validation offering the two fixed choices
    refs["@applies"] = f"={quote_sheet(LISTS_SHEET)}!{a1(2, len(fixed) + 1, True)}:{a1(3, len(fixed) + 1, True)}"
    counts = {}
    # field / value sheets
    for sname, sections_ in FV_SECTIONS.items():
        w = ws[sname]
        w.write(0, 0, sname, F["title"])
        w.merge_range(1, 0, 1, 3, FV_HELP[sname], F["help"])
        w.set_row(1, 30)
        for j, (h, wd) in enumerate(zip(FV_HEADER, (18, 28, 70, 40))):
            w.write(HEADER_ROW, j, h, F["hdr"])
            w.set_column(j, j, wd)
        r = HEADER_ROW + 1
        for sec, fields in sections_:
            for code, desc, kind, lst in fields:
                w.write_string(r, 0, sec, F["ro"])
                w.write_string(r, 1, code, F["ro"])
                w.write_string(r, 2, desc, F["ro"])
                v = data["fv"][sec].get(code)
                if isinstance(v, bool):
                    v = "Y" if v else "N"
                if v is None:
                    w.write_blank(r, 3, None, F["in"])
                elif isinstance(v, (int, float)):
                    w.write_number(r, 3, v, F["num"])
                else:
                    w.write_string(r, 3, str(v), F["in"])
                if lst:
                    w.data_validation(r, 3, r, 3, {"validate": "list", "source": refs[lst], "error_type": "stop"})
                r += 1
        w.freeze_panes(HEADER_ROW + 1, 0)
    # long tables
    for sname, cols in T.items():
        if cols is None:
            continue
        w = ws[sname]
        rws = data["rows"][sname]
        counts[sname] = len(rws)
        w.write(0, 0, sname, F["title"])
        w.merge_range(1, 0, 1, max(len(cols) - 1, 3), HELP[sname], F["help"])
        w.set_row(1, 60)
        for j, (h, kind, lst, mand, wd) in enumerate(cols):
            w.write(HEADER_ROW, j, h + (" *" if mand else ""), F["hdr"])
            w.set_column(j, j, wd)
        last = HEADER_ROW + len(rws) + SPARE
        for i in range(len(rws) + SPARE):
            r = HEADER_ROW + 1 + i
            row = rws[i] if i < len(rws) else {}
            for j, (h, kind, lst, mand, wd) in enumerate(cols):
                fmt = F["ro"] if kind == "ro" else (F["num"] if kind in ("int", "num") else F["in"])
                v = row.get(h)
                if v is None:
                    w.write_blank(r, j, None, fmt)
                elif isinstance(v, bool):
                    w.write_string(r, j, "Y" if v else "N", fmt)
                elif isinstance(v, (int, float)):
                    w.write_number(r, j, v, fmt)
                else:
                    w.write_string(r, j, str(v), fmt)
        for j, (h, kind, lst, mand, wd) in enumerate(cols):
            if lst:
                warn = lst[0] in "@~"
                w.data_validation(HEADER_ROW + 1, j, last, j,
                                  {"validate": "list", "source": refs[lst.lstrip("~")], "error_type": "warning" if warn else "stop",
                                   "error_title": h, "error_message": ("Not in the list: add it on the Lists sheet, or "
                                                                       "check the spelling." if warn else "Pick a value from the list.")})
            elif kind in ("int", "num"):
                w.data_validation(HEADER_ROW + 1, j, last, j,
                                  {"validate": "integer" if kind == "int" else "decimal", "criteria": ">=", "value": 0,
                                   "error_type": "stop", "error_message": "A number >= 0."})
        w.freeze_panes(HEADER_ROW + 1, 0)
        w.autofilter(HEADER_ROW, 0, last, len(cols) - 1)
    # cover
    c = ws["Cover"]
    c.set_column(0, 0, 24)
    c.set_column(1, 1, 120)
    c.write(0, 0, "Spec Builder - template_spec.json in Excel" + (f" - {spec['tender_id']}" if spec else " (blank)"), F["title"])
    lines = [("What this is", "An editing view of the tender's template_spec.json (SPEC §8.6). The JSON is the master: "
                              "the tool reads this workbook back into it, and every other step uses the JSON."),
             ("How to use", "1. Fill or change the yellow cells (dropdowns where a list applies; * = mandatory). "
                            "2. Save the workbook as inputs\\spec_builder.xlsx in the tender folder. "
                            "3. Run step 0b (t spec-read). SPEC READ PASS writes inputs\\template_spec.json (the previous "
                            "one is kept as template_spec.previous.json). SPEC READ FAIL lists each problem with its sheet "
                            "and cell: fix them and run 0b again."),
             ("Check the result", "Run steps 1 and 3: the cedant workbook and the tender template are made from the new "
                                  "spec. Look at them before issuing."),
             ("Built from", f"{spec['tender_id']} (spec hash {src_hash[:16]})" if spec else "blank"),
             ("Builder version", f"{BUILDER_VERSION} / {MODE}"),
             ("Kept from the spec, not shown here", "; ".join(nc) if nc else "nothing")]
    for i, (k, v) in enumerate(lines):
        c.write(2 + i, 0, k, F["bold"])
        c.write(2 + i, 1, v, F["wrap"])
    r = 3 + len(lines)
    c.write(r, 0, "Sheet directory", F["bold"])
    for i, n in enumerate(ORDER_SHEETS[1:]):
        c.write(r + 1 + i, 0, n)
        c.write(r + 1 + i, 1, (HELP.get(n) or FV_HELP.get(n) or "").split(". ")[0] + "."
                + (f" ({counts[n]} row(s))" if n in counts else ""))
    meta = [("mode", MODE), ("builder_version", BUILDER_VERSION), ("tender_id", spec["tender_id"] if spec else ""),
            ("source_spec_hash", src_hash), ("source_cedant_spec_hash", src_chash)]
    for i, (k, v) in enumerate(meta):
        ws_meta.write_string(i, 0, k)
        ws_meta.write_string(i, 1, v)
    ws_meta.hide()
    ws["Cover"].activate()
    wb.close()
    return GeneratedBuilder(path=out_path, counts=counts, not_carried=nc, source_hash=src_hash)


# ============================================================================================== read
class _Table:
    """A long table read by header scan: rows with values by header and the Excel row number."""

    def __init__(self, wb, name: str, rep: BuilderReadReport):
        self.name, self.rows, self.cols = name, [], {}
        cols = T[name]
        want = [h for h, *_ in cols]
        if name not in wb.sheetnames:
            rep.add(name, "", "sheet_missing", f"sheet '{name}' is missing (do not delete or rename builder sheets)")
            return
        ws = wb[name]
        hdr_row = None
        for r in range(1, 16):
            vals = [_s(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            clean = [v.rstrip(" *") if v else v for v in vals]
            if all(h in clean for h in want if dict((x[0], x[3]) for x in cols)[h]):
                hdr_row = r
                self.cols = {h: clean.index(h) + 1 for h in want if h in clean}
                break
        if hdr_row is None:
            rep.add(name, "", "header_missing", f"header row not found on '{name}' (expected columns {want})")
            return
        for r in range(hdr_row + 1, ws.max_row + 1):
            vals = {h: ws.cell(r, c).value for h, c in self.cols.items()}
            if all(_s(v) is None for v in vals.values()):
                continue
            self.rows.append((r, vals))
        self.spec = {x[0]: x for x in cols}

    def cell(self, r: int, h: str) -> str:
        return a1(r, self.cols[h]) if h in self.cols else ""


class _Reader:
    def __init__(self, wb, rep: BuilderReadReport):
        self.wb, self.rep = wb, rep
        self.fixed = _fixed_lists()

    # -------------------------------------------------------------------------- cell parsing
    def val(self, t: _Table, r: int, row: dict, h: str):
        """Typed value of one cell, with the mandatory / list / number gates. Returns None for blank or bad."""
        if h not in t.cols:
            return None
        _, kind, lst, mand, _ = t.spec[h]
        raw = row.get(h)
        s = _s(raw)
        if s is None:
            if mand:
                self.rep.add(t.name, t.cell(r, h), "blank_mandatory", f"'{h}' is mandatory")
            return None
        if kind in ("int", "num"):
            try:
                x = float(s.replace(",", ""))
            except ValueError:
                self.rep.add(t.name, t.cell(r, h), "not_a_number", f"'{h}' must be a number", raw)
                return None
            if kind == "int":
                if not x.is_integer():
                    self.rep.add(t.name, t.cell(r, h), "not_whole", f"'{h}' must be a whole number", raw)
                    return None
                return int(x)
            return int(x) if x.is_integer() else x
        if lst and lst[0] not in "@~" and s not in self.fixed[lst]:
            self.rep.add(t.name, t.cell(r, h), "not_in_list", f"'{h}' value '{s}' is not in the dropdown list", raw)
            return None
        return s

    def fv(self, sname: str) -> dict:
        """Field / value sheet -> {section: {code: typed value}} (blank = not set)."""
        out = {}
        if sname not in self.wb.sheetnames:
            self.rep.add(sname, "", "sheet_missing", f"sheet '{sname}' is missing")
            return out
        ws = self.wb[sname]
        known = {(sec, code): (kind, lst) for sec, fields in FV_SECTIONS[sname] for code, _, kind, lst in fields}
        seen = set()
        for r in range(1, ws.max_row + 1):
            sec, code = _s(ws.cell(r, 1).value), _s(ws.cell(r, 2).value)
            if (sec, code) not in known:
                continue
            seen.add((sec, code))
            kind, lst = known[(sec, code)]
            raw = ws.cell(r, 4).value
            s = _s(raw)
            if s is None:
                continue
            ref = a1(r, 4)
            if kind == "bool":
                if s.upper() not in ("Y", "N"):
                    self.rep.add(sname, ref, "not_in_list", f"{sec}.{code} must be Y or N", raw)
                    continue
                v = s.upper() == "Y"
            elif kind in ("int", "num"):
                try:
                    x = float(s)
                except ValueError:
                    self.rep.add(sname, ref, "not_a_number", f"{sec}.{code} must be a number", raw)
                    continue
                if kind == "int" and not x.is_integer():
                    self.rep.add(sname, ref, "not_whole", f"{sec}.{code} must be a whole number", raw)
                    continue
                v = int(x) if x.is_integer() else x
            elif kind == "date":
                try:
                    _dt.date.fromisoformat(s[:10])
                except ValueError:
                    self.rep.add(sname, ref, "not_a_date", f"{sec}.{code} must be a date YYYY-MM-DD", raw)
                    continue
                v = s[:10] if len(s) >= 10 and s[10:] in ("", "T00:00:00") else s
            else:
                if lst and s not in self.fixed[lst]:
                    self.rep.add(sname, ref, "not_in_list", f"{sec}.{code} value '{s}' is not in the list {self.fixed[lst]}", raw)
                    continue
                v = s
            out.setdefault(sec, {})[code] = v
        missing = [f"{s}.{c}" for s, c in known if (s, c) not in seen]
        if missing:
            self.rep.add(sname, "", "field_missing", f"field rows missing (do not delete rows on '{sname}'): {missing}")
        return out


def read_builder(xlsx: Path, base_spec: dict | None = None, previous_spec: dict | None = None
                 ) -> tuple[dict | None, BuilderReadReport]:
    """Read a filled Spec Builder back into a template spec. Returns (spec, report); spec is None on any problem.
    base_spec: the current template_spec.json (with _base_dir): fields the builder does not carry are kept from it.
    previous_spec: the spec an earlier read of this builder replaced (template_spec.previous.json): when it is the
    builder's source, the builder is being read again after a fix, which is allowed."""
    rep = BuilderReadReport()
    xlsx = Path(xlsx)
    try:
        wb = load_workbook(xlsx, data_only=True)
    except Exception as e:  # noqa: BLE001
        rep.add("", "", "unreadable", f"cannot open {xlsx.name}: {e}")
        return None, rep
    meta = {}
    if META_SHEET in wb.sheetnames:
        for row in wb[META_SHEET].iter_rows(values_only=True):
            if row and row[0]:
                meta[str(row[0])] = "" if row[1] is None else str(row[1])
    if meta.get("mode") != MODE:
        rep.add(META_SHEET, "", "not_a_builder", f"{xlsx.name} is not a Spec Builder workbook (no {META_SHEET} mode {MODE})")
        return None, rep
    if base_spec is not None:
        cur = spec_hash(normalise(base_spec))
        src = meta.get("source_spec_hash") or ""
        again = previous_spec is not None and src and spec_hash(normalise(previous_spec)) == src
        if src and src != cur and again:
            rep.warn("this builder was read before (its source is template_spec.previous.json): the earlier read's "
                     "result is replaced")
        elif src and src != cur:
            rep.add(META_SHEET, "", "spec_changed", "the template spec changed since this builder was made (its spec hash "
                    f"{cur[:16]} is not the builder's source {src[:16]}): make a new builder from the current spec (step "
                    "0a) and copy your changes into it")
            return None, rep
        if not src:
            rep.warn("blank builder read against an existing spec: the spec is replaced by the builder's content "
                     "(fields the builder does not carry are still kept)")
    R = _Reader(wb, rep)
    tabs = {n: _Table(wb, n, rep) for n in T if T[n] is not None}
    if rep.problems:
        return None, rep
    spec = _assemble(R, tabs, rep, base_spec)
    if rep.problems or spec is None:
        return None, rep
    spec = _merge_base(spec, base_spec, rep)
    # the spec validator and the cedant model are the final judges (one implementation, spec.py)
    errs = validate_template_spec(spec)
    for e in errs:
        rep.add("", "", "spec_invalid", e)
    if not errs:
        try:
            normalise(spec)
        except Exception as e:  # noqa: BLE001
            rep.add("", "", "spec_invalid", str(e))
    if rep.problems:
        return None, rep
    rep.counts = {"sheets": len(spec["sheets"]), "blocks": sum(len(s["blocks"]) for s in spec["sheets"]),
                  "plans": sum(len(b["plans"]) for s in spec["sheets"] for b in s["blocks"]),
                  "factor tables": sum(len(s.get("factor_tables") or []) for s in spec["sheets"]),
                  "riders": sum(len(s.get("riders") or []) for s in spec["sheets"])}
    return spec, rep


# ============================================================================================== rows -> spec
def _assemble(R: _Reader, tabs: dict, rep: BuilderReadReport, base: dict | None) -> dict | None:
    reg = load_registry()
    fvs = R.fv("Tender")
    fvc = R.fv("Cedant workbook")
    spec: dict = {}
    for c, *_ in TENDER_FIELDS:
        if c in fvs.get("tender", {}):
            spec[c] = fvs["tender"][c]
    # age bands
    file_sets = _file_band_sets(base) if base else {}
    bsets, bfile, bcell = {}, {}, {}
    t = tabs["Age bands"]
    for r, row in t.rows:
        name, lab = R.val(t, r, row, "Band set"), R.val(t, r, row, "Label")
        lo, hi = R.val(t, r, row, "Age from"), R.val(t, r, row, "Age to")
        if name is None or lab is None:
            continue
        if (lo is None) != (hi is None):
            rep.add(t.name, t.cell(r, "Age to" if hi is None else "Age from"), "age_range", "give both Age from and Age to, "
                    "or neither (label-only row)")
            continue
        if lo is not None and lo > hi:
            rep.add(t.name, t.cell(r, "Age from"), "age_range", f"Age from {lo} is above Age to {hi}")
            continue
        b = {"label": lab, "range": [lo, hi]} if lo is not None else {"label": lab, "range": None}
        src = _s(row.get("Source")) or ""
        (bfile if src.startswith("file") and name in file_sets else bsets).setdefault(name, []).append(b)
        bcell.setdefault(name, []).append((r, b))
        if src.startswith("file") and name not in file_sets:
            rep.warn(f"age band set '{name}' is marked as coming from a file the base spec does not name: written into the spec")
    for name, items in bcell.items():
        seen_l = set()
        for (r, b) in items:
            if b["label"] in seen_l:
                rep.add(t.name, t.cell(r, "Label"), "duplicate", f"band set '{name}': label '{b['label']}' twice")
            seen_l.add(b["label"])
        rngs = sorted([(b["range"], r) for r, b in items if b["range"]])
        for (a, _), (b2, r2) in zip(rngs, rngs[1:]):
            if b2[0] <= a[1]:
                rep.add(t.name, t.cell(r2, "Age from"), "overlap", f"band set '{name}': ranges {a} and {b2} overlap")
    if base and base.get("age_bands_file"):
        for name, bands in bfile.items():
            fb = [{"label": b.get("label"), "range": b.get("range")} for b in file_sets[name]]
            if bands != fb:
                rep.add(t.name, "", "file_bands_changed", f"band set '{name}' differs from {base['age_bands_file']}: edit "
                        "that file (the builder does not overwrite it), or rename the set to make a new one in the spec")
        gone = [n for n in file_sets if n not in bfile]
        if gone:
            rep.warn(f"band sets {gone} of {base['age_bands_file']} are not on the Age bands sheet: still in the file")
    known_bands = set(bsets) | (set(file_sets) if base and base.get("age_bands_file") else set())
    # column sets
    csets: dict = {}
    t = tabs["Columns"]
    for r, row in t.rows:
        name, lab = R.val(t, r, row, "Column set"), R.val(t, r, row, "Column label")
        g, m, mc = R.val(t, r, row, "Gender"), R.val(t, r, row, "Marital"), R.val(t, r, row, "Member category")
        if name is None or lab is None:
            continue
        if name in DEFAULT_COLUMN_SETS:
            rep.add(t.name, t.cell(r, "Column set"), "builtin", f"'{name}' is a built-in column set: use another name")
            continue
        key = {k: v for k, v in (("gender", g), ("marital", m), ("member_category", mc)) if v is not None}
        cols = csets.setdefault(name, [])
        if cols and cols[-1]["label"] == lab:
            cols[-1]["keys"].append(key)
        elif any(c["label"] == lab for c in cols):
            rep.add(t.name, t.cell(r, "Column label"), "duplicate", f"column set '{name}': column '{lab}' appears twice "
                    "(rows of one column must be next to each other)")
        else:
            cols.append({"label": lab, "keys": [key]})
    known_cols = set(csets) | set(DEFAULT_COLUMN_SETS)
    if csets or (base and "column_sets" in base):
        spec["column_sets"] = csets
    if bsets or (base and "age_band_sets" in base) or not (base and base.get("age_bands_file")):
        spec["age_band_sets"] = bsets
    if base and base.get("age_bands_file"):
        spec["age_bands_file"] = base["age_bands_file"]

    def ref_band(tab, r, h, v):
        if v is not None and v not in known_bands:
            rep.add(tab.name, tab.cell(r, h), "undefined", f"age band set '{v}' is not defined on the Age bands sheet")

    def ref_cols(tab, r, h, v):
        if v is not None and v not in known_cols:
            rep.add(tab.name, tab.cell(r, h), "undefined", f"column set '{v}' is not defined on the Columns sheet")
    # sheets
    sheets: dict = {}
    t = tabs["Sheets"]
    for r, row in t.rows:
        n = R.val(t, r, row, "Sheet name")
        vals = {h: R.val(t, r, row, h) for h in t.cols if h != "Sheet name"}
        if n is None:
            continue
        if n in sheets:
            rep.add(t.name, t.cell(r, "Sheet name"), "duplicate", f"sheet name '{n}' twice")
            continue
        sh = {"sheet_name": n}
        for h, k in (("Title", "title"), ("Product type", "product_type"), ("Regulator", "regulator"),
                     ("Rating region", "rating_region"), ("Age band set", "age_bands"), ("Column set", "column_set"),
                     ("Refer text", "refer_text"), ("Premium basis", "premium_basis"), ("Scope id", "scope_id")):
            if vals.get(h) is not None:
                sh[k] = vals[h]
        if vals.get("Mandatory") is not None:
            sh["mandatory"] = vals["Mandatory"] == "Y"
        ref_band(t, r, "Age band set", vals.get("Age band set"))
        ref_cols(t, r, "Column set", vals.get("Column set"))
        ref_cols(t, r, "Marital split columns", vals.get("Marital split columns"))
        if vals.get("Marital split columns") is not None:
            sh.setdefault("split_options", {})["marital"] = {"column_set": vals["Marital split columns"]}
        f = {}
        if vals.get("Formula order") is not None:
            f["order"] = list(ORDERS[vals["Formula order"]])
        for h, k in (("Formula rounding", "rounding"), ("Formula min premium", "min_premium"),
                     ("Formula max premium", "max_premium")):
            if vals.get(h) is not None:
                f[k] = vals[h]
        if f:
            sh["formula"] = f
        sh["copay_basis"] = {}
        sh["blocks"] = []
        sh["_row"] = r
        sheets[n] = sh

    def sheet_of(tab, r, row):
        n = R.val(tab, r, row, "Sheet")
        if n is None:
            return None
        if n not in sheets:
            rep.add(tab.name, tab.cell(r, "Sheet"), "undefined", f"sheet '{n}' is not on the Sheets sheet")
            return None
        return sheets[n]
    t = tabs["Co-pay basis"]
    for r, row in t.rows:
        sh, fam, opt = sheet_of(t, r, row), R.val(t, r, row, "Family"), R.val(t, r, row, "Option")
        if sh is None or fam is None or opt is None:
            continue
        if fam in sh["copay_basis"]:
            rep.add(t.name, t.cell(r, "Family"), "duplicate", f"sheet '{sh['sheet_name']}': family '{fam}' twice")
            continue
        sh["copay_basis"][fam] = opt
    # blocks
    blocks: dict = {}
    t = tabs["Blocks"]
    for r, row in t.rows:
        sh = sheet_of(t, r, row)
        vals = {h: R.val(t, r, row, h) for h in t.cols if h != "Sheet"}
        if sh is None or vals.get("Block id") is None or vals.get("TPA") is None:
            continue
        bid = vals["Block id"]
        if (sh["sheet_name"], bid) in blocks:
            rep.add(t.name, t.cell(r, "Block id"), "duplicate", f"block id '{bid}' twice on sheet '{sh['sheet_name']}'")
            continue
        b = {"block_id": bid, "tpa": vals["TPA"]}
        for h, k in (("Label", "label"), ("Age band set", "age_bands"), ("Column set", "column_set"),
                     ("Refer text", "refer_text"), ("Comment", "comment")):
            if vals.get(h) is not None:
                b[k] = vals[h]
        sp = {}
        for h, k in (("Marital split", "marital"), ("Member split", "member_section")):
            if vals.get(h) is not None:
                sp[k] = vals[h] == "Y"
        if sp:
            b["splits"] = sp
        ref_band(t, r, "Age band set", vals.get("Age band set"))
        ref_cols(t, r, "Column set", vals.get("Column set"))
        b["plans"] = []
        b["_row"] = r
        sh["blocks"].append(b)
        blocks[(sh["sheet_name"], bid)] = b

    def block_of(tab, r, row, h="Block id"):
        sh = sheet_of(tab, r, row)
        bid = R.val(tab, r, row, h)
        if sh is None or bid is None:
            return None
        b = blocks.get((sh["sheet_name"], bid))
        if b is None:
            rep.add(tab.name, tab.cell(r, h), "undefined", f"block '{bid}' is not a block of sheet '{sh['sheet_name']}' "
                    "on the Blocks sheet")
        return b
    t = tabs["Plans"]
    for r, row in t.rows:
        b = block_of(t, r, row)
        pid = R.val(t, r, row, "Plan id")
        if b is None or pid is None:
            continue
        if any(p["plan_id"] == pid for p in b["plans"]):
            rep.add(t.name, t.cell(r, "Plan id"), "duplicate", f"block '{b['block_id']}': plan '{pid}' twice")
            continue
        p = {"plan_id": pid}
        lab = R.val(t, r, row, "Plan label")
        if lab is not None:
            p["label"] = lab
        p["network"] = R.val(t, r, row, "Network")
        num = R.val(t, r, row, "Plan number")
        if num is not None:
            p["plan_number"] = num
        b["plans"].append(p)
    for b in blocks.values():
        if not b["plans"]:
            rep.warn(f"block '{b['block_id']}' has no plans on the Plans sheet")

    def target(tab, r, row, allow_split: bool):
        sh = sheet_of(tab, r, row)
        a = R.val(tab, r, row, "Applies to")
        if sh is None or a is None:
            return None, None
        if a == APPLIES_SHEET:
            return sh, "sheet"
        if a == APPLIES_SPLIT:
            if allow_split:
                return sh, "split"
            rep.add(tab.name, tab.cell(r, "Applies to"), "not_allowed", f"'{APPLIES_SPLIT}' is for the Sections sheet only")
            return None, None
        b = blocks.get((sh["sheet_name"], a))
        if b is None:
            rep.add(tab.name, tab.cell(r, "Applies to"), "undefined", f"'{a}' is not '{APPLIES_SHEET}', "
                    f"{'' if not allow_split else repr(APPLIES_SPLIT) + ', '}or a block of sheet '{sh['sheet_name']}'")
            return None, None
        return b, "block"
    t = tabs["Sections"]
    for r, row in t.rows:
        obj, kind = target(t, r, row, True)
        lab = R.val(t, r, row, "Section label")
        mc, bs, rl = R.val(t, r, row, "Member category"), R.val(t, r, row, "Age band set"), R.val(t, r, row, "Row label (no bands)")
        if obj is None or lab is None:
            continue
        ref_band(t, r, "Age band set", bs)
        if bs is not None and rl is not None:
            rep.add(t.name, t.cell(r, "Row label (no bands)"), "conflict", "give an age band set or a row label, not both")
            continue
        sec = {"label": lab, "key": {"member_category": mc} if mc is not None else {}}
        if bs is not None:
            sec["age_bands"] = bs
        if rl is not None:
            sec["row_label"] = rl
        if kind == "split":
            obj.setdefault("split_options", {}).setdefault("member_section", {"sections": []})["sections"].append(sec)
        else:
            obj.setdefault("sections", []).append(sec)
            obj.setdefault("_srow", r)
    t = tabs["Custom rows"]
    for r, row in t.rows:
        obj, kind = target(t, r, row, False)
        lab, sec = R.val(t, r, row, "Row label"), R.val(t, r, row, "Section label")
        mc, lo, hi = R.val(t, r, row, "Member category"), R.val(t, r, row, "Age from"), R.val(t, r, row, "Age to")
        if obj is None or lab is None:
            continue
        if (lo is None) != (hi is None) or (lo is not None and lo > hi):
            rep.add(t.name, t.cell(r, "Age from"), "age_range", "give Age from <= Age to, or neither")
            continue
        key = {}
        if mc is not None:
            key["member_category"] = mc
        if lo is not None:
            key["age"] = [lo, hi]
        rws = obj.setdefault("rows", [])
        obj.setdefault("_rrow", r)
        if rws and rws[-1]["label"] == lab and rws[-1].get("section") == sec:
            prev = rws[-1]
            if "keys" not in prev:
                prev["keys"] = [prev.pop("key")]
            prev["keys"].append(key)
        else:
            nr = {"label": lab}
            if sec is not None:
                nr["section"] = sec
            nr["key"] = key
            rws.append(nr)
    for obj in list(sheets.values()) + list(blocks.values()):
        has = [k for k in ("age_bands", "sections", "rows") if obj.get(k)]
        if len(has) > 1:
            where = "sheet" if "sheet_name" in obj else "block"
            name = obj.get("sheet_name") or obj.get("block_id")
            tab = "Sheets" if where == "sheet" else "Blocks"
            rep.add(tab, tabs[tab].cell(obj["_row"], "Age band set"), "row_source", f"{where} '{name}' has "
                    f"{' and '.join(has)}: give one row source per {where} (age band set, sections or custom rows)")
    t = tabs["Refer cells"]
    for r, row in t.rows:
        b = block_of(t, r, row)
        if b is None:
            continue
        k = {}
        for h, d in (("Plan", "plan"), ("Gender", "gender"), ("Marital", "marital"), ("Member category", "member_category")):
            v = R.val(t, r, row, h)
            if v is not None:
                k[d] = v
        lo, hi = R.val(t, r, row, "Age from"), R.val(t, r, row, "Age to")
        if (lo is None) != (hi is None) or (lo is not None and lo > hi):
            rep.add(t.name, t.cell(r, "Age from"), "age_range", "give Age from <= Age to, or neither")
            continue
        if lo is not None:
            k["age"] = [lo, hi]
        if not k:
            rep.add(t.name, t.cell(r, "Plan"), "blank_mandatory", "a refer cell needs at least one field (all blank = every cell)")
            continue
        if "plan" in k and k["plan"] not in [p["plan_id"] for p in b["plans"]]:
            rep.add(t.name, t.cell(r, "Plan"), "undefined", f"plan '{k['plan']}' is not a plan of block '{b['block_id']}'")
        b.setdefault("refer_cells", []).append(k)
    # cedant sections: tables, riders, options
    tables: dict = {}
    for tname, kind in (("Factor tables", "factor_tables"), ("Riders", "riders")):
        t = tabs[tname]
        for r, row in t.rows:
            sh = sheet_of(t, r, row)
            vals = {h: R.val(t, r, row, h) for h in t.cols if h != "Sheet"}
            if sh is None or vals.get("Component id") is None or vals.get("Family") is None:
                continue
            cid = vals["Component id"]
            if (sh["sheet_name"], cid) in tables:
                rep.add(t.name, t.cell(r, "Component id"), "duplicate", f"component '{cid}' twice on sheet '{sh['sheet_name']}'")
                continue
            tb = {"component_id": cid, "family": vals["Family"]}
            for h, k in (("Option dimension", "option_dim"), ("Key on", "key_dim"), ("Default option", "default_option"),
                         ("Label", "label")):
                if vals.get(h) is not None:
                    tb[k] = vals[h]
            if vals.get("Key order") is not None:
                tb["keys"] = _csv(vals["Key order"])
            if vals.get("TPAs") is not None:
                tb["tpas"] = _csv(vals["TPAs"])
            if kind == "riders":
                for h, k in (("Unit", "unit"), ("Mode", "mode"), ("Premium basis", "premium_basis")):
                    if vals.get(h) is not None:
                        tb[k] = vals[h]
                if vals.get("Mandatory") is not None:
                    tb["mandatory"] = vals["Mandatory"] == "Y"
                if vals.get("Eligibility") is not None:
                    el = {}
                    for part in _csv(vals["Eligibility"]):
                        if "=" not in part or part.split("=", 1)[0].strip() not in reg.dims:
                            rep.add(t.name, t.cell(r, "Eligibility"), "eligibility", "eligibility must be dimension=value, "
                                    "comma separated (e.g. maternity_eligible=Y)", vals["Eligibility"])
                            el = None
                            break
                        d, v = part.split("=", 1)
                        el[d.strip()] = v.strip()
                    if el is None:
                        continue
                    tb["eligibility"] = el
            tb["_kind"], tb["_rows"], tb["_labels"], tb["_row"], tb["_tab"] = kind, [], {}, r, tname
            sh.setdefault(kind, []).append(tb)
            tables[(sh["sheet_name"], cid)] = tb
    t = tabs["Options"]
    for r, row in t.rows:
        sh = sheet_of(t, r, row)
        cid, key, opt = R.val(t, r, row, "Component id"), R.val(t, r, row, "Key"), R.val(t, r, row, "Option")
        refy, lab = R.val(t, r, row, "Reference"), R.val(t, r, row, "Key label")
        if sh is None or cid is None or opt is None:
            continue
        tb = tables.get((sh["sheet_name"], cid))
        if tb is None:
            rep.add(t.name, t.cell(r, "Component id"), "undefined", f"component '{cid}' is not on the Factor tables or "
                    f"Riders sheet for sheet '{sh['sheet_name']}'")
            continue
        k = ANY if key is None else key
        if key is not None and not tb.get("key_dim"):
            rep.add(t.name, t.cell(r, "Key"), "no_key_dim", f"component '{cid}' has no 'Key on': leave Key blank")
            continue
        if any(x["key"] == k and x["option"] == opt for x in tb["_rows"]):
            rep.add(t.name, t.cell(r, "Option"), "duplicate", f"component '{cid}': key {k} option '{opt}' twice")
            continue
        if refy == "Y" and tb["_kind"] == "riders":
            rep.add(t.name, t.cell(r, "Reference"), "not_allowed", "riders have no reference option")
            continue
        tb["_rows"].append({"key": k, "option": opt, "ref": refy == "Y", "r": r})
        if lab is not None:
            if tb["_labels"].get(k, lab) != lab:
                rep.add(t.name, t.cell(r, "Key label"), "conflict", f"component '{cid}': key {k} has two different labels")
            tb["_labels"][k] = lab
    for (sname, cid), tb in tables.items():
        _finish_table(tb, rep, tabs)
    t = tabs["Plan defaults"]
    for r, row in t.rows:
        sh, fam = sheet_of(t, r, row), R.val(t, r, row, "Family")
        if sh is None or fam is None:
            continue
        pd_ = sh.setdefault("plan_defaults", [])
        if fam in pd_:
            rep.add(t.name, t.cell(r, "Family"), "duplicate", f"sheet '{sh['sheet_name']}': family '{fam}' twice")
            continue
        pd_.append(fam)
    t = tabs["Benefits"]
    for r, row in t.rows:
        sh, code, plans = sheet_of(t, r, row), R.val(t, r, row, "Benefit code"), R.val(t, r, row, "Plans")
        if sh is None or code is None:
            continue
        bl = sh.setdefault("benefits", [])
        if any((x if isinstance(x, str) else x["code"]) == code for x in bl):
            rep.add(t.name, t.cell(r, "Benefit code"), "duplicate", f"sheet '{sh['sheet_name']}': benefit '{code}' twice")
            continue
        bl.append(code if plans is None else {"code": code, "plans": _csv(plans)})
    # tender sections
    if fvs.get("reference_sheets") or (base and "reference_sheets" in base):
        spec["reference_sheets"] = dict(fvs.get("reference_sheets") or {})
    opts = dict(fvs.get("options") or {})
    lo, hi = opts.pop("outlier_band_low", None), opts.pop("outlier_band_high", None)
    if (lo is None) != (hi is None):
        rep.add("Tender", "", "outlier_band", "give both outlier band values, or neither")
    elif lo is not None:
        opts["outlier_band"] = [lo, hi]
    if opts or (base and "options" in base):
        spec["options"] = opts
    cw = dict(fvc.get("cedant_workbook") or {})
    notes = []
    t = tabs["Cedant notes"]
    for r, row in t.rows:
        v = R.val(t, r, row, "Note")
        if v is not None:
            notes.append(v)
    if cw or notes or (base and "cedant_workbook" in base):
        if notes or (base and "notes" in (base.get("cedant_workbook") or {})):
            cw["notes"] = notes
        spec["cedant_workbook"] = cw
    spec["sheets"] = list(sheets.values())
    if not spec["sheets"]:
        rep.add("Sheets", "", "blank_mandatory", "no rate sheet on the Sheets sheet")
    for sh in spec["sheets"]:
        if not sh["blocks"]:
            rep.warn(f"sheet '{sh['sheet_name']}' has no blocks on the Blocks sheet")
    # strip helpers
    for sh in spec["sheets"]:
        for o in [sh, *sh["blocks"], *(sh.get("factor_tables") or []), *(sh.get("riders") or [])]:
            for k in [k for k in o if k.startswith("_")]:
                del o[k]
    return spec


def _finish_table(tb: dict, rep: BuilderReadReport, tabs: dict):
    """Options rows -> the table's canonical form: options x key_values when the rows are a full grid in key order,
    else rows; reference as one option when every key shares it, else by key."""
    rows, kind = tb["_rows"], tb["_kind"]
    tab = tabs[tb["_tab"]]
    if not rows:
        rep.add(tab.name, tab.cell(tb["_row"], "Component id"), "no_options",
                f"component '{tb['component_id']}' has no rows on the Options sheet")
        return
    keys = list(dict.fromkeys(r["key"] for r in rows))
    by_key = {k: [r["option"] for r in rows if r["key"] == k] for k in keys}
    grid = (all(by_key[k] == by_key[keys[0]] for k in keys)
            and [(r["key"], r["option"]) for r in rows] == [(k, o) for k in keys for o in by_key[keys[0]]])
    if grid and keys == [ANY]:
        tb["options"] = list(by_key[ANY])
    elif grid and ANY not in keys:
        tb["options"] = list(by_key[keys[0]])
        tb["key_values"] = keys
    else:
        tb["rows"] = [{"key": r["key"], "option": r["option"]} for r in rows]
    if tb["_labels"]:
        tb["key_labels"] = dict(tb["_labels"])
    if kind == "factor_tables":
        refs = {}
        for r in rows:
            if r["ref"]:
                if r["key"] in refs:
                    rep.add("Options", tabs["Options"].cell(r["r"], "Reference"), "reference_count",
                            f"component '{tb['component_id']}': key {r['key']} has two reference rows")
                refs[r["key"]] = r["option"]
        missing = [k for k in keys if k not in refs and not (k != ANY and ANY in refs)]
        if missing:
            rep.add(tab.name, tab.cell(tb["_row"], "Component id"), "reference_count",
                    f"component '{tb['component_id']}': key(s) {missing} have no reference option (Reference Y on the "
                    "Options sheet)")
        if refs:
            vals = set(refs.values())
            one = next(iter(vals)) if len(vals) == 1 else None
            # one option for every key when that marks exactly the same rows as the Y flags, else by key
            same = one is not None and all((r["option"] == one) == r["ref"] for r in rows)
            tb["reference"] = one if same else dict(refs)


# ============================================================================================== keep what the builder does not carry
def _semantic_block(nspec: dict, sname: str, bid: str):
    for sh in nspec["sheets"]:
        if sh["sheet_name"] == sname:
            for b in sh["blocks"]:
                if b["block_id"] == bid:
                    return b
    return None


def _table_model(sheet: dict, kind: str, table: dict):
    """The cedant model of one table on its sheet (the other tables, defaults and benefits left out, so their own
    problems do not hide this one's meaning). None when the table itself is invalid."""
    mini = {k: sheet.get(k) for k in ("sheet_name", "product_type", "regulator", "rating_region", "blocks", "scope_id")}
    mini[kind] = [table]
    m, errs = cedant_model({"sheets": [mini]})
    if not m["sheets"] or any(f"{kind}[0]" in e for e in errs):
        return None
    return next((t for t in m["sheets"][0]["factor_tables"] + m["sheets"][0]["riders"]
                 if t["component_id"] == table.get("component_id")), None)


def _merge_base(spec: dict, base: dict | None, rep: BuilderReadReport) -> dict:
    """Keep the base spec's fields the builder does not carry, and the base's raw form of every block and table whose
    meaning is unchanged (SPEC §8.6 round trip)."""
    if base is None:
        return spec
    base = copy.deepcopy(base)
    base_dir = base.pop("_base_dir", None)
    out = {k: v for k, v in spec.items()}
    for k, v in base.items():
        if k not in TOP_KEYS:
            out[k] = v
    if "options" in out:
        out["options"] = _order_like({**{k: v for k, v in (base.get("options") or {}).items() if k not in OPTIONS_KEYS},
                                      **out["options"]}, base.get("options"))
    if "reference_sheets" in out:
        out["reference_sheets"] = _order_like(out["reference_sheets"], base.get("reference_sheets"))
    if "cedant_workbook" in out:
        bcw = base.get("cedant_workbook") or {}
        out["cedant_workbook"] = _order_like({**{k: v for k, v in bcw.items() if k not in CW_KEYS}, **out["cedant_workbook"]}, bcw)
    for k in ("column_sets", "age_band_sets"):
        if k in out and k in base:
            out[k] = _order_like(out[k], base[k])
    try:
        cand = normalise({**out, "_base_dir": base_dir})
        bnorm = normalise({**base, "_base_dir": base_dir})
    except Exception:  # noqa: BLE001  (the final validation reports it)
        cand = bnorm = None
    bsheets = {s.get("sheet_name"): s for s in base.get("sheets") or []}
    new_sheets = []
    for sh in out["sheets"]:
        bsh = bsheets.get(sh["sheet_name"])
        if bsh is None:
            new_sheets.append(sh)
            continue
        sh = dict(sh)
        for k, v in bsh.items():
            if k not in SHEET_KEYS:
                sh[k] = v
        if "copay_basis" not in bsh and sh.get("copay_basis") == {}:
            del sh["copay_basis"]
        if "formula" in sh and isinstance(bsh.get("formula"), dict):
            sh["formula"] = _order_like({**{k: v for k, v in bsh["formula"].items() if k not in FORMULA_KEYS},
                                         **sh["formula"]}, bsh["formula"])
        bblocks = {b.get("block_id"): b for b in bsh.get("blocks") or []}
        nb = []
        for b in sh["blocks"]:
            bb = bblocks.get(b["block_id"])
            if bb is None:
                nb.append(b)
                continue
            b = dict(b)
            for k, v in bb.items():
                if k not in BLOCK_KEYS:
                    b[k] = v
            bplans = {p.get("plan_id"): p for p in bb.get("plans") or []}
            b["plans"] = [_order_like({**{k: v for k, v in (bplans.get(p["plan_id"]) or {}).items() if k not in PLAN_KEYS}, **p},
                                      bplans.get(p["plan_id"])) for p in b["plans"]]
            if cand is not None and bnorm is not None:
                x, y = _semantic_block(cand, sh["sheet_name"], b["block_id"]), _semantic_block(bnorm, sh["sheet_name"], b["block_id"])
                if x is not None and x == y:
                    b = copy.deepcopy(bb)            # unchanged meaning: keep the base's own form
            nb.append(_order_like(b, bb))
        sh["blocks"] = nb
        for kind, keys in (("factor_tables", TABLE_KEYS), ("riders", RIDER_KEYS)):
            btab = {t.get("component_id"): t for t in bsh.get(kind) or []}
            nt = []
            for t in sh.get(kind) or []:
                bt = btab.get(t["component_id"])
                if bt is None:
                    nt.append(t)
                    continue
                t = {**{k: v for k, v in bt.items() if k not in keys}, **t}
                cm, bm = _table_model(sh, kind, t), _table_model(bsh, kind, bt)
                if cm is not None and cm == bm:
                    t = copy.deepcopy(bt)
                nt.append(_order_like(t, bt))
            if kind in sh:
                sh[kind] = nt
        for k in ("benefits", "plan_defaults"):
            if k not in sh and k in bsh and not bsh[k]:
                sh[k] = copy.deepcopy(bsh[k])          # an empty list in the base stays an empty list
        new_sheets.append(_order_like(sh, bsh))
    out["sheets"] = new_sheets
    gone = [n for n in bsheets if n not in {s["sheet_name"] for s in out["sheets"]}]
    if gone:
        rep.warn(f"sheets removed from the spec: {gone}")
    out = _order_like(out, base)
    out["_base_dir"] = base_dir
    return out


def write_spec(spec: dict, path: Path):
    """Write a spec as the tool writes JSON (indent 1, UTF-8, LF, trailing newline); internal keys left out."""
    d = {k: v for k, v in spec.items() if k != "_base_dir"}
    Path(path).write_text(json.dumps(d, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
