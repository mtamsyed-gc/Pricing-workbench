"""Cedant rate workbook (SPEC §8.5; work order WS5): the cedant's current (expiring) catalogue in the tool's own layout.

    generate_cedant_workbook(spec, out)            -> GeneratedCedantWorkbook   (py tools\\pw.py t cedant ...)
    read_cedant(xlsx, spec, ...)                   -> (Catalogue | None, CedantReadReport)
    read_cedant_workbook(xlsx, spec, ...)          -> Catalogue (kind EXPIRING); raises GateFailed on any problem
    fill_from_catalogue(xlsx, out, catalogue)      -> fill a generated workbook from an existing catalogue (fixtures)

One layout, two modes, one parser (SPEC §8.5): the rate sheets are the tender's (same blocks, keyed cells in _KEYS);
cedant mode adds four long tables - Factors, Riders, Plan defaults, Benefits - whose labels come from the template
spec (cedant sections, see spec.cedant_model) and whose yellow cells GC fills. Co-pay variants are never entered: the
Nil / 0% base and the factor tables generate them. The read-back is gated (stops, never defaults) and lists every
problem with sheet and cell. The filled copy holds client rates and lives only in runs/.

Read-back scopes: one RatingScope per sheet x group of TPAs sharing the same factor tables and riders; the base table
keeps the TPA as a key (each TPA block is entered separately, before the TPA factor - SPEC §8.0 TPA basis).
"""
from __future__ import annotations

import datetime as _dt
import json
import math
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .. import GateFailed, SpecError
from ..io_excel import a1, new_workbook, quote_sheet
from ..model import BenefitLine, Catalogue, Component, Plan, RatingFormula, RatingScope, canonical_key_str
from ..registry import ANY, load_registry
from . import GENERATOR_VERSION, KEYS_SHEET, META_SHEET, TEMPLATE_SPEC_VERSION
from .generator import InputCell, Layout, _Fmt, compute_layout, hash_keys_rows, write_keys_meta, write_rate_sheets
from .parser import _keys_table, _locate, fill_template, read_meta
from .spec import ALL_TPAS, NONE_OPTION, cedant_model, cedant_spec_hash, normalise, spec_hash

MODE = "CEDANT"
FACTORS, RIDERS, DEFAULTS, BENEFITS, LISTS = "Factors", "Riders", "Plan defaults", "Benefits", "_LISTS"
LONG_SHEETS = (FACTORS, RIDERS, DEFAULTS, BENEFITS)
HEADER_ROW = 3                                    # 0-based row of the long-table header
COLS = {
    FACTORS: ["Product", "Emirate", "TPAs", "Family", "Key on", "Key", "Option", "Factor", "Reference"],
    RIDERS: ["Product", "Emirate", "TPAs", "Rider", "Limit", "Key on", "Key", "Unit", "Eligibility", "Rate"],
    DEFAULTS: ["Product", "Emirate", "TPAs", "Plan", "Family", "Option"],
    BENEFITS: ["Product", "Emirate", "TPAs", "Plan", "Benefit code", "Value type", "Allowed", "Value", "Condition text"],
}
KIND_OF = {FACTORS: "FACTOR", RIDERS: "RIDER", DEFAULTS: "DEFAULT", BENEFITS: "BENEFIT"}
SPARE_EDITABLE = {FACTORS: ["Product", "Emirate", "TPAs", "Family", "Key", "Option", "Factor", "Reference"],
                  RIDERS: ["Product", "Emirate", "TPAs", "Rider", "Limit", "Key", "Rate"],
                  DEFAULTS: ["Product", "Emirate", "TPAs", "Plan", "Family", "Option"],
                  BENEFITS: ["Product", "Emirate", "TPAs", "Plan", "Benefit code", "Value", "Condition text"]}
NUMERIC = {FACTORS: ["Factor"], RIDERS: ["Rate"]}
YES = ("Y", "YES", "TRUE")
BOOL_TEXT = {"Y": True, "YES": True, "TRUE": True, "N": False, "NO": False, "FALSE": False}
_EPS = 1e-12


# ============================================================================================== layout
def _num_fmt(wb):
    return {"txt": wb.add_format({"border": 1}),
            "ro_txt": wb.add_format({"border": 1, "bg_color": "#F2F2F2"}),
            "inp_txt": wb.add_format({"bg_color": "#FFFF99", "border": 1, "locked": False}),
            "inp_num": wb.add_format({"bg_color": "#FFFF99", "border": 1, "locked": False, "num_format": "0.0000##"}),
            "inp_rate": wb.add_format({"bg_color": "#FFFF99", "border": 1, "locked": False, "num_format": "#,##0.00"}),
            "inp_gen": wb.add_format({"bg_color": "#FFFF99", "border": 1, "locked": False, "num_format": "General"})}


class _Lists:
    """Dropdown sources on the hidden _LISTS sheet (one column per list; Excel list validation by range)."""

    def __init__(self):
        self.cols: list[tuple[str, list]] = []
        self.ref: dict[str, str] = {}

    def add(self, name: str, values: list) -> str | None:
        vals = [str(v) for v in values if v is not None and str(v) != ""]
        vals = list(dict.fromkeys(vals))
        if not vals:
            return None
        if name in self.ref:
            return self.ref[name]
        c = len(self.cols)
        self.cols.append((name, vals))
        self.ref[name] = (f"={quote_sheet(LISTS)}!{a1(2, c + 1, absolute=True)}:"
                          f"{a1(len(vals) + 1, c + 1, absolute=True)}")
        return self.ref[name]

    def write(self, ws):
        for c, (name, vals) in enumerate(self.cols):
            ws.write_string(0, c, name)
            for i, v in enumerate(vals):
                ws.write_string(1 + i, c, v)
        ws.hide()


def _row_key_dict(**kw) -> dict:
    return {k: v for k, v in kw.items()}


def _ck(kind: str, d: dict) -> str:
    return kind.lower() + "|" + "|".join(f"{k}={'+'.join(v) if isinstance(v, list) else v}" for k, v in d.items())


@dataclass
class GeneratedCedantWorkbook:
    path: Path
    structure_hash: str
    spec_hash: str
    n_inputs: int
    n_refer: int
    counts: dict
    layout: Layout


def _key_on(t: dict, r: dict) -> str:
    lab = t.get("key_labels", {}).get(r["key"])
    return (t["key_dim"] or ANY) + (f" ({lab})" if lab else "")


def long_rows(model: dict) -> dict:
    """Pre-filled long-table rows from the cedant model: {sheet: [{"cols": {col: value}, "input": {col: (kind, ck, keys,
    mandatory)}, "lists": {col: list name}, "msh": ...}]}. The generator writes these; the fill helper uses the keys."""
    reg = load_registry()
    out = {s: [] for s in LONG_SHEETS}
    for msh in model["sheets"]:
        pe = {"product_type": msh["product_type"], "rating_region": msh["rating_region"]}
        prod, emi = msh["product_type"], msh["rating_region"] or ""
        for t in msh["factor_tables"]:
            for r in t["rows"]:
                k = {**pe, "sheet": msh["sheet_name"], "tpas": t["tpas"], "component_id": t["component_id"],
                     "family": t["family"], "key": r["key"], "option": r["option"]}
                out[FACTORS].append({"cols": {"Product": prod, "Emirate": emi, "TPAs": t["tpas_label"], "Family": t["family"],
                                              "Key on": _key_on(t, r), "Key": r["key"], "Option": r["option"],
                                              "Reference": "Y" if r["reference"] else ""},
                                     "input": {"Factor": ("FACTOR", _ck("factor", {"sheet": msh["sheet_name"], "tpas": t["tpas_label"],
                                                                                   "component": t["component_id"], "key": r["key"],
                                                                                   "option": r["option"]}), k, True)}})
        for t in msh["riders"]:
            elig = ", ".join(f"{d}={v}" for d, v in sorted(t["eligibility"].items())) or "all members"
            for r in t["rows"]:
                k = {**pe, "sheet": msh["sheet_name"], "tpas": t["tpas"], "component_id": t["component_id"],
                     "family": t["family"], "key": r["key"], "option": r["option"]}
                out[RIDERS].append({"cols": {"Product": prod, "Emirate": emi, "TPAs": t["tpas_label"], "Rider": t["family"],
                                             "Limit": r["option"], "Key on": _key_on(t, r), "Key": r["key"],
                                             "Unit": t["unit"] + ("" if t["mode"] == "ADD" else " (multiplier)"),
                                             "Eligibility": elig},
                                    "input": {"Rate": ("RIDER", _ck("rider", {"sheet": msh["sheet_name"], "tpas": t["tpas_label"],
                                                                              "component": t["component_id"], "key": r["key"],
                                                                              "option": r["option"]}), k, t["mandatory"])}})
        for g in msh["groups"]:
            for p in g["plans"]:
                for fam in g["default_families"]:
                    k = {**pe, "sheet": msh["sheet_name"], "tpas": g["tpas"], "plan": p["plan_id"], "family": fam}
                    out[DEFAULTS].append({"cols": {"Product": prod, "Emirate": emi, "TPAs": g["label"], "Plan": p["plan_id"],
                                                   "Family": fam},
                                          "lists": {"Option": ("opt", msh["sheet_name"], g["label"], fam)},
                                          "input": {"Option": ("DEFAULT", _ck("default", {"sheet": msh["sheet_name"],
                                                                                          "tpas": g["label"], "plan": p["plan_id"],
                                                                                          "family": fam}), k, True)}})
                for b in msh["benefits"]:
                    if b["plans"] is not None and p["plan_id"] not in b["plans"]:
                        continue
                    line = reg.benefits[b["code"]]
                    allowed = ", ".join(line.get("ladder") or []) or {"bool": "Y / N", "money_or_bool": "amount or Y / N"}.get(
                        line["value_type"], "")
                    k = {**pe, "sheet": msh["sheet_name"], "tpas": g["tpas"], "plan": p["plan_id"], "code": b["code"]}
                    base = {"sheet": msh["sheet_name"], "tpas": g["label"], "plan": p["plan_id"], "code": b["code"]}
                    out[BENEFITS].append({"cols": {"Product": prod, "Emirate": emi, "TPAs": g["label"], "Plan": p["plan_id"],
                                                   "Benefit code": b["code"], "Value type": line["value_type"],
                                                   "Allowed": allowed},
                                          "lists": {"Value": ("ben", b["code"])},
                                          "input": {"Value": ("BENEFIT", _ck("benefit", base), k, False),
                                                    "Condition text": ("BENEFIT_TEXT", _ck("benefit_text", base), k, False)}})
    return out


def _option_lists(model: dict) -> dict:
    """(sheet, group label, family) -> options a plan default may name (rider families add NONE = not purchased)."""
    out = {}
    for msh in model["sheets"]:
        for g in msh["groups"]:
            for fam in g["default_families"]:
                opts, rider = [], False
                for t in msh["factor_tables"] + msh["riders"]:
                    if t["family"] == fam and set(g["tpas"]) <= set(t["tpas"]):
                        opts += [r["option"] for r in t["rows"]]
                        rider = rider or t in msh["riders"]
                out[(msh["sheet_name"], g["label"], fam)] = list(dict.fromkeys(opts)) + ([NONE_OPTION] if rider else [])
    return out


def _cover_lines(nspec: dict, model: dict) -> list:
    notes = list(model["options"].get("notes") or [])
    return [("Purpose", "GC fills the cedant's CURRENT (expiring) catalogue in this workbook; the tool reads it back into "
                        "the EXPIRING catalogue that drives Existing Rates, V5, B2, B1 and A. Holds client rates: keep the "
                        "filled copy in runs/ only (never in the repo, never sent to reinsurers)."),
            ("Currency / basis", f"{nspec['currency']} / {nspec['premium_basis']} (a sheet states its own basis where it differs)"),
            ("Rate sheets", "Base rates only: Nil / 0% co-pay at the reference options stated on each sheet, annual per member. "
                            "Co-pay variants are NEVER entered: the Nil base and the factor tables generate them."),
            ("TPA basis", "Each TPA block holds the rate BEFORE the TPA factor. Enter the TPA factor in Factors (family tpa); "
                          "do not build the TPA difference into the block as well. A TPA whose source table is shared with "
                          "another TPA gets the same values in its own block."),
            ("Factors", "One row per option (and per key where the factor varies by network / plan / TPA, else ANY). Enter "
                        "the factor; the row marked Reference = Y is the option the base assumes and must read 1.0, exactly "
                        "one per key (a key without its own reference row uses the ANY row's)."),
            ("Riders", "Rider (add-on) premium per limit option and key, in the unit shown (per member / per eligible "
                       "female); a multiplier where the unit says so."),
            ("Plan defaults", "The option each plan is sold with (catalogue default, SPEC §3.4): used where the exposure does "
                              "not carry the option. Riders: NONE = not purchased."),
            ("Benefits", "Table of benefits per plan: value as the code needs (amount, number, Y / N, or a listed value); "
                         "keep the source wording in Condition text. Leave the value blank where it does not parse "
                         "(MANUAL_REVIEW) or the row blank where the benefit does not apply."),
            ("Spare rows", "Use the yellow spare rows at the bottom of a long table to add an option, key or benefit line "
                           "to an existing table (pick product, emirate, TPAs and family from the lists)."),
            ("Read back", "py tools\\pw.py t cedant-read <filled.xlsx> --spec <template_spec.json> --out runs\\<RunID>\\expiring.json"
                          "  - the read stops and lists every problem (sheet and cell) until the workbook is complete."),
            *[("Note", n) for n in notes]]


def generate_cedant_workbook(spec: dict, out_path: Path) -> GeneratedCedantWorkbook:
    nspec = normalise(spec)
    model, errs = cedant_model(spec)
    if errs:
        raise SpecError("template_spec cedant sections invalid:\n  " + "\n  ".join(errs))
    reg = load_registry()
    lay = compute_layout(nspec)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    created = _dt.datetime.fromisoformat(nspec["issue_date"]) if nspec.get("issue_date") else None
    wb = new_workbook(out_path, created=created, title=f"Cedant rates {nspec['tender_id']}")
    f = _Fmt(wb)
    nf = _num_fmt(wb)
    protect = nspec["options"]["protect_structure"]
    ws_cover = wb.add_worksheet("Cover")
    rate_ws = {s["sheet"]: wb.add_worksheet(s["sheet"]) for s in lay.rate_sheets}
    ws_long = {s: wb.add_worksheet(s) for s in LONG_SHEETS}
    ws_lists = wb.add_worksheet(LISTS)
    ws_keys = wb.add_worksheet(KEYS_SHEET)
    ws_meta = wb.add_worksheet(META_SHEET)
    msh_by = {m["sheet_name"]: m for m in model["sheets"]}

    def subtitle(sh, ns):
        basis = ", ".join(f"{k} = {v}" for k, v in sorted((sh.get("copay_basis") or {}).items()))
        pb = msh_by.get(sh["sheet_name"], {}).get("premium_basis") or ns["premium_basis"]
        return (f"CURRENT (expiring) base rates ({ns['currency']}, {pb} basis, annual per member) at Nil / 0% co-pay and the "
                f"reference options ({basis or 'see Factors'}), BEFORE the TPA factor. Fill the yellow cells from the "
                "cedant's catalogue; co-pay variants come from Factors.")
    write_rate_sheets(wb, f, lay, nspec, rate_ws, protect, subtitle=subtitle)

    lists = _Lists()
    lst_products = lists.add("products", [m["product_type"] for m in model["sheets"]])
    lst_emirates = lists.add("emirates", [m["rating_region"] or "" for m in model["sheets"]])
    lst_tpas = lists.add("tpas", [ALL_TPAS] + [t["tpas_label"] for m in model["sheets"] for t in m["factor_tables"] + m["riders"]]
                         + [g["label"] for m in model["sheets"] for g in m["groups"]])
    lst_ffam = lists.add("factor_families", [t["family"] for m in model["sheets"] for t in m["factor_tables"]])
    lst_rfam = lists.add("rider_families", [t["family"] for m in model["sheets"] for t in m["riders"]])
    lst_dfam = lists.add("default_families", [fam for m in model["sheets"] for fam in m["plan_defaults"]])
    lst_plans = lists.add("plans", [p["plan_id"] for m in model["sheets"] for g in m["groups"] for p in g["plans"]])
    lst_codes = lists.add("benefit_codes", sorted(reg.benefits))
    lst_y = lists.add("reference", ["Y"])
    opt_lists = _option_lists(model)
    spare_lists = {"Product": lst_products, "Emirate": lst_emirates, "TPAs": lst_tpas, "Plan": lst_plans,
                   "Benefit code": lst_codes, "Reference": lst_y}
    rows_by = long_rows(model)
    spare = int(model["options"]["spare_rows"])
    titles = {FACTORS: ("Factors - co-pay, deductible, limit, network option, nationality and TPA factors",
                        "Enter the factor (> 0). Reference = Y marks the option the base rates assume (factor 1.0)."),
              RIDERS: ("Riders - add-on module premiums", "Enter the rider rate (> 0) per limit option and key."),
              DEFAULTS: ("Plan defaults - the option each plan is sold with (SPEC §3.4)",
                         "Pick the option from the list; riders: NONE = not purchased."),
              BENEFITS: ("Benefits - table of benefits per plan (registry codes)",
                         "Value as the code needs; source wording in Condition text.")}
    counts = {}
    for sname in LONG_SHEETS:
        ws = ws_long[sname]
        cols = COLS[sname]
        ws.write(0, 0, titles[sname][0], f.title)
        ws.write(1, 0, titles[sname][1], f.sub)
        ws.write_row(HEADER_ROW, 0, cols, f.hdr)
        for j, c in enumerate(cols):
            ws.set_column(j, j, 34 if c in ("Condition text", "Eligibility", "Allowed") else (16 if c in ("Family", "Rider", "TPAs", "Option",
                                                                                                         "Limit", "Benefit code", "Key") else 12))
        ws.freeze_panes(HEADER_ROW + 1, 0)
        r = HEADER_ROW + 1
        rows = rows_by[sname]
        counts[sname] = len(rows)
        for row in rows:
            for j, c in enumerate(cols):
                if c in row["input"]:
                    kind, ck, keys, mand = row["input"][c]
                    numeric = c in NUMERIC.get(sname, [])
                    fmt = (nf["inp_rate"] if c == "Rate" else nf["inp_num"]) if numeric else (
                        nf["inp_gen"] if c == "Value" else nf["inp_txt"])
                    ws.write_blank(r, j, None, fmt)
                    lay.other_inputs.append(InputCell(cell_key=ck, kind=kind, sheet=sname, row=r, col=j, mandatory=mand,
                                                      keys=[keys], label=" / ".join(str(row["cols"].get(x, "")) for x in cols[:5])))
                    if numeric:
                        ws.data_validation(r, j, r, j, {"validate": "decimal", "criteria": ">", "value": 0,
                                                        "error_title": c, "error_message": f"Enter a {c.lower()} greater than 0."})
                    lst = (row.get("lists") or {}).get(c)
                    if lst and lst[0] == "opt":
                        src = lists.add("opt:" + "|".join(lst[1:]), opt_lists.get(lst[1:], []))
                        if src:
                            ws.data_validation(r, j, r, j, {"validate": "list", "source": src,
                                                            "error_title": "Option", "error_message": "Pick an option of the family's table."})
                    elif lst and lst[0] == "ben":
                        line = reg.benefits[lst[1]]
                        vals = line.get("ladder") if line["value_type"] == "ladder" else (
                            ["Y", "N"] if line["value_type"] == "bool" else None)
                        if vals:
                            src = lists.add("ben:" + lst[1], vals)
                            ws.data_validation(r, j, r, j, {"validate": "list", "source": src})
                else:
                    v = row["cols"].get(c, "")
                    ws.write_string(r, j, "" if v is None else str(v), nf["txt"])
            r += 1
        fam_list = {FACTORS: lst_ffam, RIDERS: lst_rfam, DEFAULTS: lst_dfam}.get(sname)
        if spare and (rows or sname == BENEFITS and model["sheets"]):
            ws.write(r, 0, f"Spare rows ({spare}): add an option, key or line to an existing table", f.sub)
            r += 1
            for _ in range(spare):
                for j, c in enumerate(cols):
                    if c in SPARE_EDITABLE[sname]:
                        numeric = c in NUMERIC.get(sname, [])
                        ws.write_blank(r, j, None, (nf["inp_rate"] if c == "Rate" else nf["inp_num"]) if numeric else nf["inp_gen"])
                        if numeric:
                            ws.data_validation(r, j, r, j, {"validate": "decimal", "criteria": ">", "value": 0})
                        src = fam_list if c in ("Family", "Rider") else spare_lists.get(c)
                        if src and not numeric:
                            ws.data_validation(r, j, r, j, {"validate": "list", "source": src})
                    else:
                        ws.write_blank(r, j, None, nf["ro_txt"])
                r += 1
        if not rows:
            ws.write(HEADER_ROW + 1, 0, "No rows in the template spec for this table.", f.sub)
        if protect:
            ws.protect("", {"select_locked_cells": True, "select_unlocked_cells": True, "format_columns": True,
                            "format_rows": True, "insert_rows": True, "autofilter": True, "sort": False})
    lists.write(ws_lists)

    shash = hash_keys_rows(lay.keys_rows(), [s["sheet"] for s in lay.rate_sheets], f"{nspec['tender_id']}|{MODE}")
    sphash = spec_hash(nspec)
    ws_cover.set_column(0, 0, 22)
    ws_cover.set_column(1, 1, 110)
    ws_cover.write(0, 0, f"Cedant rate workbook (current catalogue) - {nspec['cedant']} - {nspec['tender_id']}", f.title)
    lines = _cover_lines(nspec, model) + [("Template version", f"{TEMPLATE_SPEC_VERSION} / generator {GENERATOR_VERSION} / {MODE}"),
                                          ("Structure hash", shash)]
    for i, (k, v) in enumerate(lines):
        ws_cover.write(2 + i, 0, k, f.bold)
        ws_cover.write(2 + i, 1, v, f.wrap)
    r = 3 + len(lines)
    ws_cover.write(r, 0, "Sheet directory", f.bold)
    directory = [(s["sheet"], f"fill yellow cells ({len(s['blocks'])} block(s): " + ", ".join(bl["block"]["label"] for bl in s["blocks"]) + ")")
                 for s in lay.rate_sheets]
    directory += [(s, f"{counts[s]} row(s) to fill") for s in LONG_SHEETS]
    for i, (n, d) in enumerate(directory):
        ws_cover.write(r + 1 + i, 0, n)
        ws_cover.write(r + 1 + i, 1, d)
    if protect:
        ws_cover.protect()

    krows = lay.keys_rows()
    n_in = sum(1 for x in krows if x[1] != "REFER")
    n_ref = sum(1 for x in krows if x[1] == "REFER")
    meta = [("tender_id", nspec["tender_id"]), ("cedant", nspec["cedant"]), ("mode", MODE), ("structure_hash", shash),
            ("spec_hash", sphash), ("cedant_spec_hash", cedant_spec_hash(spec)), ("template_spec_version", nspec["template_spec_version"]),
            ("generator_version", GENERATOR_VERSION), ("issue_date", nspec.get("issue_date") or ""),
            ("currency", nspec["currency"]), ("premium_basis", nspec["premium_basis"]),
            ("n_inputs", str(n_in)), ("n_refer", str(n_ref))]
    write_keys_meta(ws_keys, ws_meta, krows, meta, protect)
    if protect:
        ws_lists.protect()
    wb.close()
    counts["RATE"] = sum(1 for c in lay.cells if not c.refer)
    return GeneratedCedantWorkbook(path=out_path, structure_hash=shash, spec_hash=sphash, n_inputs=n_in, n_refer=n_ref,
                                   counts=counts, layout=lay)


# ============================================================================================== read back
@dataclass
class CedantReadReport:
    problems: list = field(default_factory=list)      # [{sheet, cell, check, message, value}]
    warnings: list = field(default_factory=list)
    counts: dict = field(default_factory=dict)
    hash_ok: bool | None = None
    structure_hash: str | None = None

    @property
    def status(self) -> str:
        return "PASS" if not self.problems else "FAIL"

    def add(self, sheet, cell, check, message, value=None):
        self.problems.append({"sheet": sheet, "cell": cell, "check": check, "message": message,
                              "value": "" if value is None else str(value)})

    def to_frame(self):
        import pandas as pd
        rows = [{"level": "PROBLEM", **p} for p in self.problems] + \
               [{"level": "WARNING", "sheet": "", "cell": "", "check": "warning", "message": w, "value": ""} for w in self.warnings]
        return pd.DataFrame(rows, columns=["level", "sheet", "cell", "check", "message", "value"])

    def summary(self) -> str:
        return (f"CEDANT READ {self.status}: {len(self.problems)} problem(s), {len(self.warnings)} warning(s); "
                + ", ".join(f"{k} {v}" for k, v in sorted(self.counts.items())))


def _num(v):
    """(float | None, problem text | None, warning text | None). None value = blank."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None, None
    if isinstance(v, bool):
        return None, "not a number", None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None, "not a number", None
        return float(v), None, None
    s = str(v).strip().replace(",", "")
    pct = s.endswith("%")
    try:
        x = float(s.rstrip("%")) / (100.0 if pct else 1.0)
    except ValueError:
        return None, "not a number", None
    return x, None, f"text '{v}' read as the number {x:g}"


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _scan(ws, cols: list[str]):
    """Locate the long-table header (tolerates rows inserted above it) and yield (row index 1-based, {col: value})."""
    head_r, idx = None, None
    for r in range(1, min(ws.max_row, 60) + 1):
        vals = [_txt(ws.cell(r, c).value) for c in range(1, len(cols) + 12)]
        if all(c in vals for c in cols):
            head_r, idx = r, {c: vals.index(c) + 1 for c in cols}
            break
    if head_r is None:
        return None, []
    out = []
    for r in range(head_r + 1, ws.max_row + 1):
        d = {c: ws.cell(r, idx[c]).value for c in cols}
        if all(_txt(v) == "" for v in d.values()):
            continue
        first = _txt(d[cols[0]])
        if first.startswith("Spare rows") or first.startswith("No rows in the template spec"):
            continue
        out.append((r, d))
    return idx, out


def _benefit_value(v, line: dict):
    """(value, problem). Value types from registry/benefits.json; ladder values must be listed."""
    vt = line["value_type"]
    ladder = line.get("ladder") or []

    def as_bool(x):
        if isinstance(x, bool):
            return x, None
        s = _txt(x).upper()
        if s in BOOL_TEXT:
            return BOOL_TEXT[s], None
        return None, f"'{x}' is not Y / N"

    def as_ladder(x):
        s = _txt(x)
        hit = [lv for lv in ladder if lv.upper() == s.upper()]
        if hit:
            return hit[0], None
        return None, f"'{x}' is not a ladder value of {line['code']} (allowed: {', '.join(ladder)})"

    def as_num(x, integral=False):
        n, prob, _ = _num(x)
        if prob:
            return None, f"'{x}' is not a number"
        if n < 0:
            return None, f"'{x}' is negative"
        if integral:
            if abs(n - round(n)) > 1e-9:
                return None, f"'{x}' is not a whole number"
            return int(round(n)), None
        return (int(n) if float(n).is_integer() and vt == "money" else n), None
    if vt in ("money", "pct", "copay_structure"):
        return as_num(v)
    if vt in ("int", "tier_rank"):
        return as_num(v, integral=True)
    if vt == "bool":
        return as_bool(v)
    if vt == "ladder":
        return as_ladder(v)
    if vt == "money_or_bool":
        n, p = as_num(v)
        return (n, None) if p is None else as_bool(v)
    if vt == "money_or_ladder":
        n, p = as_num(v)
        return (n, None) if p is None else as_ladder(v)
    return None, f"value type '{vt}' has no read rule"


def read_cedant(path: Path, spec: dict, catalogue_id: str | None = None, owner: str | None = None,
                version: str | None = None, effective_date: str | None = None):
    """Read a filled cedant workbook. Returns (Catalogue | None, CedantReadReport); the catalogue is None when the
    gate found any problem (SPEC §8.5: the read stops with the full cell list, never defaults)."""
    path = Path(path)
    reg = load_registry()
    nspec = normalise(spec)
    model, errs = cedant_model(spec)
    if errs:
        raise SpecError("template_spec cedant sections invalid:\n  " + "\n  ".join(errs))
    rep = CedantReadReport()
    lay = compute_layout(nspec)
    meta = read_meta(path)
    if meta and meta.get("mode") != MODE:
        rep.add(META_SHEET, "", "mode", f"workbook mode is '{meta.get('mode') or 'TENDER'}', not {MODE} (is this a "
                                         "reinsurer template?)")
    elif meta and meta.get("cedant_spec_hash") != cedant_spec_hash(spec):
        rep.add(META_SHEET, "", "spec", "the workbook was generated from a different template spec (plans, networks, "
                                        "tables or options changed): read it with the spec it was issued from, or "
                                        "regenerate and re-enter")
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    try:
        return _read(path, wb_f, wb_v, nspec, model, lay, meta, rep, reg, catalogue_id, owner, version, effective_date)
    finally:
        wb_f.close()
        wb_v.close()


def _read(path, wb_f, wb_v, nspec, model, lay, meta, rep, reg, catalogue_id, owner, version, effective_date):
    names = wb_v.sheetnames
    located = {}                                      # cell_key -> (sheet, address)
    sheet_alias = {}                                  # long-table sheet -> actual name (renames followed via _KEYS)
    if KEYS_SHEET in wb_f.sheetnames:
        krows = _keys_table(wb_f)
        recomputed = hash_keys_rows([[r["cell_key"], r["kind"], bool(r["mandatory"]), r["key_json"], r["issued_sheet"],
                                      r["issued_address"]] for r in krows],
                                    [s["sheet"] for s in lay.rate_sheets], f"{nspec['tender_id']}|{MODE}")
        rep.structure_hash = recomputed
        rep.hash_ok = bool(meta.get("structure_hash")) and meta.get("structure_hash") == recomputed
        moved = []
        for r in krows:
            if r["kind"] == "REFER":
                continue
            sheet, addr = _locate(r, names, moved)
            located[r["cell_key"]] = (sheet, addr)
            if sheet and r["issued_sheet"] in LONG_SHEETS:
                sheet_alias.setdefault(r["issued_sheet"], sheet)
        rep.warnings += moved[:50] + ([f"... {len(moved) - 50} more moved cells"] if len(moved) > 50 else [])
        if not rep.hash_ok:
            rep.warnings.append("structure hash differs from _META (the workbook's _KEYS were changed or it was issued "
                                "from another spec)")
    else:
        rep.hash_ok = False
        rep.warnings.append(f"{KEYS_SHEET} missing: rate cells read by position from the template_spec")
        for c in lay.cells:
            if not c.refer:
                located[c.cell_key] = (c.sheet if c.sheet in names else None, c.address)

    # ---- rate sheets
    base_rows = {}                                    # (sheet, tpa) -> rows
    n_rates = 0
    sheet_mand = {s["sheet"]: bool(s["spec"].get("mandatory", True)) for s in lay.rate_sheets}
    for c in lay.cells:
        tpa = c.keys[0].get("tpa")
        if c.refer:
            vals = None
        else:
            sheet, addr = located.get(c.cell_key, (None, None))
            if sheet is None:
                rep.add(c.sheet, c.address, "rate", f"rate cell {c.label} not found (sheet missing or key lost)")
                continue
            raw = wb_v[sheet][addr].value
            v, prob, warn = _num(raw)
            if warn:
                rep.warnings.append(f"{sheet}!{addr}: {warn}")
            if prob:
                rep.add(sheet, addr, "rate", f"rate {c.label}: {prob}", raw)
                continue
            if v is None:
                if sheet_mand.get(c.sheet, True):
                    rep.add(sheet, addr, "rate", f"blank mandatory rate {c.label}")
                continue
            if v <= 0:
                rep.add(sheet, addr, "rate", f"rate {c.label} must be > 0", raw)
                continue
            vals = v
            n_rates += 1
        for k in c.keys:
            if isinstance(k.get("age"), str):
                rep.add(c.sheet, c.address, "rate", f"placeholder age band '{k['age']}' has no range (SPEC §8.4)")
                continue
            key = {d: v2 for d, v2 in k.items() if d not in ("product_type", "regulator", "rating_region")}
            base_rows.setdefault((c.sheet, tpa), []).append(
                {"key": key, "value": vals, "status": "REFER" if c.refer else "VALUE", "labels": {},
                 "source_ref": f"{located.get(c.cell_key, (c.sheet, c.address))[0]}!{located.get(c.cell_key, (c.sheet, c.address))[1]}"
                 if not c.refer else f"{c.sheet}!{c.address} (refer)"})
    rep.counts["rate_cells"] = n_rates

    # ---- long tables
    by_pe = {(m["product_type"], m["rating_region"] or ""): m for m in model["sheets"]}

    def sheet_of(sname):
        nm = sheet_alias.get(sname, sname)
        return nm if nm in names else None
    tables = {}            # (sheet_name, tpas_label, family) -> {"model": t, "rows": [...], "kind"}
    for m in model["sheets"]:
        for t in m["factor_tables"]:
            tables[(m["sheet_name"], t["tpas_label"], t["family"], FACTORS)] = {"model": t, "rows": [], "msh": m}
        for t in m["riders"]:
            tables[(m["sheet_name"], t["tpas_label"], t["family"], RIDERS)] = {"model": t, "rows": [], "msh": m}

    def resolve(sname, r, d, idx, fam_col, canon=None):
        prod, emi = _txt(d["Product"]), _txt(d["Emirate"])
        msh = by_pe.get((prod, emi))
        cell = lambda c: a1(r, idx[c])  # noqa: E731
        if msh is None:
            rep.add(sname, cell("Product"), "unknown", f"no rate sheet for product '{prod}' x emirate '{emi}'")
            return None, None
        tl = _txt(d["TPAs"]) or ALL_TPAS
        if fam_col is None:
            g = next((g for g in msh["groups"] if g["label"] == tl), None)
            if g is None:
                rep.add(sname, cell("TPAs"), "unknown", f"TPAs '{tl}' is not a TPA group of {msh['sheet_name']} "
                                                       f"(groups: {', '.join(x['label'] for x in msh['groups'])})")
            return msh, g
        fam = _txt(d[fam_col])
        tb = tables.get((msh["sheet_name"], tl, fam, canon))
        if tb is None:
            what = "factor" if canon == FACTORS else "rider"
            rep.add(sname, cell(fam_col), "unknown", f"no {what} table of family '{fam}' for TPAs '{tl}' on "
                                                     f"{msh['sheet_name']} (spare rows extend existing tables only)")
        return msh, tb

    for sname, fam_col in ((FACTORS, "Family"), (RIDERS, "Rider")):
        nm = sheet_of(sname)
        if nm is None:
            if any(k[3] == sname for k in tables):
                rep.add(sname, "", "sheet", f"sheet '{sname}' is missing")
            continue
        idx, rows = _scan(wb_v[nm], COLS[sname])
        if idx is None:
            rep.add(nm, "", "sheet", f"header row of '{sname}' not found")
            continue
        seen, present = {}, set()
        for r, d in rows:
            msh, tb = resolve(nm, r, d, idx, fam_col, canon=sname)
            if tb is None:
                continue
            t = tb["model"]
            cell = lambda c: a1(r, idx[c])  # noqa: E731
            key = _txt(d["Key"]) or ANY
            opt = _txt(d["Option" if sname == FACTORS else "Limit"])
            if not opt:
                rep.add(nm, cell("Option" if sname == FACTORS else "Limit"), "option", "blank option")
                continue
            if sname == RIDERS and opt == NONE_OPTION:
                rep.add(nm, cell("Limit"), "option", f"option '{NONE_OPTION}' is reserved (rider not purchased)")
                continue
            present.add((msh["sheet_name"], t["tpas_label"], t["family"], key, opt))
            ok = True
            for dname, val, col in ((t["option_dim"], opt, "Option" if sname == FACTORS else "Limit"), (t["key_dim"], key, "Key")):
                if dname is None:
                    if val != ANY:
                        rep.add(nm, cell(col), "key", f"table {t['component_id']} has no key; leave Key = ANY", val)
                        ok = False
                    continue
                try:
                    reg.check_value(dname, val, allow_any=(col == "Key"), allow_parent=True)
                except Exception as e:  # noqa: BLE001
                    rep.add(nm, cell(col), "value", str(e), val)
                    ok = False
            if t["key_dim"] == "plan" and key != ANY and key not in {p["plan_id"] for g in msh["groups"] for p in g["plans"]}:
                rep.add(nm, cell("Key"), "key", f"plan '{key}' is not a plan of {msh['sheet_name']}", key)
                ok = False
            if t["key_dim"] == "tpa" and key != ANY and key not in t["tpas"]:
                rep.add(nm, cell("Key"), "key", f"TPA '{key}' is outside the table's TPAs", key)
                ok = False
            vcol = "Factor" if sname == FACTORS else "Rate"
            raw = d[vcol]
            v, prob, warn = _num(raw)
            if warn:
                rep.warnings.append(f"{nm}!{cell(vcol)}: {warn}")
            mand = True if sname == FACTORS else t["mandatory"]
            if prob:
                rep.add(nm, cell(vcol), vcol.lower(), f"{t['family']} {key}/{opt}: {prob}", raw)
                ok = False
            elif v is None:
                if mand:
                    rep.add(nm, cell(vcol), vcol.lower(), f"blank {vcol.lower()} for {t['family']} {key}/{opt}")
                    ok = False
                else:
                    continue
            elif v <= 0:
                rep.add(nm, cell(vcol), vcol.lower(), f"{vcol.lower()} for {t['family']} {key}/{opt} must be > 0", raw)
                ok = False
            refv = False
            if sname == FACTORS:
                rr = _txt(d["Reference"]).upper()
                if rr and rr not in YES and rr not in ("N", "NO", "FALSE"):
                    rep.add(nm, cell("Reference"), "reference", f"Reference must be Y or blank, got '{d['Reference']}'")
                    ok = False
                refv = rr in YES
            dk = (msh["sheet_name"], t["tpas_label"], t["family"], key, opt)
            if dk in seen:
                rep.add(nm, cell(vcol), "duplicate", f"duplicate row {t['family']} key={key} option={opt} (first at {seen[dk]})")
                continue
            seen[dk] = f"{nm}!{cell(vcol)}"
            if ok:
                tb["rows"].append({"key": key, "option": opt, "value": v, "reference": refv, "ref": f"{nm}!{cell(vcol)}"})
            else:
                tb.setdefault("bad", True)
        # every pre-filled (spec) row must still be there: a deleted row is a missing mandatory input
        for (shn, tl, fam, kind), tb in tables.items():
            if kind != sname or (kind == RIDERS and not tb["model"]["mandatory"]):
                continue
            for sr in tb["model"]["rows"]:
                if (shn, tl, fam, sr["key"], sr["option"]) not in present:
                    rep.add(nm, "", "missing", f"{shn} / {fam} ({tl}): row key={sr['key']} option={sr['option']} is missing "
                                               "(deleted or relabelled)")
                    tb["bad"] = True
        # references (factor tables): exactly one per key, ANY inherited; reference factor 1.0
        if sname == FACTORS:
            for (shn, tl, fam, kind), tb in tables.items():
                if kind != FACTORS or tb.get("bad"):
                    continue
                keys_present = list(dict.fromkeys(r["key"] for r in tb["rows"]))
                for k in keys_present:
                    refs = [r for r in tb["rows"] if r["key"] == k and r["reference"]]
                    inh = k != ANY and any(r["key"] == ANY and r["reference"] for r in tb["rows"])
                    if len(refs) > 1:
                        rep.add(nm, ", ".join(r["ref"].split("!")[1] for r in refs), "reference",
                                f"{shn} / {fam} ({tl}): key '{k}' has {len(refs)} reference rows (exactly one)")
                    elif not refs and not inh:
                        rep.add(nm, "", "reference", f"{shn} / {fam} ({tl}): key '{k}' has no reference row (Reference = Y)")
                    for r in refs:
                        if abs(r["value"] - 1.0) > _EPS:
                            rep.add(nm, r["ref"].split("!")[1], "reference",
                                    f"{shn} / {fam} ({tl}): reference row {k}/{r['option']} must read 1.0", r["value"])
                if not tb["rows"]:
                    rep.add(nm, "", "missing", f"{shn} / {fam} ({tl}): no rows")
        else:
            for (shn, tl, fam, kind), tb in tables.items():
                if kind == RIDERS and not tb["rows"] and not tb.get("bad") and tb["model"]["mandatory"]:
                    rep.add(nm, "", "missing", f"{shn} / rider {fam} ({tl}): no rows")
    rep.counts["factor_rows"] = sum(len(t["rows"]) for k, t in tables.items() if k[3] == FACTORS)
    rep.counts["rider_rows"] = sum(len(t["rows"]) for k, t in tables.items() if k[3] == RIDERS)

    # ---- plan defaults
    defaults = {}          # (sheet, group label, plan, family) -> option
    nm = sheet_of(DEFAULTS)
    want_defaults = {(m["sheet_name"], g["label"], p["plan_id"], fam) for m in model["sheets"] for g in m["groups"]
                     for p in g["plans"] for fam in g["default_families"]}
    if nm is None:
        if want_defaults:
            rep.add(DEFAULTS, "", "sheet", f"sheet '{DEFAULTS}' is missing")
    else:
        idx, rows = _scan(wb_v[nm], COLS[DEFAULTS])
        if idx is None:
            rep.add(nm, "", "sheet", f"header row of '{DEFAULTS}' not found")
            rows = []
        seen = {}
        for r, d in rows:
            msh, g = resolve(nm, r, d, idx, None)
            if g is None:
                continue
            cell = lambda c: a1(r, idx[c])  # noqa: E731
            plan, fam, opt = _txt(d["Plan"]), _txt(d["Family"]), _txt(d["Option"])
            if plan not in {p["plan_id"] for p in g["plans"]}:
                rep.add(nm, cell("Plan"), "unknown", f"plan '{plan}' is not a plan of {msh['sheet_name']} ({g['label']})")
                continue
            if fam not in g["default_families"]:
                rep.add(nm, cell("Family"), "unknown", f"family '{fam}' is not a plan-default family of {msh['sheet_name']} "
                                                       f"({g['label']})")
                continue
            dk = (msh["sheet_name"], g["label"], plan, fam)
            if dk in seen:
                rep.add(nm, cell("Option"), "duplicate", f"duplicate plan default {plan} / {fam} (first at {seen[dk]})")
                continue
            seen[dk] = f"{nm}!{cell('Option')}"
            if not opt:
                rep.add(nm, cell("Option"), "default", f"blank plan default {plan} / {fam}")
                continue
            opts, rider, missing = set(), False, []
            pl = next(p for p in g["plans"] if p["plan_id"] == plan)
            for (shn, tl, f2, kind), tb in tables.items():
                if shn == msh["sheet_name"] and f2 == fam and set(g["tpas"]) <= set(tb["model"]["tpas"]):
                    rider = rider or kind == RIDERS
                    # the option must exist for every key the plan's cells take (plan / its networks / its TPAs), or ANY
                    for kv in _plan_keys(tb["model"], pl, g):
                        here = {x["option"] for x in tb["rows"] if x["key"] in (kv, ANY)}
                        opts |= here
                        if opt not in here and not (kind == RIDERS and opt == NONE_OPTION):
                            missing.append(kv)
            if rider and opt == NONE_OPTION:
                defaults[dk] = opt
            elif missing or opt not in opts:
                where = f" for key {', '.join(str(m) for m in missing if m is not None)}" if any(m is not None for m in missing) else ""
                rep.add(nm, cell("Option"), "default", f"plan default {plan} / {fam} = '{opt}' is not an option of its table"
                                                       f"{where} ({', '.join(sorted(opts)) or 'no rows'})", opt)
            else:
                defaults[dk] = opt
        for dk in sorted(want_defaults - set(seen)):
            rep.add(nm, "", "missing", f"no plan-default row for {dk[0]} / {dk[1]} / plan {dk[2]} / {dk[3]}")
    rep.counts["plan_defaults"] = len(defaults)

    # ---- benefits
    benefits = {}          # (sheet, group label) -> [BenefitLine]
    nm = sheet_of(BENEFITS)
    n_blank = 0
    if nm is None and any(m["benefits"] for m in model["sheets"]):
        rep.add(BENEFITS, "", "sheet", f"sheet '{BENEFITS}' is missing")
    if nm is not None:
        idx, rows = _scan(wb_v[nm], COLS[BENEFITS])
        if idx is None:
            rep.add(nm, "", "sheet", f"header row of '{BENEFITS}' not found")
            rows = []
        seen = {}
        for r, d in rows:
            raw, text = d["Value"], _txt(d["Condition text"])
            code = _txt(d["Benefit code"])
            if (raw is None or _txt(raw) == "") and not text and _txt(d["Value type"]):
                n_blank += 1                            # pre-filled line left blank: the benefit does not apply
                continue
            msh, g = resolve(nm, r, d, idx, None)
            if g is None:
                continue
            cell = lambda c: a1(r, idx[c])  # noqa: E731
            plan = _txt(d["Plan"])
            if plan not in {p["plan_id"] for p in g["plans"]}:
                rep.add(nm, cell("Plan"), "unknown", f"plan '{plan}' is not a plan of {msh['sheet_name']} ({g['label']})")
                continue
            if code not in reg.benefits:
                rep.add(nm, cell("Benefit code"), "benefit", f"unknown benefit code '{code}' (registry/benefits.json)", code)
                continue
            dk = (msh["sheet_name"], g["label"], plan, code)
            if dk in seen:
                rep.add(nm, cell("Value"), "duplicate", f"duplicate benefit {plan} / {code} (first at {seen[dk]})")
                continue
            seen[dk] = f"{nm}!{cell('Value')}"
            if raw is None or _txt(raw) == "":
                if not text:
                    rep.add(nm, cell("Value"), "benefit", f"{plan} / {code}: added row has neither a value nor condition text")
                    continue
                val, prob = None, None
            else:
                val, prob = _benefit_value(raw, reg.benefits[code])
            if prob:
                rep.add(nm, cell("Value"), "benefit", f"{plan} / {code}: {prob}", raw)
                continue
            benefits.setdefault((msh["sheet_name"], g["label"]), []).append(
                BenefitLine(plan_id=plan, benefit_code=code, value=val, structure=None, condition_text=text or None,
                            source_ref=f"{nm}!{cell('Value')}", status="PARSED" if val is not None else "MANUAL_REVIEW"))
    rep.counts["benefit_lines"] = sum(len(v) for v in benefits.values())
    if n_blank:
        rep.warnings.append(f"{n_blank} benefit row(s) left blank (benefit not applicable or not entered)")

    if rep.problems:
        return None, rep

    # ---- catalogue
    cw = model["options"]
    cat = Catalogue(catalogue_id=catalogue_id or cw.get("catalogue_id") or f"{nspec['cedant']}_EXPIRING",
                    owner=owner or nspec["cedant"], kind="EXPIRING", version=str(version or cw.get("version") or "1"),
                    effective_date=effective_date or cw.get("effective_date"), currency=nspec["currency"],
                    premium_basis=nspec["premium_basis"], source_file=path.name, spec_id=nspec["tender_id"],
                    age_band_maps={n: [list(b["range"]) for b in bl if b.get("range")]
                                   for n, bl in sorted((nspec.get("age_band_sets") or {}).items())},
                    meta={"mode": MODE, "structure_hash": rep.structure_hash, "hash_ok": rep.hash_ok,
                          "spec_hash": spec_hash(nspec), "read_counts": dict(rep.counts), "warnings": rep.warnings[:200],
                          "tpa_factor_basis": "PRE", "not_carried": list(cw.get("not_carried") or [])})
    order = reg.order
    for msh in model["sheets"]:
        for g in msh["groups"]:
            rows = []
            for t in g["tpas"]:
                rows += base_rows.get((msh["sheet_name"], t), [])
            dims = reg.sort_dims(sorted({d for r in rows for d in r["key"]}))
            for r in rows:
                r["key"] = {d: r["key"].get(d, ANY) for d in dims}
            rows.sort(key=lambda r: canonical_key_str(r["key"], order))
            fopts = msh.get("formula") or {}
            scope = RatingScope(scope_id=g["scope_id"], product_type=msh["product_type"], regulator=msh["regulator"],
                                rating_region=msh["rating_region"], tpa_set=list(g["tpas"]),
                                formula=RatingFormula(order=list(fopts.get("order") or ["BASE", "FACTORS", "MODULES"]),
                                                      rounding=int(fopts.get("rounding", nspec["options"]["rounding"])),
                                                      min_premium=fopts.get("min_premium"),
                                                      max_premium=fopts.get("max_premium"),
                                                      refer_rules=list(fopts.get("refer_rules") or [])))
            for p in g["plans"]:
                pdef = {fam: defaults[(msh["sheet_name"], g["label"], p["plan_id"], fam)] for fam in g["default_families"]
                        if (msh["sheet_name"], g["label"], p["plan_id"], fam) in defaults}
                scope.plans.append(Plan(plan_id=p["plan_id"], label=p["label"], tier_rank=p["tier_rank"],
                                        network_by_tpa=dict(p["network_by_tpa"]), attributes={}, defaults=pdef))
            scope.components.append(Component(component_id="base", type="BASE_TABLE", keys=dims, rows=rows,
                                              premium_basis=msh.get("premium_basis")))
            for cid in g["factor_tables"]:
                t = next(x for x in msh["factor_tables"] if x["component_id"] == cid)
                tb = tables[(msh["sheet_name"], t["tpas_label"], t["family"], FACTORS)]
                scope.components.append(_factor_component(t, tb["rows"], g))
            for cid in g["riders"]:
                t = next(x for x in msh["riders"] if x["component_id"] == cid)
                tb = tables[(msh["sheet_name"], t["tpas_label"], t["family"], RIDERS)]
                scope.components.append(Component(
                    component_id=t["component_id"], type="MODULE", family=t["family"], keys=list(t["keys"]),
                    rows=[{"key": _row_key(t, r), "value": r["value"], "status": "VALUE", "labels": {}, "source_ref": r["ref"]}
                          for r in tb["rows"]],
                    unit=t["unit"], mode=t["mode"], eligibility=dict(t["eligibility"]), premium_basis=t.get("premium_basis"),
                    default_option=t.get("default_option")))
            scope.benefits = list(benefits.get((msh["sheet_name"], g["label"]), []))
            cat.scopes.append(scope)
    cat.validate()
    return cat, rep


def _row_key(t: dict, r: dict) -> dict:
    vals = {t["option_dim"]: r["option"]}
    if t["key_dim"]:
        vals[t["key_dim"]] = r["key"]
    return {k: vals[k] for k in t["keys"]}


def _plan_keys(t: dict, plan: dict, g: dict) -> list:
    """Key values a plan's cells take in a table keyed on plan / network / tpa; [None] = keyless or other key."""
    kd = t["key_dim"]
    if kd == "plan":
        return [plan["plan_id"]]
    if kd == "network":
        return sorted({n for tp, n in plan["network_by_tpa"].items() if tp in g["tpas"] and n is not None}) or [None]
    if kd == "tpa":
        return [tp for tp in g["tpas"] if tp in plan["network_by_tpa"]] or list(g["tpas"])
    return [None]


def _factor_component(t: dict, rows: list, g: dict) -> Component:
    comp = Component(component_id=t["component_id"], type="FACTOR_TABLE", family=t["family"], keys=list(t["keys"]),
                     rows=[{"key": _row_key(t, r), "value": r["value"], "status": "VALUE", "labels": {}, "source_ref": r["ref"]}
                           for r in rows],
                     default_option=t.get("default_option"))
    own = {r["key"]: r["option"] for r in rows if r["reference"]}
    # every key a cell can take gets its effective reference (own row, else the ANY row's): the evaluator looks the
    # key up literally (fill = reference), so an inherited reference is written out per key
    domain = {r["key"] for r in rows if r["key"] != ANY}
    for p in g["plans"]:
        domain |= {k for k in _plan_keys(t, p, g) if k is not None}
    eff = {k: own.get(k, own.get(ANY)) for k in sorted(domain)}
    vals = set(eff.values()) | ({own[ANY]} if ANY in own else set())
    if len(vals) == 1 and None not in vals:
        comp.reference_level = next(iter(vals))
    else:
        comp.reference_level_by = t["key_dim"]
        comp.reference_level_by_key = {k: v for k, v in eff.items() if v is not None}
    return comp


def read_cedant_workbook(path: Path, spec: dict, **kw) -> Catalogue:
    """Read back into the EXPIRING catalogue; raises GateFailed (details = the report) on any problem."""
    cat, rep = read_cedant(path, spec, **kw)
    if cat is None:
        first = "; ".join(f"{p['sheet']}!{p['cell']} {p['message']}" for p in rep.problems[:5])
        raise GateFailed("CEDANT_READ", f"{len(rep.problems)} problem(s) in {Path(path).name}: {first}", details=rep)
    return cat


# ============================================================================================== fill helper
def fill_from_catalogue(template_path: Path, out_path: Path, cat: Catalogue, value_fn=None) -> dict:
    """Fill a generated cedant workbook from an existing catalogue (fixtures, migration of a mapped catalogue): rate,
    factor, rider, plan-default and benefit cells by their _KEYS keys. value_fn(cell_key, keys, kind, value) may
    alter or drop (None) a value. Returns {cell_key: value written}."""
    from ..evaluator import lookup_component
    reg = load_registry()

    def scope_for(k):
        tpas = k.get("tpas") or [k.get("tpa")]
        for s in cat.scopes:
            if s.product_type == k.get("product_type") and s.rating_region == k.get("rating_region") and \
                    (tpas[0] in s.tpa_set or ANY in s.tpa_set):
                return s
        return None
    base_cache = {}

    def base_value(k):
        s = scope_for(k)
        if s is None:
            return None
        base = next(c for c in s.components if c.type == "BASE_TABLE")
        probe = {d: (v[0] if isinstance(v, list) else v) for d, v in k.items() if d in base.keys}
        ck = (s.scope_id, json.dumps(probe, sort_keys=True))
        if ck not in base_cache:
            res = lookup_component(cat, s.scope_id, base, [probe], reg)
            base_cache[ck] = None if res.empty or res["__status"].iloc[0] != "VALUE" else float(res["__value"].iloc[0])
        return base_cache[ck]

    def value(ck, keys, kind):
        k = keys[0]
        if kind == "RATE":
            return base_value(k)
        s = scope_for(k)
        if s is None:
            return None
        if kind in ("FACTOR", "RIDER"):
            comp = next((c for c in s.components if c.component_id == k["component_id"]), None)
            if comp is None:
                return None
            for r in comp.rows:
                kd = r["key"]
                opt_dim = comp.family if comp.family in comp.keys else comp.keys[0]
                other = [d for d in comp.keys if d != opt_dim]
                if str(kd[opt_dim]) == k["option"] and (str(kd[other[0]]) if other else ANY) == k["key"]:
                    return None if r.get("value") is None else float(r["value"])
            return None
        if kind == "DEFAULT":
            try:
                p = s.plan(k["plan"])
            except SpecError:
                return None
            return p.defaults.get(k["family"], s.defaults.get(k["family"]))
        if kind in ("BENEFIT", "BENEFIT_TEXT"):
            line = next((b for b in s.benefits if b.plan_id == k["plan"] and b.benefit_code == k["code"]), None)
            if line is None:
                return None
            if kind == "BENEFIT_TEXT":
                return line.condition_text
            v = line.value
            if isinstance(v, bool):
                return "Y" if v else "N"
            return v
        return None

    def fn(ck, keys, kind):
        v = value(ck, keys, kind)
        if value_fn is not None:
            v = value_fn(ck, keys, kind, v)
        return v
    return fill_template(template_path, out_path, value_fn=fn)
