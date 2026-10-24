"""Excel I/O: declarative grid reader (openpyxl) and deterministic workbook writer (xlsxwriter).

Never round-trip a generated workbook through LibreOffice (CLAUDE.md).
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from . import MappingRequired, SpecError

FIXED_CREATED = _dt.datetime(2026, 1, 1, 0, 0, 0)


def a1(row: int, col: int, absolute: bool = False) -> str:
    """1-based row/col -> A1."""
    c = get_column_letter(col)
    return f"${c}${row}" if absolute else f"{c}{row}"


def parse_a1(ref: str):
    col, row = coordinate_from_string(ref.replace("$", ""))
    return row, column_index_from_string(col)


def quote_sheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


class WorkbookReader:
    """Thin cached wrapper: values only (data_only=False keeps formulas as text, needed for _KEYS refs)."""

    def __init__(self, path: Path, data_only: bool = True):
        self.path = Path(path)
        self.wb = load_workbook(self.path, data_only=data_only, read_only=False)

    def sheet(self, name: str):
        if name not in self.wb.sheetnames:
            raise SpecError(f"{self.path.name}: sheet '{name}' not found (sheets: {self.wb.sheetnames})")
        return self.wb[name]

    def value(self, sheet: str, ref: str):
        return self.sheet(sheet)[ref.replace("$", "")].value


def _map_label(label, entry: dict, maps: dict, plans_by_label: dict, band_maps: dict, where: str):
    """Map one header/row label to a list of partial keys (expansion) plus the kept label."""
    if label is None or (isinstance(label, str) and not label.strip()):
        raise SpecError(f"{where}: empty label")
    lab = str(label).strip() if not isinstance(label, float) else (str(int(label)) if label.is_integer() else str(label))
    dims = entry.get("dims") or [entry["dim"]]
    if entry.get("band_map"):
        bm = band_maps.get(entry["band_map"])
        if bm is None:
            raise SpecError(f"{where}: unknown band_map '{entry['band_map']}'")
        lo, hi = bm.range_of(lab)
        return [{dims[0]: [lo, hi]}], lab
    m = entry.get("map")
    if m is None and entry.get("labels") is not None:
        table = entry["labels"]
    elif m == "@plans":
        table = plans_by_label
    elif m is None:
        if entry.get("as_is"):
            return [{dims[0]: lab}], lab
        raise SpecError(f"{where}: dimension '{dims[0]}' needs a map, labels, band_map or as_is")
    else:
        if m not in maps:
            raise SpecError(f"{where}: unknown map '{m}'")
        table = maps[m]
    if lab not in table:
        raise MappingRequired(f"{where}: label '{lab}' has no mapping for {dims}", [{"where": where, "label": lab}])
    v = table[lab]
    if isinstance(v, list):
        return [dict(x) for x in v], lab
    if isinstance(v, dict):
        return [dict(v)], lab
    return [{dims[0]: v}], lab


def _cross(parts: list[list[dict]]) -> list[dict]:
    out = [{}]
    for alts in parts:
        out = [{**o, **a} for o in out for a in alts]
    return out


def read_grid(reader: WorkbookReader, src: dict, maps: dict, plans_by_label: dict, band_maps: dict,
              text_values: dict | None = None):
    """Read a declared grid. Returns list of dicts:
    {"keys": [partial key, ...], "labels": {dim: label}, "raw": value, "ref": "Sheet!B5", "status": VALUE|REFER}
    """
    sheet_name = src["sheet"]
    ws = reader.sheet(sheet_name)
    r0, c0 = parse_a1(src["anchor"])
    row_dims = src.get("row_dims", [])
    col_dims = src.get("col_dims", [])
    nr_h = len(col_dims)
    nc_h = len(row_dims)
    data_r0 = r0 + nr_h
    data_c0 = c0 + nc_h
    tv = dict(text_values or {})
    tv.update(src.get("text_values", {}))
    where = f"{reader.path.name}:{sheet_name}"

    n_cols = src.get("n_cols")
    if n_cols is None:
        n_cols = 0
        while ws.cell(row=r0 + max(nr_h - 1, 0), column=data_c0 + n_cols).value not in (None, ""):
            n_cols += 1
        if nr_h == 0:
            n_cols = 1
    row_keys = src.get("row_keys")            # explicit per-row partial keys (labels not read from the sheet)
    if row_keys is not None:
        row_dims = []
        nc_h = 0
        data_c0 = c0
    scale = float(src.get("value_scale", 1.0))  # e.g. 12 for monthly rates annualised
    n_rows = len(row_keys) if row_keys is not None else src.get("n_rows")
    if n_rows is None:
        n_rows = 0
        probe_c = c0 + max(nc_h - 1, 0) if nc_h else data_c0
        while ws.cell(row=data_r0 + n_rows, column=probe_c).value not in (None, ""):
            n_rows += 1

    # column headers (forward-filled to the right: merged cells read blank)
    col_keys = []
    last = [None] * nr_h
    for j in range(n_cols):
        parts, labels = [], {}
        for i, entry in enumerate(col_dims):
            v = ws.cell(row=r0 + i, column=data_c0 + j).value
            if v in (None, ""):
                v = last[i]
            last[i] = v
            k, lab = _map_label(v, entry, maps, plans_by_label, band_maps, f"{where} header r{r0 + i}")
            parts.append(k)
            labels["/".join(entry.get("dims") or [entry["dim"]])] = lab
        col_keys.append((parts, labels))

    out = []
    lastr = [None] * nc_h
    for i in range(n_rows):
        rparts, rlabels = [], {}
        if row_keys is not None:
            rk = dict(row_keys[i])
            rparts.append([rk])
            rlabels.update({d: (f"{v[0]}-{v[1]}" if isinstance(v, list) else str(v)) for d, v in rk.items()})
        for j, entry in enumerate(row_dims):
            v = ws.cell(row=data_r0 + i, column=c0 + j).value
            if v in (None, ""):
                v = lastr[j]
            lastr[j] = v
            k, lab = _map_label(v, entry, maps, plans_by_label, band_maps, f"{where} row {data_r0 + i}")
            rparts.append(k)
            rlabels["/".join(entry.get("dims") or [entry["dim"]])] = lab
        for j in range(n_cols):
            cparts, clabels = col_keys[j] if col_keys else ([], {})
            raw = ws.cell(row=data_r0 + i, column=data_c0 + j).value
            ref = f"{sheet_name}!{a1(data_r0 + i, data_c0 + j)}"
            if raw in (None, ""):
                if src.get("allow_blank"):
                    continue
                raise SpecError(f"{where}: blank data cell {ref}")
            status = "VALUE"
            if isinstance(raw, str):
                s = raw.strip()
                if s in tv:
                    status = tv[s]
                    raw = None
                else:
                    try:
                        raw = float(s)
                    except ValueError:
                        raise SpecError(f"{where}: non-numeric cell {ref} = '{raw}'") from None
            else:
                raw = float(raw)
            keys = _cross(rparts + cparts)
            fixed = src.get("fixed", {})
            keys = [{**fixed, **k} for k in keys]
            out.append({"keys": keys, "labels": {**rlabels, **clabels}, "raw": raw, "ref": ref, "status": status,
                        "value": None if raw is None else raw * scale})
    return out


def new_workbook(path: Path, created: _dt.datetime | None = None, title: str | None = None):
    """xlsxwriter workbook with fixed document properties so output bytes are reproducible."""
    import xlsxwriter
    wb = xlsxwriter.Workbook(str(path), {"strings_to_numbers": False, "strings_to_formulas": False})
    wb.set_properties({"created": created or FIXED_CREATED, "title": title or "", "author": "Pricing Workbench",
                       "company": "", "comments": ""})
    wb.set_calc_mode("auto")
    return wb


def make_reproducible(path: Path, when: _dt.datetime | None = None):
    """Rewrite an openpyxl-saved .xlsx so its bytes are reproducible: fixed zip entry times and a fixed
    docProps/core.xml modified/created stamp (openpyxl stamps 'now' on save)."""
    import re
    import zipfile
    when = when or FIXED_CREATED
    stamp = when.strftime("%Y-%m-%dT%H:%M:%SZ")
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rf"\g<1>{stamp}\g<2>", txt)
                txt = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rf"\g<1>{stamp}\g<2>", txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, data)
    tmp.replace(path)
